"""Converter: python-pptx emitter: one WHOLE unit with slide anchors (owner: convert).

With an OCR engine (``PptxConverter(cfg, ocr=engine)``; ``convert/ocr.py``) the text in a deck's pictures
follows each picture's ``[image…]`` line.  Without one nothing here runs: the version, the options and every
page are the ones from before OCR existed.  A failure of the helper is ``OcrError`` with fixed wording, raised
before anything is written: ``convert_file`` then converts the deck with the registry's converter that has
no engine.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from collections.abc import Iterable, Iterator, Mapping, Sequence
from pathlib import Path
from typing import Any, BinaryIO

from agentsync.config import ConvertConfig
from agentsync.convert._common import (
    _FULL_TEXT_SIDECAR,
    _cap_body,
    _cell,
    _check_ooxml_container,
    _dist_version,
    _escape_line,
    _first_line,
    _gfm_table,
    _join_limited,
)
from agentsync.convert.base import OptionValue, make_unit
from agentsync.convert.image import (
    _DOCUMENT_BUDGET_S,
    _OCR_OPTIONS,
    _PICTURE_HEAD,
    _PICTURES_CUT,
    _PICTURES_READ,
    _field_build,
    _raster_left,
    _read_pictures,
    _read_without_ocr,
)
from agentsync.convert.image import _FAILED as _OCR_FAILED
from agentsync.convert.ocr import OcrEngine, OcrError
from agentsync.errors import ConversionError
from agentsync.model import RenderedUnit, UnitKind

log = logging.getLogger(__name__)

_EMITTER_VERSION = "1.0.1"
"""1.0.1: a template's fixed text is de-duplicated a shape at a time (B13 round 5 F77), so the action key of
bytes 1.0.0 cached moves and a copy is converted again (B13 round 6 G59). Nothing is read again by itself."""
_ROW_TOLERANCE_EMU = (
    45_720  # 0.05 in: shapes whose tops differ by less are on one visual row (Docling's rule)
)
_WS_RE = re.compile(r"[ \t\f]+")
_BULLET_PLACEHOLDERS = frozenset({"BODY", "OBJECT"})
_PICTURE_SHAPES = ("Picture", "PlaceholderPicture")
_OCR_RULES = 2
"""Bumped when a rule here that decides which pictures OCR reads, or how a slide shows their text, changes.
It is in the options only with an engine, so it moves no key of a Mac without one (the emitter version
would).  2: a picture's text is one block under its head line, with no blank line in it."""


def _norm(text: str) -> str:
    """Slide text: vertical tabs (soft breaks) to spaces, runs of blanks collapsed, trimmed."""
    return _WS_RE.sub(" ", text.replace("\x0b", " ").replace("\v", " ")).strip()


def _pos(shape: Any) -> tuple[int, int]:
    """(top, left) in EMU; unknown positions sort first (0), keeping XML order via the stable sort."""
    top = getattr(shape, "top", None)
    left = getattr(shape, "left", None)
    return (int(top) if top is not None else 0, int(left) if left is not None else 0)


def _reading_order(shapes: Iterable[Any]) -> list[Any]:
    """Top-to-bottom rows (tops within 0.05 in are one row), left-to-right within a row; stable."""
    items = sorted(enumerate(shapes), key=lambda it: (_pos(it[1])[0], it[0]))
    rows: list[list[tuple[int, Any]]] = []
    row_top: int | None = None
    for i, shape in items:
        top = _pos(shape)[0]
        if row_top is None or top - row_top > _ROW_TOLERANCE_EMU:
            rows.append([])
            row_top = top
        rows[-1].append((i, shape))
    out: list[Any] = []
    for row in rows:
        out.extend(s for _i, s in sorted(row, key=lambda it: (_pos(it[1])[1], it[0])))
    return out


def _alt_text(shape: Any) -> str:
    """The picture's alt text (``descr`` on cNvPr), or ""."""
    try:
        for el in shape.element.iter():
            tag = str(el.tag)
            if tag.endswith("}cNvPr"):
                return _norm(str(el.get("descr") or ""))
    except AttributeError:
        return ""
    return ""


def _placeholder_type(shape: Any) -> str:
    """Placeholder type name (``BODY``, ``TITLE`` ...) or "" for a non-placeholder."""
    if not getattr(shape, "is_placeholder", False):
        return ""
    try:
        ptype = shape.placeholder_format.type
    except (AttributeError, ValueError):
        return ""
    return str(getattr(ptype, "name", ptype) or "")


def _text_frame_lines(shape: Any) -> list[str]:
    """Paragraphs of a text frame: bullets (with level indent) for body placeholders, else plain lines."""
    bullets = _placeholder_type(shape) in _BULLET_PLACEHOLDERS
    lines: list[str] = []
    for para in shape.text_frame.paragraphs:
        text = _norm(str(para.text))
        if not text:
            continue
        if bullets:
            level = int(getattr(para, "level", 0) or 0)
            lines.append("  " * level + "- " + text.replace("<!--", "&lt;!--"))
        else:
            lines.append(_escape_line(text))
    return lines


def _fixed_blocks(shapes: Iterable[Any]) -> list[tuple[str, ...]]:
    """The fixed text of a master's or layout's shapes in reading order, one block of lines per shape: every
    shape that is not a placeholder, and a footer placeholder that carries text.  The other placeholders hold
    the prompt text a slide replaces ("Click to edit Master title style") or a field (date, slide number).
    A group gives its shapes' blocks; a shape with no text gives none."""
    blocks: list[tuple[str, ...]] = []
    for shape in _reading_order(shapes):
        kind = _placeholder_type(shape)
        if getattr(shape, "is_placeholder", False) and kind != "FOOTER":
            continue
        if getattr(shape, "shapes", None) is not None:  # a group
            blocks += _fixed_blocks(shape.shapes)
        elif getattr(shape, "has_table", False):
            blocks.append(tuple(_table_lines(shape.table)))
        elif getattr(shape, "has_text_frame", False):
            blocks.append(tuple(_text_frame_lines(shape)))
    return [b for b in blocks if b]


def _template_lines(prs: Any) -> tuple[list[str], list[str]]:
    """A deck with no slides, usually a template: its layout names in order and the fixed text of its
    masters and layouts (:func:`_fixed_blocks`), each de-duplicated (B13 round 4 Q77).  The text is
    de-duplicated a shape at a time, never a line at a time: a footer every layout repeats is kept once, and
    a second table keeps its own delimiter row and a row an earlier table also has (B13 round 5 F77).  A
    blank line follows each shape's text, so two text boxes are two paragraphs and a table ends before the
    next shape (GFM continues a table on the next line otherwise)."""
    names: list[str] = []
    blocks: list[tuple[str, ...]] = []
    for master in prs.slide_masters:
        found = _fixed_blocks(master.shapes)
        for layout in master.slide_layouts:
            name = _norm(str(layout.name or ""))
            if name and name not in names:
                names.append(name)
            found += _fixed_blocks(layout.shapes)
        for block in found:
            if block not in blocks:
                blocks.append(block)
    return names, [line for block in blocks for line in (*block, "")][:-1]


def _table_lines(table: Any) -> list[str]:
    """A pptx table as GFM (first row as header); merged cells carry the origin cell's text."""
    rows = list(table.rows)
    ncols = len(table.columns)
    grid: list[list[str]] = [["" for _ in range(ncols)] for _ in rows]
    for r, row in enumerate(rows):
        for c, cell in enumerate(row.cells):
            if c >= ncols:
                break
            if getattr(cell, "is_spanned", False):
                continue
            text = _cell(str(cell.text))
            grid[r][c] = text
            if getattr(cell, "is_merge_origin", False):
                for dr in range(int(cell.span_height)):
                    for dc in range(int(cell.span_width)):
                        if r + dr < len(grid) and c + dc < ncols:
                            grid[r + dr][c + dc] = text
    if not grid:
        return []
    return _gfm_table(grid[0], grid[1:])


def _chart_lines(chart: Any) -> list[str]:
    """A chart as a stub line plus its category/series data as a GFM table when python-pptx exposes it."""
    title = ""
    try:
        if chart.has_title:
            title = _norm(str(chart.chart_title.text_frame.text))
    except (AttributeError, ValueError, KeyError):
        title = ""
    kind = ""
    try:
        kind = str(chart.chart_type.name)
    except (AttributeError, ValueError, KeyError):
        kind = ""
    label = ", ".join(x for x in (f'"{title}"' if title else "", kind.lower()) if x)
    lines = [f"[chart: {label}]" if label else "[chart]"]
    try:
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories]
        series = list(plot.series)
        if categories and series:
            header = ["Category", *(_cell(str(s.name or f"Series {i}")) for i, s in enumerate(series, 1))]
            columns = [_series_values(s, len(categories)) for s in series]
            data = []
            for i, cat in enumerate(categories):
                row = [_cell(cat)]
                for values in columns:
                    v = values[i]
                    row.append("" if v is None else _fmt_number(v))
                data.append(row)
            lines += ["", *_gfm_table(header, data)]
    except (AttributeError, IndexError, KeyError, ValueError, TypeError, NotImplementedError):
        pass
    return lines


def _series_values(series: Any, n: int) -> list[float | None]:
    """The series' first ``n`` cached values, read in ONE pass over its points (None where a point is absent).

    python-pptx's ``series.values`` runs one XPath per point, and calling it per table cell made a large
    chart take hours.  The result is what ``values`` returns: the first ``c:pt`` of each index below
    ``c:ptCount`` counts, and a point at or above it is ignored.
    """
    out: list[float | None] = [None] * n
    val = series._element.val
    if val is None:
        return out
    count = val.ptCount_val
    seen: set[int] = set()
    for pt in val.xpath(".//c:pt"):
        try:
            idx = int(pt.get("idx"))
        except (TypeError, ValueError):
            continue
        if not 0 <= idx < count or idx in seen:
            continue
        seen.add(idx)
        value = pt.value  # read past ``n`` too: a point with no value drops the table, as ``values`` did
        if idx < n:
            out[idx] = value
    return out


def _fmt_number(v: object) -> str:
    """Chart values: integral floats print as integers (no ``100.0``)."""
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v)


def _picture_blob(shape: Any) -> bytes | None:
    """The stored bytes of the picture a shape shows; None for a linked picture (its bytes are in another
    file), an empty picture placeholder, and a relationship that leads to no image."""
    try:
        blob = shape.image.blob
    except Exception:  # ValueError: no embedded image; KeyError: no such relationship; AttributeError
        return None
    return blob if isinstance(blob, bytes) else None


def _picture_blobs(shapes: Iterable[Any]) -> Iterator[bytes]:
    """The stored bytes of each picture among ``shapes``, groups included, in the order ``_shape_blocks``
    shows them."""
    for shape in _reading_order(shapes):
        type_name = type(shape).__name__
        if type_name == "GroupShape":
            yield from _picture_blobs(shape.shapes)
        elif type_name in _PICTURE_SHAPES:
            blob = _picture_blob(shape)
            if blob is not None:
                yield blob


def _picture_text(
    engine: OcrEngine, slides: Sequence[Any], work_dir: Path
) -> tuple[dict[bytes, list[str]], bool]:
    """(The escaped lines on-device OCR read in each picture of the deck that holds text, by the picture's
    stored bytes; whether a limit left pictures unread.)

    Each distinct picture is offered to ``image._read_pictures`` once, in the order the slides show them, so
    a logo on every slide is read once and a deck gives the same pictures on every run.  python-pptx has the
    whole package in memory already, so a picture is its part's bytes, not a second copy.  The helper works
    in a folder made inside ``work_dir`` and has ``_DOCUMENT_BUDGET_S`` seconds.

    Raises OcrError when the helper left a picture unread: a deck that kept the other pictures' text would
    look complete."""
    offered: list[bytes] = []

    def streams() -> Iterator[BinaryIO]:
        seen: set[bytes] = set()
        for slide in slides:
            for blob in _picture_blobs(slide.shapes):
                if blob not in seen:
                    seen.add(blob)
                    offered.append(blob)
                    yield io.BytesIO(blob)

    pictures = streams()
    got = _read_pictures(engine, pictures, work_dir=work_dir, budget_s=_DOCUMENT_BUDGET_S)
    if got.unread:  # _read_pictures has logged how many, and the helper's reason
        raise OcrError("the OCR helper left pictures unread")
    texts = {
        blob: got.lines[digest]
        for blob, digest in zip(offered, got.digests, strict=True)
        if digest is not None and digest in got.lines
    }
    return texts, got.over_bytes or _raster_left(pictures)


def _shape_blocks(
    shape: Any, title_id: int | None, texts: dict[bytes, list[str]] | None = None
) -> list[list[str]]:
    """Markdown blocks for one shape (groups recurse in reading order).

    ``texts`` holds the lines OCR read in the deck's pictures.  A picture's lines are taken out of it for
    the first ``[image…]`` line of that picture, so each distinct picture's text is on the page once."""
    if title_id is not None and getattr(shape, "shape_id", None) == title_id:
        return []
    blocks: list[list[str]] = []
    type_name = type(shape).__name__
    if type_name == "GroupShape":
        for child in _reading_order(shape.shapes):
            blocks += _shape_blocks(child, title_id, texts)
        return blocks
    if getattr(shape, "has_table", False):
        lines = _table_lines(shape.table)
        if lines:
            blocks.append(lines)
        return blocks
    if getattr(shape, "has_chart", False):
        blocks.append(_chart_lines(shape.chart))
        return blocks
    if type_name in _PICTURE_SHAPES:
        alt = _alt_text(shape)
        block = [f"[image: {alt}]" if alt else "[image]"]
        if texts:  # never without an engine, nor once every picture's text is on the page
            blob = _picture_blob(shape)
            if blob is not None and blob in texts:
                # One block, with no blank line in it: what follows a blank line is the next shape's own
                # text, and a second paragraph of the picture's would read as that.
                block += [_PICTURE_HEAD, *(line for line in texts.pop(blob) if line)]
        blocks.append(block)
        return blocks
    if type_name == "Movie":
        blocks.append(["[media]"])
        return blocks
    if getattr(shape, "has_text_frame", False):
        lines = _text_frame_lines(shape)
        if lines:
            bullets = _placeholder_type(shape) in _BULLET_PLACEHOLDERS
            if bullets:
                blocks.append(lines)
            else:
                blocks.extend([ln] for ln in lines)
    return blocks


_CONTENT_TYPES = "[Content_Types].xml"
_TEMPLATE_MAIN = b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
_DECK_MAIN = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


def _template_as_deck(src: Path) -> io.BytesIO | None:
    """The package of a PowerPoint template (``.potx``) with its main part typed as a deck's, in memory.

    python-pptx opens only the deck, macro-deck and slideshow main types and raises on a template's, though
    the parts are the same.  Every member is copied as it is and only ``[Content_Types].xml`` changes.  None
    when that part does not carry the template type: the file is opened as it is and python-pptx's own
    error stands."""
    with zipfile.ZipFile(src) as zin:
        try:
            types = zin.read(_CONTENT_TYPES)
        except KeyError:
            return None
        if _TEMPLATE_MAIN not in types:
            return None
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                deck = info.filename == _CONTENT_TYPES
                zout.writestr(info, types.replace(_TEMPLATE_MAIN, _DECK_MAIN) if deck else zin.read(info))
    buffer.seek(0)
    return buffer


class PptxConverter:
    """Python-pptx emitter: one WHOLE unit with slide anchors.

    ``<!-- Slide number: N -->`` anchor per slide, title as ``## ``, text frames in reading order, tables as
    GFM, speaker notes under ``### Notes:``.  Pictures become ``[image: <alt text>]`` lines.

    With an OCR engine the text read in a picture follows its ``[image…]`` line, under ``[text in the image
    above, read by on-device OCR (Apple Vision):]``, once per distinct picture.  When the engine fails,
    ``convert`` raises OcrError with fixed wording and ``convert_file`` converts the deck without OCR.
    """

    converter_id = "pptx-python-pptx"
    extensions: tuple[str, ...] = (".pptx", ".potx")  # a template is read as a deck (_template_as_deck)

    def __init__(self, cfg: ConvertConfig, ocr: OcrEngine | None = None) -> None:
        """Bind converter options from config, and the cycle's OCR engine when there is one."""
        self._cfg = cfg
        self._ocr = ocr

    def version(self) -> str:
        """Version as run (emitter version + underlying library/tool version); with an OCR engine its
        identity comes last.  Without one this is the version from before OCR existed."""
        version = f"{_EMITTER_VERSION}+python-pptx-{_dist_version('python-pptx')}"
        return version if self._ocr is None else f"{version}+{self._ocr.identity}"

    def options(self) -> Mapping[str, OptionValue]:
        """Normalised options hashed into options_hash; the OCR ones only with an engine."""
        opts: dict[str, OptionValue] = {
            "max_page_bytes": self._cfg.max_page_bytes,
            "row_tolerance_emu": _ROW_TOLERANCE_EMU,
        }
        if self._ocr is not None:
            opts.update(_OCR_OPTIONS)
            opts["ocr_pptx_rules"] = _OCR_RULES
        return opts

    def outdated(self, produced: str, reason: str | None = None) -> bool:
        """True when the page this converter wrote under version ``produced`` is worth reading the deck
        again for: it was written without OCR and there is an engine now, or it was written by the field
        build of OCR (``image._field_build``; engine or not).  Never for a stub (``reason`` is the
        stub's)."""
        if reason is not None:
            return False
        return _field_build(produced) or (self._ocr is not None and _read_without_ocr(produced))

    def convert(self, src: Path, *, name: str) -> tuple[RenderedUnit, ...]:
        """Convert one staged file; see the Converter protocol for pre/postconditions and errors.

        With an OCR engine the helper works in a folder made beside ``src``, and OcrError (fixed wording)
        means the engine failed on a deck that converts without it."""
        _check_ooxml_container(src)
        from pptx import Presentation  # noqa: PLC0415 - heavy import, only when a deck is converted

        try:
            template = _template_as_deck(src) if name.lower().endswith(".potx") else None
            prs = Presentation(template if template is not None else str(src))
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise ConversionError(f"python-pptx cannot read the deck: {type(exc).__name__}: {exc}") from exc
        except Exception as exc:  # PackageNotFoundError, lxml XMLSyntaxError
            raise ConversionError(f"python-pptx cannot read the deck: {type(exc).__name__}: {exc}") from exc
        out: list[str] = []
        titles: list[str] = []
        texts: dict[bytes, list[str]] = {}  # stays empty without an engine
        cut = False
        try:
            slides = list(prs.slides)
            if self._ocr is not None:
                try:
                    texts, cut = _picture_text(self._ocr, slides, src.parent)
                except Exception as exc:  # OcrError, or whatever else the pass let through
                    # What went wrong goes to the log, never into a page or a reason, and nothing half-read
                    # is returned: convert_file converts the deck again without OCR.
                    detail = str(exc) if isinstance(exc, OcrError) else type(exc).__name__
                    log.warning("%s: on-device OCR failed: %s", name, detail)
                    raise OcrError(_OCR_FAILED) from None
            found = len(texts)
            for n, slide in enumerate(slides, start=1):
                out += self._slide(n, slide, titles, texts)
        except ConversionError:
            raise
        except Exception as exc:
            raise ConversionError(f"cannot read slide content: {type(exc).__name__}: {exc}") from exc
        layouts: list[str] = []
        if not slides:
            # A template (or an empty deck): what its layouts and master say, since no slide does.
            try:
                layouts, fixed = _template_lines(prs)
            except Exception as exc:
                raise ConversionError(f"cannot read the slide layouts: {type(exc).__name__}: {exc}") from exc
            out = ["[empty presentation: no slides]"]
            if layouts:
                out += ["", _escape_line("Layouts: " + _join_limited(layouts, limit=600))]
            if fixed:
                out += ["", "Fixed text in the master and layouts:", "", *fixed]
            # The first fixed text that can be a title (never a table row), else the first layout name.
            first = _first_line("\n".join(fixed)) or (layouts[0][:120] if layouts else "")
            titles.append(first or "Untitled presentation")  # never the marker
        body = "\n".join(out)
        # A picture's text follows its own ``[image…]`` line, so the first line is never one OCR read.
        title = next((t for t in titles if t), "") or _first_line(body) or "Untitled presentation"
        named = [t for t in titles if t] if slides else []
        summary = f"Presentation: {len(slides)} slide(s)"
        if not slides:
            summary += f"; {len(layouts)} layout(s)"
        # Counts only: a summary is front matter, above the banner, and never holds text OCR read.
        if found > len(texts):
            summary += "; " + _PICTURES_READ.format(found - len(texts))
        if cut:
            summary += "; " + _PICTURES_CUT
        if named:
            summary += "; titles: " + _join_limited(named, limit=170)
        body, sidecars = _cap_body(body, self._cfg.max_page_bytes, sidecar_name=_FULL_TEXT_SIDECAR)
        return (
            make_unit(
                unit_id="whole",
                kind=UnitKind.WHOLE,
                index=0,
                of=1,
                name="",
                file_stem="",
                title=title,
                summary=summary,
                body=body,
                sidecars=sidecars,
            ),
        )

    def _slide(self, n: int, slide: Any, titles: list[str], texts: dict[bytes, list[str]]) -> list[str]:
        """Markdown lines for one slide (anchor, title, blocks, notes); ``texts`` as for ``_shape_blocks``."""
        out = [f"<!-- Slide number: {n} -->", ""]
        hidden = str(slide.element.get("show", "1")) in ("0", "false")
        title_shape = slide.shapes.title
        title = (
            _norm(str(title_shape.text_frame.text))
            if title_shape is not None and title_shape.has_text_frame
            else ""
        )
        titles.append(title)
        if title:
            out += [f"## {title.replace('<!--', '&lt;!--')}", ""]
        if hidden:
            out += ["[hidden slide]", ""]
        title_id = None if title_shape is None else int(title_shape.shape_id)
        for shape in _reading_order(slide.shapes):
            for block in _shape_blocks(shape, title_id, texts):
                out += [*block, ""]
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = [] if frame is None else [_escape_line(_norm(str(p.text))) for p in frame.paragraphs]
            notes = [x for x in notes if x]
            if notes:
                out += ["### Notes:", "", *notes, ""]
        return out
