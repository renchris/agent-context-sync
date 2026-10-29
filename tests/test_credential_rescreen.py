"""N5: credential stubs are screened again when the secret-scan rules change (CONTRACTS 16.39, N5)."""

from __future__ import annotations

import dataclasses
import hashlib
import inspect
import json
import os
import subprocess
from collections.abc import Sequence
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import Workbook

from agentsync import arm_local, cli, cycle, governance, lints
from agentsync import materialise as mat
from agentsync.config import Config
from agentsync.convert import pdf as pdf_mod
from agentsync.manifest import Manifest
from agentsync.model import ConversionStatus, LintFinding, RowState
from agentsync.ops.lock import read_heartbeat
from test_convert_builders import build_annotated_pdf, build_picture_pdf, page_picture, shade_engine
from test_convert_image import picture
from test_ocr import fake_engine
from test_review_fixes import (
    AWS_KEY,
    committed_blobs_containing,
    git,
    item_id,
    make_env,
    page_fm,
    queued,
    run,
)

INVITE = (
    "Contoso weekly sync\nJoin: https://teams.microsoft.com/l/meetup-join/19%3ameeting_x?pwd=AbCdEfGh1234\n"
)


@pytest.fixture(autouse=True)
def _builtin_scan(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(lints, "_gitleaks", lambda: None)  # the builtin rules, whatever this Mac has


def _old_rules(monkeypatch: pytest.MonkeyPatch) -> None:
    """The rules before de081ae: a Teams join link's pwd= counts as a password."""
    monkeypatch.setattr(lints, "_strip_teams_join_pwd", lambda line: line)
    monkeypatch.setattr(lints, "SECRET_RULES_VERSION", 0, raising=False)


def _state(config, name: str) -> tuple[RowState, str | None]:  # type: ignore[no-untyped-def]
    with Manifest(config.state_paths.db) as m:
        row = m.get_item("src", item_id(config, name))
        assert row is not None
        return row.state, row.state_reason


def _count_screens(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    """The page count of each pre-write screen a run makes (``_Cycle._still_credential``)."""
    screened: list[int] = []
    real = cycle._Cycle._still_credential

    def counted(self: cycle._Cycle, pages: Sequence[cycle.PlannedPage]) -> bool:
        screened.append(len(pages))
        return real(self, pages)

    monkeypatch.setattr(cycle._Cycle, "_still_credential", counted)
    return screened


def test_a_false_credential_stub_is_published_once_the_rules_change(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, src = make_env(tmp_path)
    (src / "invite.txt").write_text(INVITE, encoding="utf-8")
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _old_rules(mp)
        assert run(config).exit_code == 0
    assert _state(config, "invite.txt") == (RowState.QUARANTINED, "contains a credential")
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert _state(config, "invite.txt") == (RowState.LIVE, None)
    assert "Contoso weekly sync" in (config.docs_repo / "mirror/src/invite.txt.md").read_text(
        encoding="utf-8"
    )
    assert page_fm(config, "mirror/src/keys.txt.md")["reason"] == "contains a credential"
    assert not committed_blobs_containing(config.docs_repo, AWS_KEY.encode())
    assert [(c.op.value, c.path) for c in again.changes] == [("M", "mirror/src/invite.txt.md")]
    assert run(config).commit_sha is None  # once: the next run re-screens nothing


def test_a_stub_screened_again_and_stubbed_again_is_no_commit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, src = make_env(tmp_path)
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _old_rules(mp)
        run(config)
    head = git(config.docs_repo, "rev-parse", "HEAD")
    screened = _count_screens(monkeypatch)
    again = run(config)
    assert screened == [1]  # the stub was converted and screened again, not left alone
    assert again.exit_code == 0 and again.commit_sha is None and again.changes == ()
    assert git(config.docs_repo, "rev-parse", "HEAD") == head
    assert _state(config, "keys.txt") == (RowState.QUARANTINED, "contains a credential")


def test_a_run_without_the_gitleaks_of_the_last_screen_screens_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, src = make_env(tmp_path)
    (src / "invite.txt").write_text(INVITE, encoding="utf-8")
    with monkeypatch.context() as mp:
        _old_rules(mp)
        mp.setattr(lints, "secret_rules_identity", lambda: {"rules": "old", "gitleaks": "/x/gitleaks:1:1"})
        run(config)
    monkeypatch.setattr(lints, "secret_rules_identity", lambda: {"rules": "new", "gitleaks": ""})
    assert run(config).commit_sha is None  # launchd's PATH: no gitleaks, so no re-screen
    assert _state(config, "invite.txt") == (RowState.QUARANTINED, "contains a credential")
    monkeypatch.setattr(
        lints, "secret_rules_identity", lambda: {"rules": "new", "gitleaks": "/x/gitleaks:1:1"}
    )
    assert run(config).commit_sha is not None
    assert _state(config, "invite.txt") == (RowState.LIVE, None)


def test_an_online_only_stub_waits_and_is_screened_once_it_is_on_this_mac(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, src = make_env(tmp_path)
    invite = src / "invite.txt"
    invite.write_text(INVITE, encoding="utf-8")
    ino, real = invite.stat().st_ino, mat.is_dataless
    with monkeypatch.context() as mp:
        _old_rules(mp)
        run(config)
        mp.setattr(arm_local, "is_dataless", lambda st: st.st_ino == ino or real(st))
        mp.setattr(mat, "is_dataless", lambda st: st.st_ino == ino or real(st))
        run(config, budget_bytes=0)  # evicted under the old rules: the walk records it online-only
    with monkeypatch.context() as mp:
        mp.setattr(arm_local, "is_dataless", lambda st: st.st_ino == ino or real(st))
        mp.setattr(mat, "is_dataless", lambda st: st.st_ino == ino or real(st))
        run(config)  # the rules moved: it waits, and nothing is downloaded although the budget allows it
        with Manifest(config.state_paths.db) as m:
            row = m.get_item("src", item_id(config, "invite.txt"))
            assert row is not None and row.content_sha256 is not None  # not marked
        assert _state(config, "invite.txt") == (RowState.QUARANTINED, "contains a credential")
        with Manifest(config.state_paths.db) as m:
            assert item_id(config, "invite.txt") in (m.get_meta("secret_rescreen_waiting") or "")
    run(config)  # on this Mac again: the walk sees it
    run(config)  # and the re-screen reads it
    assert _state(config, "invite.txt") == (RowState.LIVE, None)


def test_a_published_page_the_new_rules_match_becomes_a_stub_and_queues_a_purge(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, src = make_env(tmp_path)
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        mp.setattr(
            lints, "SECRET_PATTERNS", tuple(r for r in lints.SECRET_PATTERNS if r[0] != "aws-access-key")
        )
        run(config)
    assert _state(config, "keys.txt") == (RowState.LIVE, None)
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert page_fm(config, "mirror/src/keys.txt.md")["reason"] == "contains a credential"
    assert AWS_KEY not in git(config.docs_repo, "show", "HEAD:mirror/src/keys.txt.md")
    assert queued(config) == [(item_id(config, "keys.txt"), "dlp-remediation")]
    alarms = [a for s in again.sources for a in s.alarms]
    hit = [a for a in alarms if a.startswith("a published page matches the secret scan's new rules")]
    # B13 R46/R47/R57: the alarm that queued a purge ends with the preview first (§16.38)
    assert hit and all(a.endswith(cycle._PURGE_QUEUE_STEPS) for a in hit)


def _without_aws(mp: pytest.MonkeyPatch) -> None:
    """Earlier rules: the aws-access-key pattern is not one of them."""
    mp.setattr(lints, "SECRET_PATTERNS", tuple(r for r in lints.SECRET_PATTERNS if r[0] != "aws-access-key"))


def _purged_with_alarm(config, report, name: str) -> None:  # type: ignore[no-untyped-def]
    assert queued(config) == [(item_id(config, name), "dlp-remediation")]
    alarms = [a for s in report.sources for a in s.alarms]
    assert any(a.startswith("a published page matches the secret scan's new rules") for a in alarms)


def test_a_page_edited_in_the_run_the_rules_move_still_queues_the_purge_of_its_earlier_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R48: the run that rewrites a page is the run whose rules moved; the committed text, which the new
    rules match, is read from HEAD, so its purge is queued although the page is not read as published."""
    config, src = make_env(tmp_path)
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _without_aws(mp)
        run(config)
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\nrotated by Contoso ops\n", encoding="utf-8")
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert page_fm(config, "mirror/src/keys.txt.md")["reason"] == "contains a credential"
    _purged_with_alarm(config, again, "keys.txt")


def test_a_page_read_again_in_the_run_the_rules_move_still_queues_the_purge_of_its_earlier_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R48: an update usually moves the rules and an emitter at once, so the rescan run reads the file
    again (CONTRACTS 16.27) and rewrites its page; the committed page still gets its purge."""
    config, src = make_env(tmp_path)
    note = "/Subtype /Text /Rect [400 700 420 720] /T (Contoso reviewer) /Contents (Check the key)"
    build_annotated_pdf(src / "keys.pdf", [([f"deploy key {AWS_KEY}"], [note])])
    with monkeypatch.context() as mp:
        _without_aws(mp)
        mp.setattr(pdf_mod, "_EMITTER_VERSION", "2.0.0")
        mp.setattr(pdf_mod, "_render_comments", lambda _found: [])
        run(config)
    assert _state(config, "keys.pdf") == (RowState.LIVE, None)
    assert committed_blobs_containing(config.docs_repo, AWS_KEY.encode())
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert page_fm(config, "mirror/src/keys.pdf.md")["reason"] == "contains a credential"
    _purged_with_alarm(config, again, "keys.pdf")


_HISTORY_ALARM = (
    "a published page matches the secret scan's new rules: the text committed before this run does"
)


def test_a_key_the_edit_removes_in_the_run_the_rules_move_still_queues_the_purge_of_its_earlier_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R48 fix pass: the edit removes the key, so the page written in the run the rules move is clean and
    published, but the committed text holds the key: its purge is queued, with an alarm that names no file."""
    config, src = make_env(tmp_path)
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\nowner: Contoso ops\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _without_aws(mp)
        run(config)
    (src / "keys.txt").write_text("deploy key rotated\nowner: Contoso ops\n", encoding="utf-8")
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert _state(config, "keys.txt") == (RowState.LIVE, None)
    assert queued(config) == [(item_id(config, "keys.txt"), "dlp-remediation")]
    alarms = [a for s in again.sources for a in s.alarms if a.startswith(_HISTORY_ALARM)]
    assert alarms and all(a.endswith(cycle._PURGE_QUEUE_STEPS) and "keys" not in a for a in alarms)


def test_a_file_deleted_in_the_run_the_rules_move_still_queues_the_purge_of_its_earlier_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R48 fix pass: the run the rules move is the run that turns the page into a deletion notice; the
    committed page held the key, so its purge is queued as a credential's."""
    config, src = make_env(tmp_path)
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _without_aws(mp)
        run(config)
        stable = item_id(config, "keys.txt")
        (src / "keys.txt").unlink()
        run(config)  # the first absence marks it; the second, in the next run, removes it
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert page_fm(config, "mirror/src/keys.txt.md")["status"] == "deleted"
    assert (stable, "dlp-remediation") in queued(config)
    assert any(a.startswith(_HISTORY_ALARM) for s in again.sources for a in s.alarms)


def test_a_key_the_edit_brings_in_the_run_the_rules_move_is_a_stub_with_no_purge(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R48: the committed text is clean and the edit brings the key: the page is a stub, but no history
    holds the key, so no purge is queued and no purge alarm is raised."""
    config, src = make_env(tmp_path)
    (src / "keys.txt").write_text("deploy key to come\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _without_aws(mp)
        run(config)
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert page_fm(config, "mirror/src/keys.txt.md")["reason"] == "contains a credential"
    assert queued(config) == []
    alarms = [a for s in again.sources for a in s.alarms]
    assert not any(a.startswith("a published page matches the secret scan's new rules") for a in alarms)
    assert not committed_blobs_containing(config.docs_repo, AWS_KEY.encode())


def test_an_archive_page_the_new_rules_match_is_deleted_and_queues_a_purge(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R74: the ``archive/`` copy of a deleted page ([governance] archive) is scanned again when the rules
    move.  A hit deletes the copy and its sidecars, the tombstone stops naming it, and a purge of the item's
    history is queued, named by the archive page's own ids."""
    config, src = make_env(tmp_path)
    with config.config_path.open("a", encoding="utf-8") as fh:
        fh.write("\n[governance]\narchive = true\n")
    repo = config.docs_repo
    page, kept = "mirror/src/keys.txt.md", "archive/src/keys.txt.md"
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _without_aws(mp)
        run(config)
        (src / "keys.txt").unlink()
        for _ in range(2):  # a local file must be absent from two complete passes
            run(config)
        assert (repo / kept).is_file() and f"kept, searchable, at {kept}" in (repo / page).read_text("utf-8")
        (repo / "archive/src/keys.txt.files").mkdir()  # a text sidecar goes with its page
        (repo / "archive/src/keys.txt.files/notes.txt").write_text("rotate quarterly\n", encoding="utf-8")
        git(repo, "add", "archive")
        git(repo, "commit", "-q", "-m", "sidecar")
    durable = page_fm(config, kept)["stable_id"]
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert not (repo / kept).exists() and not (repo / "archive/src/keys.txt.files").exists()
    assert git(repo, "ls-files", "--", "archive").strip() == ""
    assert "kept, searchable" not in (repo / page).read_text(encoding="utf-8")
    assert queued(config) == [(durable, "dlp-remediation")]
    alarms = [a for s in again.sources for a in s.alarms]
    hit = [a for a in alarms if a.startswith("a published page matches the secret scan's new rules")]
    assert hit and all(a.endswith(cycle._PURGE_QUEUE_STEPS) and "keys" not in a for a in hit)


def test_an_archive_pages_secret_hit_keeps_its_secret_finding(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 round 30 ruling 37 (C328): the first scan's findings are dropped for a page stubbed or redacted,
    which its alarm says; an ``archive/`` copy is neither (it is deleted and its purge queued), so its
    SECRET finding is kept."""
    monkeypatch.setattr(lints, "_gitleaks", lambda: None)
    config, src = make_env(tmp_path)
    with config.config_path.open("a", encoding="utf-8") as fh:
        fh.write("\n[governance]\narchive = true\n")
    kept = "archive/src/keys.txt.md"
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _without_aws(mp)
        run(config)
        (src / "keys.txt").unlink()
        for _ in range(2):  # a local file must be absent from two complete passes
            run(config)
        assert (config.docs_repo / kept).is_file()
    again = run(config)
    assert again.exit_code == 0
    secret = [f for f in again.lint_findings if f.code == "SECRET"]
    assert [f.path for f in secret] == [kept], again.lint_findings
    assert secret[0].message.startswith("aws-access-key at line ")


def test_a_published_page_stubbed_again_names_its_page_and_rules_in_its_alarm(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C328: a published page the new rules match is stubbed, so its first-scan findings are dropped, and its
    one alarm, the purge's, named neither the page nor the rule. It names the page as it is after any
    redaction (hashed when its name carries the key, never that name) and the rule, after the fixed words
    the report and the tests match on, and no SECRET finding is left."""
    monkeypatch.setattr(lints, "_gitleaks", lambda: None)
    config, src = make_env(tmp_path)
    (src / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    keyed = f"{AWS_KEY}.txt"
    (src / keyed).write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    with monkeypatch.context() as mp:
        _without_aws(mp)
        run(config)
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert not [f for f in again.lint_findings if f.code == "SECRET"], again.lint_findings
    alarms = [a for s in again.sources for a in s.alarms]
    head = "a published page matches the secret scan's new rules: "
    hit = sorted(a for a in alarms if a.startswith(head))
    assert len(hit) == 2, alarms
    said = "is a `contains a credential` stub now: its converted text held one (aws-access-key); "
    assert [a for a in hit if a.startswith(f"{head}mirror/src/keys.txt.md {said}")], hit
    (hashed,) = [a for a in hit if "mirror/src/keys.txt.md" not in a]
    assert hashed.startswith(f"{head}mirror/src/redacted-") and f".md {said}" in hashed, hashed
    assert AWS_KEY.lower() not in hashed.lower()
    assert all(a.endswith(cycle._PURGE_QUEUE_STEPS) for a in hit)


def test_the_rules_identity_moves_with_a_pattern_the_version_and_gitleaks(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    base = lints.secret_rules_identity()
    assert base["gitleaks"] == ""
    with monkeypatch.context() as mp:
        mp.setattr(lints, "SECRET_PATTERNS", lints.SECRET_PATTERNS[1:])
        assert lints.secret_rules_identity()["rules"] != base["rules"]
    with monkeypatch.context() as mp:
        mp.setattr(lints, "SECRET_RULES_VERSION", lints.SECRET_RULES_VERSION + 1)
        assert lints.secret_rules_identity()["rules"] != base["rules"]
    exe = tmp_path / "gitleaks"
    exe.write_bytes(b"#!/bin/sh\n")
    monkeypatch.setattr(lints, "_gitleaks", lambda: exe)
    assert lints.secret_rules_identity()["gitleaks"].startswith(str(exe.resolve()))


_SCAN_CODE_SHA = "aacd4ee4dc5903f0e9b508ac32903385df9bfc542696161918b927c7b79f969a"
"""sha256 of the source of the functions that decide a secret-scan hit.  A change there that decides
differently with the same patterns must bump ``lints.SECRET_RULES_VERSION`` (every install then screens its
credential stubs again); a change that decides the same (a comment, a docstring) only updates this pin."""


def test_a_change_to_the_scan_code_bumps_the_rules_version() -> None:
    code = "".join(
        inspect.getsource(f)
        for f in (
            lints._strip_teams_join_pwd,
            lints._builtin_secret_scan,
            lints.lint_secrets,
            cycle._Cycle._still_credential,
        )
    )
    digest = hashlib.sha256(f"{lints.SECRET_RULES_VERSION}\n{code}".encode()).hexdigest()
    assert digest == _SCAN_CODE_SHA, (
        "the secret scan's code changed: if it decides differently now, bump lints.SECRET_RULES_VERSION; "
        "then set _SCAN_CODE_SHA to " + digest
    )


def test_a_sync_from_the_shell_records_the_gitleaks_its_path_finds(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R75: the next LaunchAgent run, whose PATH is launchd's, uses the gitleaks this run's PATH found."""
    config, _ = make_env(tmp_path)
    bin_dir = tmp_path / "shell-bin"
    bin_dir.mkdir()
    exe = bin_dir / "gitleaks"
    exe.write_text("#!/bin/sh\nexit 0\n")
    exe.chmod(0o755)
    monkeypatch.setenv("PATH", f"{bin_dir}:{os.environ['PATH']}")
    assert run(config).exit_code == 0
    assert lints.recorded_gitleaks() == exe


PASSWORD_PAGE = "Contoso VPN\nUsername: jdoe\nPassword: Welcome123\n"


def _blind_gitleaks(mp: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """A gitleaks that finds nothing: 8.30.1 finds nothing in PASSWORD_PAGE (measured, B13 N5 skeptic)."""
    exe = tmp_path / "gitleaks"
    exe.write_bytes(b"#!/bin/sh\n")

    def scan(_exe: Path, _repo: Path, _paths: Sequence[str]) -> list[LintFinding]:
        return []

    mp.setattr(lints, "_gitleaks", lambda: exe)
    mp.setattr(lints, "_gitleaks_scan", scan)


def test_the_builtin_patterns_read_every_page_when_gitleaks_scans(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R76: gitleaks misses a plain password, so the builtin patterns read every page beside it."""
    (tmp_path / "vpn.md").write_text(PASSWORD_PAGE, encoding="utf-8")
    (tmp_path / "other.md").write_text(PASSWORD_PAGE, encoding="utf-8")
    assert [f.path for f in lints.lint_secrets(tmp_path, ["vpn.md"])] == ["vpn.md"]  # builtin catches it
    _blind_gitleaks(monkeypatch, tmp_path)
    both = lints.lint_secrets(tmp_path, ["vpn.md", "other.md"])
    assert [(f.path, "generic-password" in f.message) for f in both] == [("other.md", True), ("vpn.md", True)]


def test_a_new_file_with_a_password_gitleaks_misses_is_a_credential_stub(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R76: a file never stubbed before, on a Mac with gitleaks, is read by the builtin patterns too."""
    config, src = make_env(tmp_path)
    _blind_gitleaks(monkeypatch, tmp_path)
    (src / "vpn.txt").write_text(PASSWORD_PAGE, encoding="utf-8")
    assert run(config).exit_code == 0
    assert _state(config, "vpn.txt") == (RowState.QUARANTINED, "contains a credential")
    assert not committed_blobs_containing(config.docs_repo, b"Welcome123")


SHORTCUT = "[InternetShortcut]\nURL=https://jdoe:NotARealPassw0rd@files.contoso.example/q.docx\n"


def test_a_url_shortcut_with_a_password_is_a_credential_stub(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 round 28 I40: a ``.url`` whose target carries ``user:password@`` is withheld as a credential, on a
    Mac whose gitleaks misses it too; the password is in no committed file."""
    config, src = make_env(tmp_path)
    _blind_gitleaks(monkeypatch, tmp_path)
    (src / "q.url").write_text(SHORTCUT, encoding="utf-8")
    assert run(config).exit_code == 0
    assert _state(config, "q.url") == (RowState.QUARANTINED, "contains a credential")
    assert not committed_blobs_containing(config.docs_repo, b"NotARealPassw0rd")
    assert not [
        p for p in config.docs_repo.rglob("*") if p.is_file() and b"NotARealPassw0rd" in p.read_bytes()
    ]


def _cached_with(config: Config, secret: bytes) -> list[Path]:
    """Every file under the converter cache whose bytes hold ``secret``."""
    root = config.cache_dir
    return [p for p in root.rglob("*") if p.is_file() and secret in p.read_bytes()] if root.exists() else []


@pytest.mark.parametrize(
    ("name", "text", "secret"),
    [
        ("credential.url", SHORTCUT, b"NotARealPassw0rd"),
        ("keys.txt", f"Contoso staging\naws_access_key_id = {AWS_KEY}\n", AWS_KEY.encode()),
    ],
)
def test_a_first_conversion_stubbed_for_a_credential_leaves_no_cached_text_c625(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, name: str, text: str, secret: bytes
) -> None:
    """B13 round 32 C625: a file whose first conversion held a credential became a stub, but its converted
    text, the made-up password or key included, stayed in the converter cache (0600) until a reconcile's
    gc, since the stub's output keeps no action key. The stub deletes it at once, as the held-once path
    and a published page's alarm do; a second sync leaves the stub and the cache clean."""
    config, src = make_env(tmp_path)
    _blind_gitleaks(monkeypatch, tmp_path)
    (src / name).write_text(text, encoding="utf-8")
    assert run(config).exit_code == 0
    assert _state(config, name) == (RowState.QUARANTINED, "contains a credential")
    assert _cached_with(config, secret) == []
    assert run(config).exit_code == 0
    assert _state(config, name) == (RowState.QUARANTINED, "contains a credential")
    assert _cached_with(config, secret) == []


def test_a_url_shortcut_published_before_the_url_credential_rule_becomes_a_stub(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 round 28 I40: the new rule moves the rules digest, so a page published before it is read again and
    stubbed, with its purge queued, and no SECRET_RULES_VERSION bump."""
    config, src = make_env(tmp_path)
    (src / "q.url").write_text(SHORTCUT, encoding="utf-8")
    with monkeypatch.context() as mp:
        mp.setattr(
            lints, "SECRET_PATTERNS", tuple(r for r in lints.SECRET_PATTERNS if r[0] != "url-credential")
        )
        run(config)
    assert _state(config, "q.url") == (RowState.LIVE, None)
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    assert page_fm(config, "mirror/src/q.url.md")["reason"] == "contains a credential"
    assert "NotARealPassw0rd" not in git(config.docs_repo, "show", "HEAD:mirror/src/q.url.md")
    _purged_with_alarm(config, again, "q.url")


def test_a_copy_of_a_credential_stub_under_a_new_name_is_a_stub_too(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """B13 R76: a copy is a new file id, which the once-a-credential record does not name; the scan of
    every new page still holds it, once gitleaks has arrived."""
    config, src = make_env(tmp_path)
    (src / "vpn.txt").write_text(PASSWORD_PAGE, encoding="utf-8")
    assert run(config).exit_code == 0  # the builtin patterns: no gitleaks yet
    stub = (RowState.QUARANTINED, "contains a credential")
    assert _state(config, "vpn.txt") == stub
    _blind_gitleaks(monkeypatch, tmp_path)
    (src / "vpn-copy.txt").write_text(PASSWORD_PAGE, encoding="utf-8")
    assert run(config).exit_code == 0
    assert item_id(config, "vpn-copy.txt") != item_id(config, "vpn.txt")
    assert _state(config, "vpn-copy.txt") == stub and _state(config, "vpn.txt") == stub
    assert not committed_blobs_containing(config.docs_repo, b"Welcome123")


@pytest.mark.parametrize("first_run_on_build", [False, True])
def test_a_password_the_builtin_rules_stubbed_stays_a_stub_when_gitleaks_screens_it(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, first_run_on_build: bool
) -> None:
    config, src = make_env(tmp_path)
    (src / "vpn.txt").write_text(PASSWORD_PAGE, encoding="utf-8")
    (src / "invite.txt").write_text(INVITE, encoding="utf-8")
    with monkeypatch.context() as mp:
        mp.setattr(lints, "_gitleaks", lambda: None)  # a launchd run, or before gitleaks was installed
        mp.setattr(lints, "_strip_teams_join_pwd", lambda line: line)  # the rules before de081ae
        mp.setattr(lints, "SECRET_RULES_VERSION", 0)
        assert run(config).exit_code == 0
    assert _state(config, "vpn.txt") == (RowState.QUARANTINED, "contains a credential")
    if first_run_on_build:  # a manifest from before this build: no identity stored
        with Manifest(config.state_paths.db) as m:
            m.set_meta("secret_rules", "")
    _blind_gitleaks(monkeypatch, tmp_path)
    screened = _count_screens(monkeypatch)
    again = run(config)
    assert again.exit_code == 0
    assert sorted(screened) == [1, 1]  # both stubs were converted and screened again
    assert _state(config, "vpn.txt") == (RowState.QUARANTINED, "contains a credential")
    assert not committed_blobs_containing(config.docs_repo, b"Welcome123")
    assert _state(config, "invite.txt") == (RowState.LIVE, None)  # the de081ae fix still reaches it
    assert [(c.op.value, c.path) for c in again.changes] == [("M", "mirror/src/invite.txt.md")]


def test_a_workbook_stub_screened_again_and_stubbed_again_is_no_commit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(lints, "_gitleaks", lambda: None)
    config, src = make_env(tmp_path)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Keys"
    ws["A1"], ws["B1"] = "deploy key", AWS_KEY
    wb.create_sheet("Notes")["A1"] = "Contoso notes"
    wb.save(src / "book.xlsx")
    assert run(config).exit_code == 0
    assert _state(config, "book.xlsx") == (RowState.QUARANTINED, "contains a credential")
    head = git(config.docs_repo, "rev-parse", "HEAD")
    monkeypatch.setattr(lints, "SECRET_RULES_VERSION", lints.SECRET_RULES_VERSION + 1)  # the rules move
    screened = _count_screens(monkeypatch)
    again = run(config)
    assert len(screened) == 1 and screened[0] > 1  # every sheet's page was screened, before a write
    assert again.exit_code == 0 and again.commit_sha is None and again.changes == ()
    assert git(config.docs_repo, "rev-parse", "HEAD") == head
    assert _state(config, "book.xlsx") == (RowState.QUARANTINED, "contains a credential")
    with Manifest(config.state_paths.db) as m:
        row = m.get_item("src", item_id(config, "book.xlsx"))
        assert row is not None and row.content_sha256 is not None  # settled again: the next run is quiet
    assert run(config).commit_sha is None and len(screened) == 1


def test_a_run_without_the_gitleaks_of_the_last_screen_leaves_the_waiting_list(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, src = make_env(tmp_path)
    (src / "invite.txt").write_text(INVITE, encoding="utf-8")
    monkeypatch.setattr(lints, "_gitleaks", lambda: None)
    with monkeypatch.context() as mp:
        mp.setattr(lints, "_strip_teams_join_pwd", lambda line: line)
        mp.setattr(lints, "secret_rules_identity", lambda: {"rules": "old", "gitleaks": "/x/gitleaks:1:1"})
        run(config)
    with Manifest(config.state_paths.db) as m:  # the online-only stub a gitleaks run left waiting
        m.set_meta("secret_rules", '{"gitleaks": "/x/gitleaks:1:1", "rules": "new"}')
        m.set_meta("secret_rescreen_waiting", f'[["src", "{item_id(config, "invite.txt")}"]]')
    monkeypatch.setattr(lints, "secret_rules_identity", lambda: {"rules": "new", "gitleaks": ""})
    assert run(config).commit_sha is None  # launchd's PATH: the waiting stub is not screened by a weaker scan
    assert _state(config, "invite.txt") == (RowState.QUARANTINED, "contains a credential")
    with Manifest(config.state_paths.db) as m:
        assert item_id(config, "invite.txt") in (m.get_meta("secret_rescreen_waiting") or "")


# ---------------------------------------------------------------------------------------------------------
# the hold outlives a run that could not read the whole file (review of N5's first build)
# ---------------------------------------------------------------------------------------------------------


def _ocr(mp: pytest.MonkeyPatch, engine: Any) -> None:
    """Every cycle from now on finds ``engine`` (None: a Mac without the OCR helper)."""
    mp.setattr(cycle.ocr, "engine", lambda _convert, _cache_dir: engine)


def test_a_picture_of_a_password_is_a_stub_again_after_a_run_without_ocr(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The re-screen run cannot read the pictures (no OCR helper): each gets the ``no converter`` stub, and
    its row no longer says it was a credential.  The run that reads them screens them all the same, by the
    builtin patterns too.  Before, gitleaks alone scanned the password and it was committed."""
    config, src = make_env(tmp_path)
    engine = fake_engine(tmp_path / "ocr-bin")
    _ocr(monkeypatch, engine)
    picture(src / "vpn.png", "Contoso VPN", "Password: Welcome123")
    picture(src / "invite.png", *INVITE.splitlines())
    with monkeypatch.context() as mp:
        _old_rules(mp)
        assert run(config).exit_code == 0
    stub = (RowState.QUARANTINED, "contains a credential")
    assert _state(config, "vpn.png") == stub and _state(config, "invite.png") == stub
    _blind_gitleaks(monkeypatch, tmp_path)  # gitleaks arrives: the rules moved
    _ocr(monkeypatch, None)
    assert run(config).exit_code == 0
    assert _state(config, "vpn.png") == (RowState.REFUSED, "no converter for .png")
    _ocr(monkeypatch, engine)
    screened = _count_screens(monkeypatch)
    read = run(config)
    assert read.exit_code == 0 and sorted(screened) == [1, 1]
    assert _state(config, "vpn.png") == stub
    assert not committed_blobs_containing(config.docs_repo, b"Welcome123")
    assert _state(config, "invite.png") == (RowState.LIVE, None)  # the false stub is published all the same
    assert run(config).commit_sha is None


def test_a_scan_with_a_password_is_a_stub_again_once_its_picture_page_is_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A PDF converts without OCR, and its cover page alone is clean: that page is published.  The run
    that reads the scanned page screens the file as the credential stub it was."""
    config, src = make_env(tmp_path)
    engine = shade_engine(tmp_path / "ocr-bin", {90: ["Password: Welcome123"]})
    _ocr(monkeypatch, engine)
    build_picture_pdf(src / "scan.pdf", [(["Contoso cover page"], []), ([], [page_picture(90)])])
    assert run(config).exit_code == 0
    stub = (RowState.QUARANTINED, "contains a credential")
    assert _state(config, "scan.pdf") == stub
    _blind_gitleaks(monkeypatch, tmp_path)
    _ocr(monkeypatch, None)
    assert run(config).exit_code == 0
    assert _state(config, "scan.pdf") == (RowState.LIVE, None)
    assert "Contoso cover page" in (config.docs_repo / "mirror/src/scan.pdf.md").read_text("utf-8")
    _ocr(monkeypatch, engine)
    assert run(config).exit_code == 0
    assert _state(config, "scan.pdf") == stub
    assert not committed_blobs_containing(config.docs_repo, b"Welcome123")
    assert run(config).commit_sha is None


def test_a_stub_whose_conversion_failed_is_screened_as_a_stub_by_the_run_that_converts_it(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, src = make_env(tmp_path)
    (src / "vpn.txt").write_text(PASSWORD_PAGE, encoding="utf-8")
    assert run(config).exit_code == 0
    stub = (RowState.QUARANTINED, "contains a credential")
    assert _state(config, "vpn.txt") == stub
    _blind_gitleaks(monkeypatch, tmp_path)
    real = cycle.convert_file
    with monkeypatch.context() as mp:

        def failing(path: Path, **kw: Any) -> Any:
            return dataclasses.replace(
                real(path, **kw), status=ConversionStatus.FAILED, units=(), reason="probe"
            )

        mp.setattr(cycle, "convert_file", failing)
        run(config)
    assert _state(config, "vpn.txt")[1] != "contains a credential"  # the row forgot; the record did not
    (src / "vpn.txt").write_text(PASSWORD_PAGE + "Contoso gateway\n", encoding="utf-8")
    screened = _count_screens(monkeypatch)
    assert run(config).exit_code == 0
    assert screened == [1] and _state(config, "vpn.txt") == stub
    assert not committed_blobs_containing(config.docs_repo, b"Welcome123")


def test_a_stub_saved_under_a_new_file_id_is_still_screened_as_a_stub(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A safe-save (a new file at the same path) rekeys the row; the once-a-credential record moves with
    it, so the run that reads the file after a failed conversion still screens it with the builtin rules."""
    config, src = make_env(tmp_path)
    (src / "vpn.txt").write_text(PASSWORD_PAGE, encoding="utf-8")
    assert run(config).exit_code == 0
    stub = (RowState.QUARANTINED, "contains a credential")
    assert _state(config, "vpn.txt") == stub
    old_id = item_id(config, "vpn.txt")
    _blind_gitleaks(monkeypatch, tmp_path)
    real = cycle.convert_file
    with monkeypatch.context() as mp:

        def failing(path: Path, **kw: Any) -> Any:
            return dataclasses.replace(
                real(path, **kw), status=ConversionStatus.FAILED, units=(), reason="probe"
            )

        mp.setattr(cycle, "convert_file", failing)
        run(config)
    assert _state(config, "vpn.txt")[1] != "contains a credential"
    saved = src / "vpn.txt.tmp"
    saved.write_text(PASSWORD_PAGE + "Contoso gateway\n", encoding="utf-8")
    saved.replace(src / "vpn.txt")  # the same path, another file id
    screened = _count_screens(monkeypatch)
    assert run(config).exit_code == 0
    assert item_id(config, "vpn.txt") != old_id
    assert screened == [1] and _state(config, "vpn.txt") == stub
    assert not committed_blobs_containing(config.docs_repo, b"Welcome123")


def test_a_stub_named_after_a_credential_is_redacted_when_the_rules_change(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, src = make_env(tmp_path)
    (src / "password=Spring2026.bin").write_bytes(b"\x00\x01 Contoso")
    with monkeypatch.context() as mp:
        mp.setattr(
            lints, "SECRET_PATTERNS", tuple(r for r in lints.SECRET_PATTERNS if r[0] != "generic-password")
        )
        assert run(config).exit_code == 0
    stable = item_id(config, "password=Spring2026.bin")
    with Manifest(config.state_paths.db) as m:
        assert not m.is_redacted("src", stable)
    again = run(config)
    assert again.exit_code == 0 and again.commit_sha is not None
    with Manifest(config.state_paths.db) as m:
        assert m.is_redacted("src", stable)
    assert git(config.docs_repo, "grep", "-il", "spring2026", "HEAD", "--", "mirror", check=False) == ""
    assert any("named after a credential" in a for s in again.sources for a in s.alarms)  # history: purge


def test_a_stub_whose_name_carries_a_key_is_announced_by_its_alarm_hashed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C328: an existing stub whose key is only in its name is redacted and its first-scan findings dropped,
    and its alarm says so, by the hashed page and the rule, never by the name that carries the key."""
    monkeypatch.setattr(lints, "_gitleaks", lambda: None)
    config, src = make_env(tmp_path)
    (src / "password=Spring2026.bin").write_bytes(b"\x00\x01 Contoso")
    with monkeypatch.context() as mp:
        mp.setattr(
            lints, "SECRET_PATTERNS", tuple(r for r in lints.SECRET_PATTERNS if r[0] != "generic-password")
        )
        assert run(config).exit_code == 0
    again = run(config)
    assert again.exit_code == 0
    assert not [f for f in again.lint_findings if f.code == "SECRET"], again.lint_findings
    alarms = [a for s in again.sources for a in s.alarms]
    said = [a for a in alarms if "is a `contains a credential` stub now" in a]
    assert len(said) == 1, alarms
    assert said[0].startswith("mirror/src/redacted-"), said
    assert said[0].endswith(".md is a `contains a credential` stub now: its name held one (generic-password)")
    assert "spring2026" not in " ".join(alarms).lower()


def test_a_rescan_hit_in_a_source_the_run_did_not_select_writes_it_no_heartbeat(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    two = tmp_path / "two"
    two.mkdir()
    (two / "README.txt").write_text("sentinel\n", encoding="utf-8")
    (two / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    second = (
        f'\n[[source]]\nid = "two"\nkind = "local"\npath = "{two}"\nsentinel = "README.txt"\n'
        'max_materialise_bytes = "50MB"\nmax_files = 100\n'
    )
    config, _src = make_env(tmp_path, source_extra=second)
    with monkeypatch.context() as mp:
        mp.setattr(
            lints, "SECRET_PATTERNS", tuple(r for r in lints.SECRET_PATTERNS if r[0] != "aws-access-key")
        )
        assert run(config).exit_code == 0
    before = dict(read_heartbeat(config.state_paths.heartbeat)["two"])
    assert before["ok"] is True
    again = run(config, only=["src"])
    assert again.exit_code == 0
    assert page_fm(config, "mirror/two/keys.txt.md")["reason"] == "contains a credential"
    hit = [a for s in again.sources for a in s.alarms if "matches the secret scan's new rules" in a]
    assert hit and all(a.endswith(cycle._PURGE_QUEUE_STEPS) for a in hit)
    assert dict(read_heartbeat(config.state_paths.heartbeat)["two"]) == before  # never attempted: no beat


def test_a_source_the_run_did_not_select_reads_not_in_this_run_with_no_completeness_word(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """§16.45 R51: the rescan's alarm about a page of a source the run did not list gives that source a line;
    the line says why it was skipped, not ``skipped (-) · INCOMPLETE · cursor held``, and keeps the alarm."""
    two = tmp_path / "two"
    two.mkdir()
    (two / "README.txt").write_text("sentinel\n", encoding="utf-8")
    (two / "keys.txt").write_text(f"deploy key {AWS_KEY}\n", encoding="utf-8")
    second = (
        f'\n[[source]]\nid = "two"\nkind = "local"\npath = "{two}"\nsentinel = "README.txt"\n'
        'max_materialise_bytes = "50MB"\nmax_files = 100\n'
    )
    config, _src = make_env(tmp_path, source_extra=second)
    with monkeypatch.context() as mp:
        mp.setattr(
            lints, "SECRET_PATTERNS", tuple(r for r in lints.SECRET_PATTERNS if r[0] != "aws-access-key")
        )
        assert run(config).exit_code == 0
    again = run(config, only=["src"])
    assert again.exit_code == 0
    (rep,) = [s for s in again.sources if s.source_id == "two"]
    assert (rep.pass_kind, rep.skipped_reason) == (None, "not in this run")
    capsys.readouterr()
    cli._print_report(again)
    out = capsys.readouterr().out
    (line,) = [x for x in out.splitlines() if x.startswith("  two: ")]
    assert line.startswith("  two: skipped (not in this run)"), line
    assert "INCOMPLETE" not in line and "complete" not in line and "cursor" not in line, line
    assert "    alarm: a published page matches the secret scan's new rules" in out
    (month,) = (config.docs_repo / "CHANGELOG").glob("*.md")
    (entry,) = [x for x in month.read_text(encoding="utf-8").splitlines() if x.startswith("- two: ")][-1:]
    assert "skipped (not in this run)" in entry and "complete" not in entry.lower(), entry


def test_gitleaks_reads_the_pre_write_screen_inside_its_private_folder(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A root that is no git repository (the folder ``_still_credential`` makes under staging): the copy
    gitleaks reads is made in it, not in the system's temporary folder."""
    root = tmp_path / "private"
    root.mkdir()
    (root / "vpn.md").write_text(PASSWORD_PAGE, encoding="utf-8")
    exe = tmp_path / "gitleaks"
    exe.write_bytes(b"#!/bin/sh\n")
    scanned: list[Path] = []

    def spy(argv: Sequence[str], **_kw: Any) -> subprocess.CompletedProcess[str]:
        scanned.append(Path(argv[2]))
        assert (Path(argv[2]) / "vpn.md").read_text(encoding="utf-8") == PASSWORD_PAGE
        Path(argv[argv.index("--report-path") + 1]).write_text(json.dumps([]), encoding="utf-8")
        return subprocess.CompletedProcess(list(argv), 0, "", "")

    monkeypatch.setattr(lints, "_gitleaks", lambda: exe)
    monkeypatch.setattr(lints.subprocess, "run", spy)
    hits = lints.lint_secrets(root, ["vpn.md"])
    assert [f.path for f in hits] == ["vpn.md"]
    assert [d.parent for d in scanned] == [root.resolve()]
    assert sorted(f.name for f in root.iterdir()) == ["vpn.md"]  # and it is gone afterwards


def test_an_unreadable_once_a_credential_record_is_rebuilt_from_the_manifest_and_logged(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    """B13 R77: a ``secret_once_credential`` value that does not parse is not read as empty: the log says so,
    the record is rebuilt from the rows quarantined as a credential, and the rebuilt value is written back."""
    config, src = make_env(tmp_path)
    (src / "vpn.txt").write_text(PASSWORD_PAGE, encoding="utf-8")
    (src / "notes.txt").write_text("Contoso standup notes\n", encoding="utf-8")
    assert run(config).exit_code == 0
    assert _state(config, "vpn.txt") == (RowState.QUARANTINED, "contains a credential")
    pair = ("src", item_id(config, "vpn.txt"))
    # a file whose credential queued a purge and whose stub has another reason now: found through the queue
    queued_pair = ("src", item_id(config, "notes.txt"))
    assert governance.enqueue_purge(
        config.state_paths.root,
        governance.PurgeSelector(source_id=queued_pair[0], stable_id=queued_pair[1]),
        governance.PurgeReason.DLP_REMEDIATION,
    )
    with Manifest(config.state_paths.db) as m:
        m.set_meta(cycle._SECRET_ONCE_META, "not json")
        reader = SimpleNamespace(_once=None, manifest=m, config=config)
        reader._rebuilt_once = lambda exc: cycle._Cycle._rebuilt_once(reader, exc)  # type: ignore[arg-type]
        with caplog.at_level("WARNING", logger="agentsync.cycle"):
            once = cycle._Cycle._once_credential(reader)  # type: ignore[arg-type]
        stored = json.loads(m.get_meta(cycle._SECRET_ONCE_META) or "null")
    assert pair in once and queued_pair in once
    assert sorted(tuple(p) for p in stored) == sorted([pair, queued_pair])
    assert any("secret_once_credential is unreadable" in r.getMessage() for r in caplog.records)
