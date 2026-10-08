"""The generated fixtures are real files their libraries can read (guards the fixture builder itself)."""

from __future__ import annotations

import email
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pdfminer.high_level import extract_text
from pptx import Presentation


def test_every_fixture_exists(fixture_files: dict[str, Path]) -> None:
    assert set(fixture_files) == {
        "2026-09.teams.json",
        "sample-resaved.xlsx",
        "sample.csv",
        "sample.docx",
        "sample.eml",
        "sample.html",
        "sample.md",
        "sample.pdf",
        "sample.pptx",
        "sample.txt",
        "sample.vtt",
        "sample.xlsx",
    }
    assert all(p.stat().st_size > 0 for p in fixture_files.values())


def test_office_fixtures_open(fixture_files: dict[str, Path]) -> None:
    assert zipfile.is_zipfile(fixture_files["sample.docx"])
    wb = load_workbook(fixture_files["sample.xlsx"])
    assert wb.sheetnames == ["Q3 Budget", "Notes"]
    assert wb["Q3 Budget"]["D2"].value == "=C2-B2"
    assert "A6:D6" in {str(r) for r in wb["Q3 Budget"].merged_cells.ranges}
    prs = Presentation(str(fixture_files["sample.pptx"]))
    assert len(prs.slides) == 2


def test_resaved_xlsx_differs_only_outside_content(fixture_files: dict[str, Path]) -> None:
    a = fixture_files["sample.xlsx"].read_bytes()
    b = fixture_files["sample-resaved.xlsx"].read_bytes()
    assert a != b
    with (
        zipfile.ZipFile(fixture_files["sample.xlsx"]) as za,
        zipfile.ZipFile(fixture_files["sample-resaved.xlsx"]) as zb,
    ):
        differing = sorted(n for n in za.namelist() if za.read(n) != zb.read(n))
    assert differing == ["docProps/core.xml"]


def test_pdf_eml_teams_fixtures(fixture_files: dict[str, Path]) -> None:
    text = extract_text(str(fixture_files["sample.pdf"]))
    assert "purchase order" in text and "page 2" in text
    msg = email.message_from_bytes(fixture_files["sample.eml"].read_bytes())
    assert msg["Subject"] == "Acme kickoff recap"
    doc = json.loads(fixture_files["2026-09.teams.json"].read_text())
    assert doc["schema"] == "agentsync.teams-month/1"


def test_local_source_tree(local_source_dir: Path) -> None:
    assert (local_source_dir / "README.txt").is_file()
    assert (local_source_dir / "projects" / "acme" / "Kickoff Notes.docx").is_file()
