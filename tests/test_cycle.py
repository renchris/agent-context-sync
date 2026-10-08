"""cycle.py mechanics: recovery after crashes, cursor ordering, source isolation, the lock, dry-run,
retirement.  Crashes are simulated with a BaseException the cycle does not catch (like a SIGKILL mid-cycle:
no cleanup code runs), so the next cycle's ``recover`` sees exactly what a dead process leaves behind."""

from __future__ import annotations

import contextlib
import dataclasses
import errno
import hashlib
import io
import itertools
import json
import logging
import os
import re
import shutil
import struct
import subprocess
import sys
import threading
import time
import zlib
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from unittest.mock import ANY

import httpx
import pytest
from pptx import Presentation
from pptx.util import Inches

from agentsync import arm_local as al
from agentsync import curate, gitops, governance, loop, materialise, slug
from agentsync import cycle as cycle_mod
from agentsync.arm_local import LocalArm
from agentsync.config import Config, parse_config
from agentsync.convert import image as image_mod
from agentsync.convert import ocr
from agentsync.convert import pandoc as pandoc_mod
from agentsync.convert import pdf as pdf_mod
from agentsync.convert import recording as recording_mod
from agentsync.convert.base import estimate_tokens, rendered_sha256
from agentsync.convert.media import MediaEngine, MediaError
from agentsync.convert.pieces import PieceStore
from agentsync.convert.recording import Allowance, RecordingConverter, RecordingNotFinished
from agentsync.convert.registry import Registry
from agentsync.cycle import RecoveryAction, recover, run_cycle
from agentsync.errors import ConversionError, DatalessRefusedError, LockHeldError
from agentsync.graph.client import GraphClient
from agentsync.graph.drive import DriveArm
from agentsync.manifest import Manifest
from agentsync.model import (
    CycleMode,
    CycleReport,
    PassKind,
    RenderedUnit,
    RowState,
    ScanResult,
    UnitKind,
    Verdict,
)
from agentsync.ops import doctor
from agentsync.ops.lock import SingleWriterLock, read_heartbeat
from agentsync.paths import DocsLayout
from agentsync.publish import Publisher, sidecar_rel
from conftest import config_text
from test_convert_builders import (
    build_annotated_pdf,
    build_picture_pdf,
    page_picture,
    pandoc_build,
    shade_engine,
)
from test_convert_image import picture, picture_bytes, reads, text_png
from test_e2e import (
    GRAPH_SOURCE,
    SID,
    FakeDrive,
    FakeTokens,
    clock,
    config_with,
    git,
    page,
    porcelain,
    source_report,
)
from test_ocr import calls, fake_engine, write_fake
from test_review_fixes import committed_blobs_containing


class Crash(BaseException):
    """Raised to kill a cycle mid-flight without running its except handlers."""


def run(config: Config, **kwargs: Any) -> CycleReport:
    return run_cycle(config, mode=kwargs.pop("mode", CycleMode.POLL), now=clock, **kwargs)


def crash(*_a: object, **_k: object) -> None:
    raise Crash


@pytest.fixture(autouse=True)
def _no_media_helper(monkeypatch: pytest.MonkeyPatch) -> None:
    """No cycle finds a media helper unless its test gives it one, as on a Mac whose installer built none."""
    monkeypatch.setattr(cycle_mod.media, "engine", lambda _convert, _cache_dir: None)


@pytest.fixture
def drive_env(tmp_path: Path, local_source_dir: Path, fixture_files: dict[str, Path]) -> Iterator[Any]:
    config = config_with(tmp_path, local_source_dir, GRAPH_SOURCE)
    drive = FakeDrive({"I1": ("notes.md", b"# Notes\n\nfirst version\n")})
    client = GraphClient(
        FakeTokens(),
        user_agent="NONISV|test|agentsync/0",
        transport=httpx.MockTransport(drive.handler),
        sleep=lambda _s: None,
    )
    yield config, drive, client
    client.close()


def edit_i1(drive: FakeDrive, text: bytes) -> None:
    drive.version["I1"] += 1
    drive.qx["I1"] = f"qx-I1-{drive.version['I1']}"
    drive.files["I1"] = ("notes.md", text)


def cursor(config: Config, sid: str = "drive") -> tuple[str | None, str | None]:
    with Manifest(config.state_paths.db) as m:
        c = m.get_cursor(sid)
        return (c.current, c.pending) if c is not None else (None, None)


# ---------------------------------------------------------------------------------------------------------
# cursor ordering and crash recovery
# ---------------------------------------------------------------------------------------------------------


def test_crash_after_commit_before_promotion_adopts_the_pending_cursor(
    drive_env: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, drive, client = drive_env
    assert run(config, client=client, only=["drive"]).exit_code == 0
    assert cursor(config)[0] is not None and cursor(config)[0].endswith("token=T1")  # type: ignore[union-attr]
    edit_i1(drive, b"# Notes\n\nsecond version\n")
    monkeypatch.setattr(Manifest, "promote_cursors", crash)
    with pytest.raises(Crash):
        run(config, client=client, only=["drive"])
    monkeypatch.undo()
    head_msg = git(config.docs_repo, "log", "-1", "--format=%B")
    assert "Agentsync-Run: 2" in head_msg  # the commit landed ...
    current, pending = cursor(config)
    assert current is not None and current.endswith("token=T1") and pending is not None  # ... unpromoted
    with Manifest(config.state_paths.db) as m:
        action = recover(config, m)
        assert action is RecoveryAction.ADOPT_PENDING
        c = m.get_cursor("drive")
        assert (
            c is not None and c.current is not None and c.current.endswith("token=T2") and c.pending is None
        )
        assert m.last_runs(1)[0][2] == "aborted"
    after = run(config, client=client, only=["drive"])
    assert after.exit_code == 0 and after.commit_sha is None  # nothing re-applied, nothing lost
    assert "second version" in page(config.docs_repo, "mirror/drive/projects/notes.md")[1]


def test_crash_before_commit_discards_the_cursor_and_the_next_cycle_commits(
    drive_env: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    config, drive, client = drive_env
    run(config, client=client, only=["drive"])
    head = git(config.docs_repo, "rev-parse", "HEAD")
    edit_i1(drive, b"# Notes\n\nthird version\n")
    monkeypatch.setattr(gitops, "commit_cycle", crash)
    with pytest.raises(Crash):
        run(config, client=client, only=["drive"])
    monkeypatch.undo()
    assert git(config.docs_repo, "rev-parse", "HEAD") == head
    assert porcelain(config.docs_repo) != ""  # published, never committed
    current, pending = cursor(config)
    assert current is not None and current.endswith("token=T1") and pending is not None
    with Manifest(config.state_paths.db) as m:
        assert recover(config, m) is RecoveryAction.DISCARD_PENDING
        c = m.get_cursor("drive")
        assert (
            c is not None and c.pending is None and c.current is not None and c.current.endswith("token=T1")
        )
    report = run(config, client=client, only=["drive"])
    assert report.exit_code == 0 and report.commit_sha is not None
    assert "third version" in page(config.docs_repo, "mirror/drive/projects/notes.md")[1]
    assert porcelain(config.docs_repo) == ""
    current, _ = cursor(config)
    assert current is not None and current.endswith("token=T3")  # the replayed round's cursor


def test_blocking_lint_commits_nothing_and_promotes_no_cursor(drive_env: Any) -> None:
    config, _drive, client = drive_env
    repo = config.docs_repo
    repo.mkdir(parents=True, exist_ok=True)
    (repo / "topics").mkdir(parents=True, exist_ok=True)
    (repo / "topics" / "escape").symlink_to("/etc")
    report = run(config, client=client)
    assert report.exit_code == 1 and report.commit_sha is None
    assert any(f.code == "SYMLINK" and f.blocking for f in report.lint_findings)
    assert all(not s.cursor_advanced for s in report.sources)
    assert cursor(config) == (None, None)
    assert gitops.head_sha(repo) is None
    assert "BLOCKING SYMLINK" in (repo / "_sync" / "STATE.md").read_text(encoding="utf-8")
    (repo / "topics" / "escape").unlink()
    healed = run(config, client=client)
    assert healed.exit_code == 0 and healed.commit_sha is not None
    current, _ = cursor(config)
    assert current is not None


def test_head_moved_by_hand_resets_generated_paths_and_republishes(sample_config: Config) -> None:
    repo = sample_config.docs_repo
    run(sample_config)
    victim = f"mirror/{SID}/projects/sample.docx.md"
    git(repo, "rm", "-q", victim)
    git(repo, "commit", "-q", "-m", "manual: someone deleted a generated page")
    assert not (repo / victim).exists()
    with Manifest(sample_config.state_paths.db) as m:
        assert recover(sample_config, m) is RecoveryAction.RESET_GENERATED
    assert (repo / victim).is_file()  # re-published from the converter cache
    report = run(sample_config)
    assert report.exit_code == 0 and report.commit_sha is not None
    assert victim in git(repo, "ls-files")
    assert porcelain(repo) == ""


def test_archive_on_upstream_delete_survives_a_head_moved_by_hand(
    sample_config: Config, local_source_dir: Path
) -> None:
    """[governance] archive: the cycle copies a deleted page to archive/, queues no purge, skips compaction,
    and the repair path re-renders the mirror tombstone still naming the archive copy."""
    cfg = sample_config.config_path
    cfg.write_text(cfg.read_text(encoding="utf-8") + "\n[governance]\narchive = true\n", encoding="utf-8")
    repo = sample_config.docs_repo
    run(sample_config)
    victim, kept = f"mirror/{SID}/projects/sample.txt.md", f"archive/{SID}/projects/sample.txt.md"
    (local_source_dir / "projects" / "sample.txt").unlink()
    run(sample_config)
    report = run(sample_config, mode=CycleMode.RECONCILE)
    assert report.exit_code == 0 and kept in git(repo, "ls-files")
    assert "\nstatus: archived\n" in (repo / kept).read_text(encoding="utf-8")
    assert "- archive on: history is kept" in (repo / "_sync" / "STATE.md").read_text(encoding="utf-8")
    assert governance.pending_purges(sample_config.state_paths.root) == []
    git(repo, "rm", "-q", victim)
    git(repo, "commit", "-q", "-m", "manual: someone deleted a tombstone")
    with Manifest(sample_config.state_paths.db) as m:
        assert recover(sample_config, m) is RecoveryAction.RESET_GENERATED
    assert f"kept, searchable, at {kept}" in (repo / victim).read_text(encoding="utf-8")
    assert run(sample_config).exit_code == 0 and porcelain(repo) == ""


def test_crash_between_observation_and_rename_publish_is_repaired(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo = sample_config.docs_repo
    run(sample_config)
    (local_source_dir / "projects" / "acme").rename(local_source_dir / "projects" / "Acme Corp")
    monkeypatch.setattr(Publisher, "rewrite_frontmatter", crash)
    with pytest.raises(Crash):
        run(sample_config)
    monkeypatch.undo()
    # The observation transaction committed the new rel_path; the page still sits at the old one.
    assert (repo / f"mirror/{SID}/projects/acme/kickoff-notes.docx.md").is_file()
    report = run(sample_config)
    assert report.exit_code == 0 and report.commit_sha is not None
    assert (repo / f"mirror/{SID}/projects/acme-corp/kickoff-notes.docx.md").is_file()
    assert not (repo / f"mirror/{SID}/projects/acme").exists()
    assert porcelain(repo) == ""


def test_cycle_level_failure_is_reported_not_raised(
    sample_config: Config, monkeypatch: pytest.MonkeyPatch
) -> None:
    def boom(*_a: object, **_k: object) -> None:
        raise OSError("disk full")

    monkeypatch.setattr(Publisher, "write_index", boom)
    report = run(sample_config)
    assert report.exit_code == 1 and report.commit_sha is None
    assert any(f.code == "CYCLE" and "disk full" in f.message for f in report.lint_findings)
    with Manifest(sample_config.state_paths.db) as m:
        assert m.last_runs(1)[0][2] == "failed"
    monkeypatch.undo()
    assert run(sample_config).commit_sha is not None


# ---------------------------------------------------------------------------------------------------------
# isolation, lock, dry run, paused/retired
# ---------------------------------------------------------------------------------------------------------


def two_source_config(tmp_path: Path, a: Path, b: Path) -> Config:
    text = config_text(tmp_path / "docs", tmp_path / "state", tmp_path / "cache", a)
    text += f'\n[[source]]\nid = "second"\nkind = "local"\npath = "{b}"\n'
    path = tmp_path / "sources.toml"
    path.write_text(text, encoding="utf-8")
    return parse_config(text, config_path=path)


def test_a_failing_source_does_not_stop_the_others(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    other = tmp_path / "other"
    other.mkdir()
    (other / "note.txt").write_text("second source\n", encoding="utf-8")
    config = two_source_config(tmp_path, local_source_dir, other)
    real_scan = LocalArm.scan

    def scan(self: LocalArm, cursor: str | None, *, full: bool) -> ScanResult:
        if self.source_id == "second":
            raise RuntimeError("provider exploded")
        return real_scan(self, cursor, full=full)

    monkeypatch.setattr(LocalArm, "scan", scan)
    report = run(config)
    assert report.exit_code == 1
    good = next(s for s in report.sources if s.source_id == SID)
    bad = next(s for s in report.sources if s.source_id == "second")
    assert good.errors == () and bad.errors and "provider exploded" in bad.errors[0]
    assert report.commit_sha is not None  # the healthy source still published
    assert (config.docs_repo / f"mirror/{SID}/projects/sample.docx.md").is_file()
    assert not (config.docs_repo / "mirror/second").exists()
    with Manifest(config.state_paths.db) as m:
        assert m.last_runs(1)[0][2] == "partial"
    hb = read_heartbeat(config.state_paths.heartbeat)
    assert hb[SID]["ok"] is True and hb["second"]["ok"] is False
    monkeypatch.undo()
    healed = run(config)
    assert healed.exit_code == 0 and (config.docs_repo / "mirror/second/note.txt.md").is_file()


def test_a_failing_item_is_a_row_not_a_crash(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    real_fetch = LocalArm.fetch

    def fetch(self: LocalArm, item: Any, dest_dir: Path, budget: Any) -> Any:
        if item.name == "sample.pdf":
            raise OSError(5, "Input/output error")
        return real_fetch(self, item, dest_dir, budget)

    monkeypatch.setattr(LocalArm, "fetch", fetch)
    report = run(sample_config)
    assert report.exit_code == 0 and report.commit_sha is not None
    s = report.sources[0]
    assert s.counts.get(Verdict.ERROR) == 1 and any("sample.pdf" in e for e in s.errors)
    with Manifest(sample_config.state_paths.db) as m:
        row = m.item_by_path(SID, "projects/sample.pdf")
        assert row is not None and row.last_verdict is Verdict.ERROR
    monkeypatch.undo()
    retry = run(sample_config)
    assert [c.path for c in retry.changes] == [f"mirror/{SID}/projects/sample.pdf.md"]


def test_lock_held_raises_lock_held_error(sample_config: Config) -> None:
    lock = SingleWriterLock(sample_config.state_paths.lock, "other")
    lock.acquire()
    try:
        with pytest.raises(LockHeldError):
            run(sample_config)
    finally:
        lock.release()
    assert run(sample_config).exit_code == 0


def test_dry_run_classifies_without_writing(sample_config: Config) -> None:
    repo = sample_config.docs_repo
    report = run(sample_config, mode=CycleMode.DRY_RUN)
    assert report.exit_code == 0 and report.commit_sha is None and report.changes == ()
    assert report.sources[0].counts.get(Verdict.CREATED) == 13
    assert gitops.head_sha(repo) is None and not (repo / "mirror").exists()
    with Manifest(sample_config.state_paths.db) as m:
        assert list(m.iter_items(SID)) == [] and m.last_runs(1) == []
    assert run(sample_config).commit_sha is not None


def test_a_cycle_stores_the_empty_cloud_folders_its_walk_found(
    tmp_path: Path, tmp_docs_repo: Path, tmp_state_dir: Path
) -> None:
    """``loop.next_step`` names the folders that keep a listing incomplete from manifest meta, so status and
    STATE.md never walk the source a second time. Each real pass rewrites the list (empty once the walk
    finds none, or cannot walk); a dry run writes nothing."""
    root = Path.home() / "Library" / "CloudStorage" / "OneDrive-Contoso" / "Work"
    (root / "Bids" / "Old").mkdir(parents=True)
    (root / "README.txt").write_text("sentinel\n", encoding="utf-8")
    text = config_text(tmp_docs_repo, tmp_state_dir, tmp_path / "cache", root)
    config = parse_config(text, config_path=tmp_path / "sources.toml")
    key = cycle_mod._EMPTY_DIRS_META + SID

    def stored() -> str | None:
        with Manifest(config.state_paths.db) as m:
            return m.get_meta(key)

    report = run(config)
    assert not report.sources[0].enumeration_complete
    assert stored() == '["Bids/Old"]'
    (root / "Plans — Draft").mkdir()
    assert run(config, mode=CycleMode.DRY_RUN).exit_code == 0
    assert stored() == '["Bids/Old"]', "a dry run stores nothing"
    run(config)
    assert stored() == '["Bids/Old", "Plans — Draft"]', "sorted, names as they are"
    (root / "Bids" / "Old" / "a.txt").write_text("x\n", encoding="utf-8")
    (root / "Plans — Draft").rmdir()
    assert run(config).sources[0].enumeration_complete
    assert stored() == ""
    (root / "Bids" / "Empty").mkdir()
    run(config)
    assert stored() == '["Bids/Empty"]'
    root.rename(root.with_name("Work moved"))
    assert not run(config).sources[0].enumeration_complete
    assert stored() == "", "a walk that could not run knows no empty folder"


SYNC_FIX = "agentsync sync (it makes these owner-only)"
"""``docs_repo.permissions``'s fix when a sync clears every path it found: a command a setup agent may run."""


def permissions(config: Config) -> doctor.CheckResult:
    [result] = doctor._check_permissions(config)
    return result


def mode(path: Path) -> int:
    return path.lstat().st_mode & 0o777


def own_config(tmp_path: Path, repo: Path) -> Config:
    """A config for ``repo`` alone: the cache and log folders are tmp folders that need not exist, and the
    config's own folder does not hold the docs repo."""
    text = config_text(repo, tmp_path / "state", tmp_path / "cache", tmp_path / "source")
    return parse_config(text, config_path=tmp_path / "config" / "sources.toml")


def test_a_cycle_makes_its_own_paths_owner_only_and_a_dry_run_does_not(sample_config: Config) -> None:
    """Field reports 2026-10-06 and 2026-10-07: the baseline draft, written by an agent under umask 022,
    left ``_eval/`` readable by group and other, and the next status ended on a ``docs_repo.permissions``
    FAIL. A cache folder and a log that something else made do the same. The fix on the FAIL names the
    sync, and one cycle clears them all."""
    repo = sample_config.docs_repo
    run(sample_config)
    eval_dir, topic_dir = repo / "_eval", repo / "topics" / "contoso" / "notes"
    old_build, log_dir = sample_config.cache_dir / "old-build", sample_config.log_dir
    eval_dir.mkdir()
    topic_dir.mkdir(parents=True)
    old_build.mkdir()
    log_dir.mkdir(exist_ok=True)
    log = log_dir / "poll.out.log"
    files = [eval_dir / "questions.md", topic_dir / "scratch.txt", repo / "NOTES.txt", log]
    for f in files:
        f.write_text("1. What did Contoso decide?\n", encoding="utf-8")
        f.chmod(0o644)
    script = topic_dir / "run.sh"
    script.write_text("#!/bin/sh\n", encoding="utf-8")
    script.chmod(0o755)
    dirs = [eval_dir, topic_dir, topic_dir.parent, old_build, log_dir, repo, repo.parent]
    for d in dirs:
        d.chmod(0o755)
    git(repo, "config", "--local", "--unset", "core.sharedRepository")  # a repo from before the setting

    def modes(paths: list[Path]) -> set[int]:
        return {mode(p) for p in paths}

    found = permissions(sample_config)
    assert not found.ok and "_eval" in found.detail and found.fix == SYNC_FIX
    assert run(sample_config, mode=CycleMode.DRY_RUN).exit_code == 0
    assert modes(files) == {0o644} and modes(dirs) == {0o755} and not permissions(sample_config).ok

    assert run(sample_config).exit_code == 0
    assert modes(files) == {0o600} and modes(dirs) == {0o700} and modes([script]) == {0o700}
    assert permissions(sample_config).ok, permissions(sample_config).detail
    assert git(repo, "config", "--local", "core.sharedRepository").strip() == "0600"
    # The draft rode in the cycle's commit and a mode change dirties nothing; the stray file is not a
    # path the cycle commits.
    assert porcelain(repo) == "?? NOTES.txt\n"


def test_tightening_follows_no_symlink_and_leaves_the_mirror_walk_to_the_publisher(tmp_path: Path) -> None:
    repo, outside = tmp_path / "docs", tmp_path / "outside"
    for d in (repo / "_eval", repo / "topics", repo / "mirror" / "src", repo / "other" / "deep", outside):
        d.mkdir(parents=True)
    loose = [outside / "kept.txt", repo / "mirror" / "src" / "page.md", repo / "other" / "deep" / "f.txt"]
    for f in loose:
        f.write_text("x\n", encoding="utf-8")
        f.chmod(0o644)
    for d in (outside, repo / "mirror" / "src", repo / "other" / "deep", repo / "other", repo / "mirror"):
        d.chmod(0o755)
    (repo / "_eval" / "file-link").symlink_to(outside / "kept.txt")
    (repo / "topics" / "dir-link").symlink_to(outside, target_is_directory=True)
    (repo / "top-link").symlink_to(outside / "kept.txt")
    config = own_config(tmp_path, repo)
    config.cache_dir.symlink_to(outside, target_is_directory=True)  # a cache folder that is a link
    assert cycle_mod._tighten_own_paths(config) == 2  # the two top-level folders themselves, nothing else
    assert mode(repo / "mirror") == mode(repo / "other") == 0o700
    assert {mode(f) for f in loose} == {0o644}
    assert mode(outside) == 0o755 and mode(repo / "mirror" / "src") == 0o755
    # What a sync leaves is the chmod's: the fix names no sync that would not clear it.
    left = cycle_mod._sync_leaves(config, [*loose[1:], repo / "mirror" / "src", repo / "_eval" / "file-link"])
    assert left == [*loose[1:], repo / "mirror" / "src", repo / "_eval" / "file-link"]
    assert "chmod -R go-rwx" in (permissions(config).fix or "")
    assert cycle_mod._tighten_own_paths(config) == 0
    assert cycle_mod._tighten_own_paths(own_config(tmp_path, tmp_path / "missing")) == 0


def test_tightening_leaves_a_log_folder_that_is_the_home_folder(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A config may name any folder for its cache and its logs. One that is the home folder, or holds it,
    is not agentsync's alone, so nothing in it is changed: the check reports it, and the chmod is the
    person's decision. The file system says which folder a path is, not its spelling (review, 2026-10-07):
    ``~/x/..``, a path through a symlink and, where the volume takes a name in any case, another case all
    passed a comparison of the paths as written, and the home folder and what is in it were made
    owner-only."""
    repo, home = tmp_path / "docs", tmp_path / "volume" / "people" / "me"
    repo.mkdir()
    (home / "x").mkdir(parents=True)
    monkeypatch.setenv("HOME", str(home))
    mine = home / "notes.txt"
    mine.write_text("x\n", encoding="utf-8")
    loose = {home.parent: 0o755, home: 0o755, home / "x": 0o755, mine: 0o644}
    for path, bits in loose.items():
        path.chmod(bits)
    (tmp_path / "link").symlink_to(home.parent, target_is_directory=True)
    spellings = [home, home.parent, home / "x" / "..", home / "..", tmp_path / "link" / "me"]
    spellings += [p for p in (home.with_name("ME"),) if p.is_dir()]  # a volume that takes any case
    for log_dir in spellings:
        config = dataclasses.replace(own_config(tmp_path, repo), log_dir=log_dir)
        assert cycle_mod._holds_home(log_dir), log_dir
        assert cycle_mod._tighten_own_paths(config) == 0, log_dir
        assert {path: mode(path) for path in loose} == loose, log_dir
        assert cycle_mod._sync_leaves(config, [mine]) == [mine]
    assert not cycle_mod._holds_home(home / "x") and not cycle_mod._holds_home(tmp_path / "missing")
    assert "chmod -R go-rwx" in (permissions(dataclasses.replace(config, log_dir=home)).fix or "")
    monkeypatch.setenv("HOME", str(tmp_path / "link" / "me"))  # a home folder reached through a symlink
    assert cycle_mod._holds_home(tmp_path / "volume"), "the folder above where it really is holds it too"
    assert cycle_mod._holds_home(home) and not cycle_mod._holds_home(repo)


def test_tightening_never_widens_a_mode(tmp_path: Path) -> None:
    """Only group and other bits are cleared: what the owner may do with a path is what it was."""
    repo = tmp_path / "docs"
    (repo / "_eval").mkdir(parents=True)
    before = {"read-only.md": 0o444, "private.md": 0o600, "sealed.md": 0o400, "tool.sh": 0o750}
    for name, bits in before.items():
        (repo / "_eval" / name).write_text("x\n", encoding="utf-8")
        (repo / "_eval" / name).chmod(bits)
    assert cycle_mod._tighten_own_paths(own_config(tmp_path, repo)) == 2
    after = {name: mode(repo / "_eval" / name) for name in before}
    assert after == {"read-only.md": 0o400, "private.md": 0o600, "sealed.md": 0o400, "tool.sh": 0o700}


def test_tightening_looks_at_a_bounded_number_of_entries_below_each_tree(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A cache holds a few entries per converted file, so a sync does not walk all of it: it looks at a
    fixed number of entries below each tree. What the walk does not reach stays the check's to report,
    and its fix is then the chmod, not a sync that would not clear it."""
    repo = tmp_path / "docs"
    repo.mkdir()
    config = own_config(tmp_path, repo)
    config.log_dir.mkdir(parents=True)
    logs = [config.log_dir / f"job-{n}.log" for n in range(5)]
    for f in logs:
        f.write_text("x\n", encoding="utf-8")
        f.chmod(0o644)
    monkeypatch.setattr(cycle_mod, "_OWN_WALK", 3)
    assert cycle_mod._tighten_own_paths(config) == 3
    still = [f for f in logs if mode(f) == 0o644]
    assert len(still) == 2 and cycle_mod._tighten_own_paths(config) == 0
    found = permissions(config)
    assert not found.ok and all(str(f) in found.detail for f in still)
    assert "chmod -R go-rwx" in (found.fix or "")
    monkeypatch.undo()
    assert permissions(config).fix == SYNC_FIX
    assert cycle_mod._tighten_own_paths(config) == 2 and permissions(config).ok
    assert cycle_mod._OWN_WALK >= doctor._PERM_SAMPLE, "a sync reaches every entry the check samples"


def test_tightening_walks_a_tree_in_the_order_the_check_does_and_leaves_no_descriptor_open(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The walk goes from descriptor to descriptor, and the check from path to path (``os.walk``). The bound
    holds only while both list a tree in one order: the first entries the check samples are then the first
    the sync reaches. A folder that is a symlink is listed and not entered by either."""
    tree, outside = tmp_path / "cache", tmp_path / "outside"
    for d in (tree / "b" / "deep" / "deeper", tree / "a", tree / "c" / "empty", outside / "inner"):
        d.mkdir(parents=True)
    for f in (tree / "z.txt", tree / "b" / "y.txt", tree / "b" / "deep" / "x.txt", outside / "kept.txt"):
        f.write_text("x\n", encoding="utf-8")
    (tree / "a" / "dir-link").symlink_to(outside, target_is_directory=True)
    (tree / "c" / "gone-link").symlink_to(tmp_path / "missing")
    as_the_check = [Path(at) / name for at, dirs, files in os.walk(tree) for name in (*dirs, *files)]
    assert len(as_the_check) == 11 and tree / "a" / "dir-link" in as_the_check

    def open_descriptors() -> int:
        return len(list(Path("/dev/fd").iterdir()))

    before = open_descriptors()
    assert [path for path, _name, _dir_fd in cycle_mod._below(tree, tree)] == as_the_check
    monkeypatch.setattr(cycle_mod, "_OWN_WALK", 4)
    stopped = cycle_mod._below(tree, tree)
    assert [path for path, _name, _dir_fd in stopped] == as_the_check[:4]
    left = cycle_mod._below(tree, tree)  # a caller that stops asking: closing it closes what it had open
    assert next(left)[0] == as_the_check[0] and open_descriptors() > before
    left.close()
    assert open_descriptors() == before


def test_a_path_of_another_users_keeps_its_fail_and_the_chmod(
    sample_config: Config, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """No false green: only its owner may change a mode, so a path of another user's is left as it is, the
    cycle says so once, and the check still fails with the chmod as its fix (a sync would not clear it)."""
    repo = sample_config.docs_repo
    run(sample_config)
    theirs = repo / "_eval" / "questions.md"
    theirs.parent.mkdir()
    theirs.write_text("1. What did Contoso decide?\n", encoding="utf-8")
    theirs.chmod(0o644)
    assert permissions(sample_config).fix == SYNC_FIX, "this user's: a sync clears it"

    def not_the_owner(_fd: int, _mode: int) -> None:
        raise PermissionError(errno.EPERM, "Operation not permitted")

    monkeypatch.setattr(os, "fchmod", not_the_owner)  # what the kernel answers for a file of another user
    monkeypatch.setattr(os, "geteuid", lambda: theirs.stat().st_uid + 1)
    with caplog.at_level(logging.WARNING, logger="agentsync.cycle"):
        assert cycle_mod._tighten_own_paths(sample_config) == 0
    assert "1 path(s) could not be made owner-only" in caplog.text
    found = permissions(sample_config)
    assert mode(theirs) == 0o644 and not found.ok and str(theirs) in found.detail
    assert (found.fix or "").startswith("chmod -R go-rwx ") and "agentsync sync" not in (found.fix or "")


@pytest.mark.skipif(os.geteuid() == 0, reason="root opens a file whatever its mode")
def test_a_path_its_owner_may_not_read_keeps_its_fail_and_the_chmod(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    """Review, 2026-10-07: a mode is changed through a descriptor, and opening one takes the owner's read
    bit. A file at 0044 and a folder at 0333 are this user's and still not paths a sync changes, so a fix
    that named the sync sent a setup run round in a circle: sync, status, the same FAIL. Their fix is the
    chmod, which needs no read bit."""
    repo = tmp_path / "docs"
    unread, sealed = repo / "_eval" / "unread.md", repo / "_eval" / "sealed"
    sealed.mkdir(parents=True)
    unread.write_text("1. What did Contoso decide?\n", encoding="utf-8")
    config = own_config(tmp_path, repo)
    unread.chmod(0o644)
    assert permissions(config).fix == SYNC_FIX, "readable by its owner: a sync clears it"
    was = {unread: 0o044, sealed: 0o333}
    for path, bits in was.items():
        path.chmod(bits)
    try:
        with caplog.at_level(logging.WARNING, logger="agentsync.cycle"):
            assert cycle_mod._tighten_own_paths(config) == 0
        assert "2 path(s) could not be made owner-only" in caplog.text
        found = permissions(config)
        assert {path: mode(path) for path in was} == was and not found.ok
        assert all(str(path) in found.detail for path in was)
        assert (found.fix or "").startswith("chmod -R go-rwx ") and "agentsync sync" not in (found.fix or "")
        for path, bits in was.items():  # what that chmod does to them
            path.chmod(bits & ~0o077)
        assert permissions(config).ok
    finally:
        sealed.chmod(0o700)  # the tmp folder can be removed again


def test_tightening_never_follows_an_entry_swapped_for_a_symlink(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The entries it tightens are the ones others can still write, so one can be replaced between the look
    and the change. The mode is changed through a descriptor opened without following a link: a symlink put
    there in between is an error for that path, and its target outside the docs repo keeps its mode."""
    repo, outside = tmp_path / "docs", tmp_path / "outside"
    (repo / "_eval").mkdir(parents=True)
    outside.mkdir()
    outside.chmod(0o755)
    before = tmp_path / "before.md"
    before.write_text("1. What did Contoso decide?\n", encoding="utf-8")
    before.chmod(0o644)
    swapped = repo / "_eval" / "notes.md"
    swapped.symlink_to(outside, target_is_directory=True)
    lstat = os.lstat

    def first_look(path: str | Path, *, dir_fd: int | None = None) -> os.stat_result:
        """What the entry was before the swap: a loose file."""
        return lstat(before) if Path(path).name == swapped.name else lstat(path, dir_fd=dir_fd)

    monkeypatch.setattr(os, "lstat", first_look)
    with pytest.raises(OSError):
        cycle_mod._clear_group_other(swapped)
    cycle_mod._tighten_own_paths(own_config(tmp_path, repo))
    assert outside.stat().st_mode & 0o777 == 0o755


@pytest.mark.parametrize("swapped", ["a folder in _eval", "the docs repo"])
def test_tightening_never_follows_a_folder_swapped_for_a_symlink_above_an_entry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, swapped: str
) -> None:
    """Review, 2026-10-07: ``O_NOFOLLOW`` covers an entry, not the folders above it. An entry opened by its
    full path after a folder above it had become a symlink was changed in the link's target: files outside
    the docs repo went from 0644 to 0600. Each entry is opened by its name from a descriptor on the folder
    that holds it, so what the walk listed is what is changed, wherever that folder is by then, and
    nothing where the link points."""

    def loose_tree(root: Path) -> tuple[list[Path], list[Path]]:
        folders = [root, root / "_eval", root / "_eval" / "sub"]
        folders[-1].mkdir(parents=True)
        files = [root / "NOTES.txt", root / "PLAN.txt", *(folders[-1] / f"{n}.txt" for n in "abc")]
        for f in files:
            f.write_text("1. What did Contoso decide?\n", encoding="utf-8")
            f.chmod(0o644)
        for d in folders:
            d.chmod(0o755)
        return folders, files

    repo, outside, moved = tmp_path / "docs", tmp_path / "outside", tmp_path / "moved"
    _, files = loose_tree(repo)
    outside_folders, outside_files = loose_tree(outside)
    if swapped == "the docs repo":  # while the entries at its top are made owner-only
        was, link_to, after = repo, outside, [moved / f.relative_to(repo) for f in files]
        watched = {p.stat().st_ino for p in repo.iterdir()}
    else:  # while the three files in it are
        was, link_to, after = repo / "_eval" / "sub", outside / "_eval" / "sub", files[:2]
        after += [moved / f.name for f in files[2:]]
        watched = {f.stat().st_ino for f in files[2:]}
    fchmod = os.fchmod

    def swap_at_the_first_change(fd: int, bits: int) -> None:
        if not was.is_symlink() and os.fstat(fd).st_ino in watched:
            was.rename(moved)
            was.symlink_to(link_to, target_is_directory=True)
        fchmod(fd, bits)

    monkeypatch.setattr(os, "fchmod", swap_at_the_first_change)
    cycle_mod._tighten_own_paths(own_config(tmp_path, repo))
    assert was.is_symlink(), "the folder was swapped while its entries were being changed"
    assert {mode(f) for f in outside_files} == {0o644} and {mode(d) for d in outside_folders} == {0o755}
    assert {mode(f) for f in after} == {0o600}, "every file the walk listed, where it is now"


def test_paused_source_is_skipped_and_retired_source_is_tombstoned_with_banners(
    tmp_path: Path, local_source_dir: Path
) -> None:
    config = config_with(tmp_path, local_source_dir)
    run(config)
    repo = config.docs_repo
    topic = repo / "topics" / "acme" / "acme-plan.md"
    topic.parent.mkdir(parents=True)
    fm, _ = page(repo, f"mirror/{SID}/projects/sample.docx.md")
    topic.write_text(
        "---\nkind: plan\nentity: acme\nsources:\n"
        f"  - path: ../../mirror/{SID}/projects/sample.docx.md\n"
        f"    at_rendered_sha256: {fm['rendered_sha256']}\n    role: primary\n---\n# Acme plan\n\nBody.\n",
        encoding="utf-8",
    )
    assert run(config).commit_sha is not None  # the curated page + DEPENDS.tsv ride in a commit
    assert "topics/acme/acme-plan.md" in (repo / "DEPENDS.tsv").read_text(encoding="utf-8")

    paused_text = config.config_path.read_text(encoding="utf-8").replace(
        'kind = "local"', 'kind = "local"\nstate = "paused"'
    )
    paused = parse_config(paused_text, config_path=config.config_path)
    report = run(paused)
    assert report.sources[0].skipped_reason == "paused" and report.changes == ()
    if report.commit_sha is not None:  # only the surface that names the state moves (INDEX/README)
        touched = set(git(repo, "show", "--name-only", "--format=", "HEAD").split())
        assert touched <= {"INDEX.md", "README.md", "_sync/STATE.snapshot.md"}, touched
    assert run(paused).commit_sha is None

    retired_text = config.config_path.read_text(encoding="utf-8").replace(
        'kind = "local"', 'kind = "local"\nstate = "retired"\nretired_reason = "project closed"'
    )
    retired = parse_config(retired_text, config_path=config.config_path)
    report = run(retired)
    assert report.commit_sha is not None
    assert report.changes and all(c.op.value == "D" for c in report.changes)
    fm2, _ = page(repo, f"mirror/{SID}/projects/sample.docx.md")
    assert fm2["status"] == "deleted" and fm2["reason"] == "retired:project closed"
    assert "SOURCE RETIRED — project closed" in topic.read_text(encoding="utf-8")
    assert porcelain(repo) == ""


def test_git_log_holds_only_real_change(sample_config: Config, local_source_dir: Path) -> None:
    repo = sample_config.docs_repo
    for _ in range(3):
        run(sample_config)
    assert int(git(repo, "rev-list", "--count", "HEAD").strip()) == 1
    (local_source_dir / "projects" / "sample.txt").write_text("new\n", encoding="utf-8")
    run(sample_config)
    run(sample_config)
    log = git(repo, "log", "--format=%s").splitlines()
    assert len(log) == 2 and log[0] == f"sync: 0a 1m 0r 0d {SID}"
    body = git(repo, "log", "-1", "--format=%B")
    assert "Agentsync-Run: 4" in body and "Agentsync-Mode: poll" in body
    assert subprocess.run(["git", "-C", str(repo), "diff", "--quiet", "HEAD"], check=False).returncode == 0


def test_a_walk_that_times_out_on_a_privacy_prompt_tombstones_nothing(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Field report N8: a listing macOS holds for an Allow prompt ends at its time limit as an incomplete
    source with a "click Allow" alarm. Nothing under the unlisted folder is read as gone."""
    assert run(sample_config).commit_sha is not None
    with Manifest(sample_config.state_paths.db) as m:
        before = {r.stable_id: r.state for r in m.iter_items("local-fixture")}
    assert before and RowState.TOMBSTONE not in before.values()
    real = al._list_dir
    release = threading.Event()

    def held(path: Path, *, dir_dataless: bool) -> list[os.DirEntry[str]]:
        if path == local_source_dir / "projects":
            release.wait(5.0)
        return real(path, dir_dataless=dir_dataless)

    monkeypatch.setattr(al, "_list_dir", held)
    monkeypatch.setattr(al, "LISTING_TIMEOUT_S", 0.2, raising=False)
    try:
        report = run(sample_config)
    finally:
        release.set()
    src = report.sources[0]
    assert not src.enumeration_complete
    assert any("click Allow" in a for a in src.alarms), src.alarms
    assert not any(c.op.value == "D" for c in report.changes), report.changes
    with Manifest(sample_config.state_paths.db) as m:
        after = {r.stable_id: r.state for r in m.iter_items("local-fixture")}
    assert after == before


def test_a_held_walk_fetches_nothing_under_the_held_root(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Field N8 review: once a walk times out, a fetch would lstat and read under the root macOS is holding
    and wait as the walk did. The pass fetches nothing, the pending row stays queued, and the pass is recorded
    as held (next-step says click Allow)."""
    assert run(sample_config).commit_sha is not None
    with Manifest(sample_config.state_paths.db) as m:
        sid = next(r.stable_id for r in m.iter_items("local-fixture") if not r.is_dir)
        m.set_verdict("local-fixture", sid, Verdict.DEFERRED)
    real = al._list_dir
    release = threading.Event()
    fetched: list[Path] = []

    def held(path: Path, *, dir_dataless: bool) -> list[os.DirEntry[str]]:
        if path == local_source_dir / "projects":
            release.wait(5.0)
        return real(path, dir_dataless=dir_dataless)

    def blocked(src: Path, dest: Path, budget: object) -> object:
        fetched.append(src)
        release.wait(5.0)
        raise AssertionError("fetched under a held root")

    monkeypatch.setattr(al, "_list_dir", held)
    monkeypatch.setattr(al, "materialise", blocked)
    monkeypatch.setattr(al, "LISTING_TIMEOUT_S", 0.2, raising=False)
    try:
        report = run(sample_config)
    finally:
        release.set()
    assert fetched == []
    assert any("click Allow" in a for a in report.sources[0].alarms), report.sources[0].alarms
    with Manifest(sample_config.state_paths.db) as m:
        row = m.get_item("local-fixture", sid)
        last = m.last_source_pass("local-fixture")
    assert row is not None and row.last_verdict is Verdict.DEFERRED
    assert last is not None and (last.skipped_reason or "").startswith(cycle_mod.LISTING_HELD)


def test_a_held_root_read_ends_in_the_same_click_allow_pass(
    sample_config: Config, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Field N8 review: the root's own metadata reads (lstat, volume UUID) share the listings' time limit, so
    a prompt holding them ends in the click-Allow pass instead of a sync that never returns."""
    assert run(sample_config).commit_sha is not None
    release = threading.Event()
    real = al.volume_uuid

    def held(path: Path) -> str:
        release.wait(5.0)
        return real(path)

    monkeypatch.setattr(al, "volume_uuid", held)
    monkeypatch.setattr(al, "LISTING_TIMEOUT_S", 0.2, raising=False)
    started = time.monotonic()
    try:
        report = run(sample_config)
    finally:
        release.set()
    assert time.monotonic() - started < 4.0
    src = report.sources[0]
    assert not src.enumeration_complete and any("click Allow" in a for a in src.alarms), src.alarms
    assert not any(c.op.value == "D" for c in report.changes), report.changes


# ---------------------------------------------------------------------------------------------------------
# performance paths: H0 fast path, batched writes, the land gate on a clean tree; inbox (name, size) dedup
# ---------------------------------------------------------------------------------------------------------


def test_noop_cycle_takes_the_h0_fast_path_for_every_file(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert run(sample_config).commit_sha is not None
    upserted: list[int] = []
    touched: list[int] = []
    real_many, real_touch = Manifest.upsert_observed_many, Manifest.touch_observed

    def many(self: Manifest, observations: Any, **kw: Any) -> None:
        upserted.append(len(observations))
        real_many(self, observations, **kw)

    def touch(self: Manifest, source_id: str, ids: Any, **kw: Any) -> int:
        n = real_touch(self, source_id, ids, **kw)
        touched.append(n)
        return n

    monkeypatch.setattr(Manifest, "upsert_observed_many", many)
    monkeypatch.setattr(Manifest, "touch_observed", touch)
    noop = run(sample_config)
    assert noop.exit_code == 0 and noop.commit_sha is None
    files = sum(1 for p in local_source_dir.rglob("*") if p.is_file() and not p.name.startswith("~$"))
    assert upserted == [0] and touched == [files]  # nothing decoded or rewritten, every row stamped
    with Manifest(sample_config.state_paths.db) as m:
        rows = [r for r in m.iter_items(SID) if not r.is_dir]
        run_id = m.last_runs(1)[0][0]
    assert {r.last_seen_run for r in rows} == {run_id} and {r.last_verdict for r in rows} == {
        Verdict.UNCHANGED
    }
    # an edit after the fast pass is still seen (stat tuple moved) and published
    (local_source_dir / "projects" / "sample.md").write_text("# Sample\n\nedited after a fast pass\n")
    edited = run(sample_config)
    assert edited.commit_sha is not None and upserted[-1] == 1


def test_land_gate_is_skipped_only_when_nothing_can_land(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from agentsync import lints  # noqa: PLC0415

    gates: list[int] = []
    real = lints.run_land_gate

    def spy(repo: Path, changed: Any, **kwargs: Any) -> Any:
        gates.append(len(changed))
        return real(repo, changed, **kwargs)

    monkeypatch.setattr(lints, "run_land_gate", spy)
    run(sample_config)
    assert len(gates) == 1  # first sync: content to land
    run(sample_config)
    assert len(gates) == 1  # clean POLL: nothing would be committed, no gate
    run(sample_config, mode=CycleMode.RECONCILE)
    assert len(gates) == 2  # RECONCILE always runs it
    (sample_config.docs_repo / "mirror" / "hand-edit.md").write_text("x\n")  # dirty tree, no change reported
    report = run(sample_config)
    assert len(gates) == 3 and report.commit_sha is None


INBOX_SOURCE = """
[[source]]
id = "inbox"
kind = "inbox"
path = "{path}"
quiescence_s = 0
"""


def test_an_atomic_re_export_converts_within_one_interactive_sync(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Field N4: a writer that follows the contract (temp name, then rename) is converted by the sync that
    sees it, not the one after: an interactive sync waits once for the inbox to settle; launchd never does."""
    sleeps: list[float] = []
    start = time.time_ns()

    class FakeClockInbox(al.InboxArm):
        def __init__(self, cfg: Any) -> None:
            super().__init__(cfg)
            self._clock = lambda: start + int(sum(sleeps) * 1e9)
            self._sleep = sleeps.append

    monkeypatch.setattr(cycle_mod, "InboxArm", FakeClockInbox)
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    config = config_with(
        tmp_path, local_source_dir, INBOX_SOURCE.format(path=inbox).replace("quiescence_s = 0", "")
    )
    page_rel = "mirror/inbox/export.md"

    def export(text: str) -> None:
        (inbox / "export.md.tmp").write_text(text, encoding="utf-8")
        (inbox / "export.md.tmp").rename(inbox / "export.md")
        nonlocal start
        start = time.time_ns()  # every export restarts the fake clock at the real one
        sleeps.clear()

    export("# Export\n\nfirst version\n")
    background = run(config, only=["inbox"])  # an explicit mode (launchd) never waits
    assert sleeps == []
    assert not next(s for s in background.sources if s.source_id == "inbox").enumeration_complete
    assert not (config.docs_repo / page_rel).exists()

    report = run(config, only=["inbox"], mode=None)
    assert report.exit_code == 0, report
    assert len(sleeps) == 1 and 0 < sleeps[0] <= 60
    assert next(s for s in report.sources if s.source_id == "inbox").enumeration_complete
    assert "first version" in page(config.docs_repo, page_rel)[1]

    export("# Export\n\nsecond version\n")  # the re-export: same name, a new inode
    report = run(config, only=["inbox"], mode=None)
    assert report.exit_code == 0, report
    assert len(sleeps) == 1
    assert "second version" in page(config.docs_repo, page_rel)[1]


def test_one_interactive_sync_waits_at_most_60s_over_every_inbox(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Field N4 review: the settle wait is bounded once per sync, not once per inbox, whatever quiescence_s
    is, so an agent's ``agentsync sync`` stays well inside its command timeout."""
    sleeps: list[float] = []
    start = time.time_ns()
    offset = 570 * 1_000_000_000  # files written now look 570 s old: 30 s left of a 600 s window

    class FakeClockInbox(al.InboxArm):
        def __init__(self, cfg: Any) -> None:
            super().__init__(cfg)
            self._clock = lambda: start + offset + int(sum(sleeps) * 1e9)
            self._sleep = sleeps.append

    monkeypatch.setattr(cycle_mod, "InboxArm", FakeClockInbox)
    sources = ""
    for sid in ("inbox-a", "inbox-b"):
        inbox = tmp_path / sid
        inbox.mkdir()
        (inbox / "older.md").write_text(f"# Older\n\n{sid}\n", encoding="utf-8")
        (inbox / "fresh.md").write_text("# Fresh\n\nstill being written\n", encoding="utf-8")
        os.utime(inbox / "fresh.md", ns=(start + offset, start + offset))  # changed just now
        sources += INBOX_SOURCE.format(path=inbox).replace('"inbox"\n', f'"{sid}"\n', 1)
    sources = sources.replace("quiescence_s = 0", "quiescence_s = 600")
    config = config_with(tmp_path, local_source_dir, sources)
    report = run(config, only=["inbox-a", "inbox-b"], mode=None)
    assert report.exit_code == 0, report
    assert 0 < sum(sleeps) <= 60, sleeps
    for sid in ("inbox-a", "inbox-b"):
        assert sid in page(config.docs_repo, f"mirror/{sid}/older.md")[1]
        assert not (config.docs_repo / "mirror" / sid / "fresh.md").exists()


def test_inbox_drop_matching_a_graph_file_by_name_and_size_is_refused_unread(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from agentsync.arm_local import InboxArm  # noqa: PLC0415
    from agentsync.model import RowState  # noqa: PLC0415

    notes = b"# Notes\n\nThe purchase order is approved.\n"
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    (inbox / "notes (1).md").write_bytes(b"#" * len(notes))  # same size, other bytes: H1 differs
    (inbox / "NOTES - Copy.md").write_bytes(b"x" * len(notes))  # case + Windows copy suffix
    (inbox / "notes.md").write_bytes(notes + b"more")  # same name, other size: not a duplicate
    (inbox / "other.md").write_bytes(b"#" * len(notes))  # same size, other name
    config = config_with(tmp_path, local_source_dir, GRAPH_SOURCE + INBOX_SOURCE.format(path=inbox))
    drive = FakeDrive({"I1": ("notes.md", notes)})
    client = GraphClient(
        FakeTokens(),
        user_agent="NONISV|test|agentsync/0",
        transport=httpx.MockTransport(drive.handler),
        sleep=lambda _s: None,
    )
    fetched: list[str] = []
    real_fetch = InboxArm.fetch

    def spy(self: InboxArm, item: Any, dest: Path, budget: Any) -> Any:
        fetched.append(item.rel_path)
        return real_fetch(self, item, dest, budget)

    monkeypatch.setattr(InboxArm, "fetch", spy)
    try:
        report = run(config, client=client, only=["drive", "inbox"])
    finally:
        client.close()
    assert report.exit_code == 0, report
    assert sorted(fetched) == ["notes.md", "other.md"]  # the two duplicates were never read
    with Manifest(config.state_paths.db) as m:
        rows = {r.rel_path: r for r in m.iter_items("inbox")}
    want = "duplicate-of drive (mirror/drive/projects/notes.md)"
    for rel in ("notes (1).md", "NOTES - Copy.md"):
        assert (rows[rel].state, rows[rel].state_reason, rows[rel].last_verdict) == (
            RowState.QUARANTINED,
            want,
            Verdict.QUARANTINED,
        )
    assert rows["notes.md"].state is RowState.LIVE and rows["other.md"].state is RowState.LIVE
    assert "duplicate-of drive" in (config.docs_repo / "_sync" / "QUARANTINE.tsv").read_text(encoding="utf-8")


# ---------------------------------------------------------------------------------------------------------
# the materialise byte budget charges downloads only (L3)
# ---------------------------------------------------------------------------------------------------------


def _mark_online_only(monkeypatch: pytest.MonkeyPatch, *paths: Path) -> None:
    """Mock SF_DATALESS on these files (a test has no File Provider): the local arm lists them as dataless
    and materialise charges them as a download; every other file stays local."""
    from agentsync import arm_local, materialise  # noqa: PLC0415

    inos = {p.stat().st_ino for p in paths}
    real = materialise.is_dataless

    def fake(st: Any) -> bool:
        return st.st_ino in inos or real(st)

    monkeypatch.setattr(materialise, "is_dataless", fake)
    monkeypatch.setattr(arm_local, "is_dataless", fake)


def test_budget_0_converts_every_local_file_and_defers_only_online_only_ones(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``budget_bytes=0`` (install.sh's first sync): already-local files never consume the byte budget, so
    all of them are converted; only the online-only files wait for a run whose budget allows the download."""
    online = [local_source_dir / "projects" / "sample.pptx", local_source_dir / "projects" / "sample.pdf"]
    _mark_online_only(monkeypatch, *online)
    first = run(sample_config, budget_bytes=0)
    [src] = first.sources
    assert first.exit_code == 0
    assert (src.deferred, src.deferred_online_only, src.materialised_bytes) == (2, 2, 0)
    assert src.converted >= 8, "every local fixture is read and converted"
    assert not any("exceeds the per-cycle budget" in a and "sample.docx" in a for a in src.alarms)
    mirror = sample_config.docs_repo / "mirror" / SID / "projects"
    assert (mirror / "sample.docx.md").is_file() and len(list(mirror.iterdir())) >= 6
    assert not any(mirror.glob("sample.pptx*")) and not any(mirror.glob("sample.pdf*"))
    second = run(sample_config)  # the per-source budget (50MB): the two downloads fit
    [src] = second.sources
    assert (src.converted, src.deferred, src.deferred_online_only) == (2, 0, 0)
    assert src.materialised_bytes == sum(p.stat().st_size for p in online), "only downloads are charged"
    assert any(mirror.glob("sample.pptx*")) and (mirror / "sample.pdf.md").is_file()


# ---------------------------------------------------------------------------------------------------------
# a page too long for any sidecar settles as one quarantined item
# ---------------------------------------------------------------------------------------------------------

APPENDIX_ROWS = 45_000
"""Enough rows to pass ``max_page_bytes``: the page of such a file needs a ``full-text.txt`` sidecar."""
_EXPORT = "Fabrikam Totals By Quarter - FY26 Forecast - All Regions - Final Export From The Data Warehouse v2"
CRAMPED = f"Contoso Travel Forms/Charlie Photo Shoots/{_EXPORT} - Appendix With Every Row.txt"
"""A made-up path whose page is over 183 characters: no ``.files/<8 hex>.txt`` fits beside it."""
MOVES = {
    # what is renamed: (a path with room for the sidecar, one without)
    "file": ("Contoso Travel Forms/appendix.txt", CRAMPED),
    "folder": (
        f"Forms/{_EXPORT}.txt",
        f"Contoso Travel Forms And Photo Shoots - Kept For Next Year/{_EXPORT}.txt",
    ),
}


def _move(root: Path, old: str, new: str, *, folder: bool) -> None:
    """Rename the file at ``old`` to ``new``, or the folder above it. A file's own rename moves its ctime,
    so the next pass reads it again; under a renamed folder its stat holds and nothing is read."""
    src, dst = root / old, root / new
    if folder:
        src, dst = src.parent, dst.parent
    dst.parent.mkdir(parents=True, exist_ok=True)
    src.rename(dst)


def _capped_at(root: Path, rel: str) -> None:
    (root / rel).parent.mkdir(parents=True, exist_ok=True)
    (root / rel).write_text("every row of the appendix\n" * APPENDIX_ROWS, encoding="utf-8")


def _appendix_rows(sidecar: Path) -> int:
    """How many of the file's rows ``sidecar`` holds (a count, so a failure prints no megabyte diff)."""
    return sidecar.read_bytes().count(b"every row of the appendix\n")


def _path_stub(config: Config, rel: str) -> bool:
    """True when ``rel`` is settled as the stub for a page with no room for its sidecar, and nothing else is
    beside that page."""
    repo, at = config.docs_repo, slug.mirror_rel_path(SID, rel)
    fm = page(repo, at)[0]
    row = _file_rows(config)[rel]
    return (
        (fm["status"], fm["source_path"]) == ("unreadable", rel)
        and "path too long" in fm["reason"]
        and not (repo / at).with_suffix(".files").exists()
        and (row.state, row.last_verdict) == (RowState.QUARANTINED, Verdict.QUARANTINED)
        and "path too long" in (row.state_reason or "")
    )


def test_a_capped_file_whose_page_leaves_no_room_for_a_sidecar_is_one_quarantined_item(
    sample_config: Config, local_source_dir: Path
) -> None:
    """The item gets a stub and is settled; every other file is converted and the commit lands."""
    rel = CRAMPED
    page = slug.mirror_rel_path(SID, rel)
    assert len(page) > 183
    _capped_at(local_source_dir, rel)
    first = run(sample_config)
    [src] = first.sources
    assert first.exit_code == 0 and first.commit_sha is not None, first
    assert src.errors == () and src.counts.get(Verdict.QUARANTINED) == 1
    assert not [f for f in first.lint_findings if f.code == "PATH"]
    repo = sample_config.docs_repo
    assert (repo / "mirror" / SID / "projects" / "sample.docx.md").is_file()
    assert "status: unreadable" in (repo / page).read_text(encoding="utf-8")
    assert not (repo / page).with_suffix(".files").exists()
    assert max(len(p.relative_to(repo).as_posix()) for p in (repo / "mirror").rglob("*")) <= 200
    with Manifest(sample_config.state_paths.db) as m:
        rows = {r.rel_path: r for r in m.iter_items(SID)}
    assert (rows[rel].state, rows[rel].last_verdict) == (RowState.QUARANTINED, Verdict.QUARANTINED)
    assert "path too long" in (rows[rel].state_reason or "")
    assert all(r.state is RowState.LIVE for name, r in rows.items() if name.startswith("projects/sample."))
    second = run(sample_config)  # settled: not read again, nothing to commit
    [src] = second.sources
    assert (second.exit_code, second.commit_sha, src.converted, src.errors) == (0, None, 0, ())


@pytest.mark.parametrize("renamed", ["file", "folder"])
def test_a_rename_that_leaves_no_room_for_the_sidecar_settles_as_the_stub_at_the_new_path(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch, renamed: str
) -> None:
    """The page cannot follow the file, because its sidecar has no name that fits there. The item becomes
    the stub a first conversion at that path gets: at the new path, nothing left at the old one, settled.
    A renamed folder needs no read for that."""
    roomy, cramped = MOVES[renamed]
    old, new = slug.mirror_rel_path(SID, roomy), slug.mirror_rel_path(SID, cramped)
    assert len(old) <= 183 < len(new)
    _capped_at(local_source_dir, roomy)
    assert run(sample_config).exit_code == 0
    repo = sample_config.docs_repo
    side = repo / sidecar_rel(old, "full-text.txt")
    assert page(repo, old)[0]["status"] == "current" and _appendix_rows(side) == APPENDIX_ROWS
    _move(local_source_dir, roomy, cramped, folder=renamed == "folder")
    fetched = _fetches(monkeypatch)
    second = run(sample_config)
    [src] = second.sources
    assert second.exit_code == 0 and second.commit_sha is not None and src.errors == ()
    assert fetched == ([cramped] if renamed == "file" else [])
    assert _path_stub(sample_config, cramped)
    assert not (repo / old).exists() and not side.parent.exists()
    assert not [f for f in second.lint_findings if f.code == "PATH"]
    third = run(sample_config)  # settled: not read again, nothing to commit
    assert (third.exit_code, third.commit_sha, third.sources[0].converted) == (0, None, 0)
    assert len(fetched) == (1 if renamed == "file" else 0) and porcelain(repo) == ""


@pytest.mark.parametrize("renamed", ["file", "folder"])
def test_shortening_the_path_of_a_stubbed_capped_file_publishes_it(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch, renamed: str
) -> None:
    """The stub says to shorten a folder or file name, and doing so clears it although the bytes did not
    change. A renamed file is read in the same pass; a file under a renamed folder is queued by that pass
    and read by the next one."""
    roomy, cramped = MOVES[renamed]
    _capped_at(local_source_dir, cramped)
    assert run(sample_config).exit_code == 0 and _path_stub(sample_config, cramped)
    _move(local_source_dir, cramped, roomy, folder=renamed == "folder")
    fetched = _fetches(monkeypatch)
    report = run(sample_config)
    if renamed == "folder":
        assert report.exit_code == 0 and fetched == []
        assert _file_rows(sample_config)[roomy].last_verdict is Verdict.MAYBE_CHANGED
        report = run(sample_config)
    assert report.exit_code == 0 and report.commit_sha is not None and fetched == [roomy]
    repo, new = sample_config.docs_repo, slug.mirror_rel_path(SID, roomy)
    fm = page(repo, new)[0]
    assert (fm["status"], fm["source_path"]) == ("current", roomy)
    assert _appendix_rows(repo / sidecar_rel(new, "full-text.txt")) == APPENDIX_ROWS
    assert not (repo / slug.mirror_rel_path(SID, cramped)).exists()
    row = _file_rows(sample_config)[roomy]
    assert (row.state, row.state_reason, row.last_verdict) == (RowState.LIVE, None, Verdict.UNCHANGED)
    settled = run(sample_config)
    assert (settled.exit_code, settled.commit_sha, fetched) == (0, None, [roomy])


def test_renaming_a_stubbed_drive_file_to_a_shorter_name_downloads_and_publishes_it(
    drive_env: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A drive rename is METADATA_ONLY (the hash holds), so nothing is downloaded for it. The stub of a page
    with no room for its sidecar is the exception: the rename queues one download, and the next pass
    publishes the content. A later change that is not a rename downloads nothing."""
    config, drive, client = drive_env
    data = b"every row of the appendix\n" * APPENDIX_ROWS
    long_name = f"{_EXPORT} - Appendix With Every Row And Every Column Of Every Sheet.txt"
    drive.files["I1"] = (long_name, data)
    stub = slug.mirror_rel_path("drive", f"Projects/{long_name}")
    assert len(stub) > 183
    assert run(config, client=client, only=["drive"]).exit_code == 0
    repo = config.docs_repo
    assert "path too long" in page(repo, stub)[0]["reason"]
    fetched = _fetches(monkeypatch, DriveArm)
    drive.version["I1"] += 1  # a metadata change that is not a rename: the stub stays, nothing is read
    assert run(config, client=client, only=["drive"]).exit_code == 0 and fetched == []
    assert page(repo, stub)[0]["status"] == "unreadable"
    drive.version["I1"] += 1
    drive.files["I1"] = ("appendix.txt", data)  # renamed: same bytes, same quickXorHash
    assert run(config, client=client, only=["drive"]).exit_code == 0 and fetched == []
    [row] = _file_rows(config, "drive").values()
    assert (row.rel_path, row.last_verdict) == ("Projects/appendix.txt", Verdict.MAYBE_CHANGED)
    report = run(config, client=client, only=["drive"])
    assert report.exit_code == 0 and fetched == ["Projects/appendix.txt"]
    new = "mirror/drive/projects/appendix.txt.md"
    assert page(repo, new)[0]["status"] == "current" and not (repo / stub).exists()
    assert _appendix_rows(repo / sidecar_rel(new, "full-text.txt")) == APPENDIX_ROWS
    [row] = _file_rows(config, "drive").values()
    assert (row.state, row.state_reason, row.last_verdict) == (RowState.LIVE, None, Verdict.UNCHANGED)


# ---------------------------------------------------------------------------------------------------------
# scope: a row the config no longer covers is never work, and is retired on an incomplete walk too
# ---------------------------------------------------------------------------------------------------------

KICKOFF = "projects/acme/Kickoff Notes.docx"


def _incomplete(config: Config, *exclude: str) -> Config:
    """``config`` with ``exclude`` added to its one source and a sentinel that is not there, so every walk
    is incomplete (as on a Mac with one folder the agent cannot list)."""
    [src] = config.sources
    narrowed = dataclasses.replace(src, exclude=(*src.exclude, *exclude), sentinel="no-such-sentinel")
    return dataclasses.replace(config, sources=(narrowed,))


def _fetches(monkeypatch: pytest.MonkeyPatch, arm: type[Any] = LocalArm) -> list[str]:
    """The rel_path of every file the cycle reads through ``arm.fetch`` from now on."""
    fetched: list[str] = []
    real = arm.fetch

    def spy(self: Any, item: Any, dest: Path, budget: Any) -> Any:
        fetched.append(item.rel_path)
        return real(self, item, dest, budget)

    monkeypatch.setattr(arm, "fetch", spy)
    return fetched


def _file_rows(config: Config, sid: str = SID) -> dict[str, Any]:
    with Manifest(config.state_paths.db) as m:
        return {r.rel_path: r for r in m.iter_items(sid) if not r.is_dir}


@pytest.mark.parametrize(
    ("queued", "exclude"),
    [
        ("projects/sample.pptx", "sample.pptx"),  # a file glob
        (KICKOFF, "acme"),  # a bare folder name: no file glob matches, only the walk's folder rule
        (KICKOFF, "projects/acme"),  # an anchored folder, no trailing slash
        (KICKOFF, "acme/"),
    ],
)
def test_a_queued_row_excluded_later_is_retired_unread_when_the_walk_is_incomplete(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch, queued: str, exclude: str
) -> None:
    """An online-only file deferred by one run and excluded before the next is never fetched or converted.
    An incomplete walk prunes no row, so the queue applies the walk's own scope and the row is retired: no
    purge, and the loop stops counting it. Rows still in scope are untouched: a queued one is worked, and
    one the incomplete walk did not list stays unknown, never gone."""
    kept, absent = "projects/sample.pdf", "projects/sample.md"
    _mark_online_only(monkeypatch, local_source_dir / queued, local_source_dir / kept)
    first = run(sample_config, budget_bytes=0)
    assert first.sources[0].deferred_online_only == 2
    config = _incomplete(sample_config, exclude)
    (local_source_dir / absent).unlink()  # in scope, and not there for the next walk to list
    fetched = _fetches(monkeypatch)
    second = run(config)
    [rep] = second.sources
    assert second.exit_code == 0 and not rep.enumeration_complete and not rep.breaker_tripped
    assert fetched == [kept] and (rep.converted, rep.deferred) == (1, 0)
    assert rep.materialised_bytes == (local_source_dir / kept).stat().st_size
    assert any("1 file(s) now outside it retired" in a for a in rep.alarms), rep.alarms
    rows = _file_rows(config)
    assert (rows[queued].state, rows[queued].state_reason) == (RowState.TOMBSTONE, "retired:scope-change")
    assert [rel for rel, r in rows.items() if r.state is RowState.TOMBSTONE] == [queued]
    repo = config.docs_repo
    assert page(repo, slug.mirror_rel_path(SID, absent))[0]["status"] == "current"
    assert not (repo / slug.mirror_rel_path(SID, queued)).exists()
    assert governance.pending_purges(config.state_paths.root) == []
    step = loop.next_step(config)
    assert step.rule != 3 and not any("online-only" in line for line in step.lines()), step.lines()
    third = run(config)  # settled: nothing is read, nothing is retired twice, the absent file is still held
    assert fetched == [kept] and not any("retired" in a for a in third.sources[0].alarms)
    assert _file_rows(config)[absent].state is RowState.LIVE


@pytest.mark.parametrize("requeued", [False, True], ids=["settled", "queued-again"])
def test_a_published_file_excluded_later_is_retired_on_an_incomplete_walk_and_comes_back(
    sample_config: Config, monkeypatch: pytest.MonkeyPatch, requeued: bool
) -> None:
    """The page of a file the config no longer covers is retired by the next pass, complete or not: marked
    retired (never deleted upstream), no purge queued, the breaker not involved. Queued again or not, the
    file is not read and the loop does not wait for it. Taking the exclude away brings the page back."""
    assert run(sample_config).exit_code == 0
    repo, rel = sample_config.docs_repo, slug.mirror_rel_path(SID, KICKOFF)
    assert page(repo, rel)[0]["status"] == "current"
    if requeued:  # on this Mac and waiting for a re-read, as after a [policy] change
        stable_id = _file_rows(sample_config)[KICKOFF].stable_id
        with Manifest(sample_config.state_paths.db) as m:
            m.set_verdict(SID, stable_id, Verdict.MAYBE_CHANGED)
    excluded = _incomplete(sample_config, "acme")
    fetched = _fetches(monkeypatch)
    second = run(excluded)
    [rep] = second.sources
    assert second.exit_code == 0 and not rep.enumeration_complete and not rep.breaker_tripped
    assert fetched == [] and [(c.op.value, c.path) for c in second.changes] == [("D", rel)]
    assert page(repo, rel)[0]["reason"] == "retired:scope-change"
    assert "[RETIRED]" in (repo / rel).read_text(encoding="utf-8")
    assert governance.pending_purges(excluded.state_paths.root) == []
    assert loop.next_step(excluded).rule != 3
    back = run(_incomplete(sample_config))  # the exclude is gone; the walk is still incomplete
    assert back.exit_code == 0 and fetched == [KICKOFF]
    assert page(repo, rel)[0]["status"] == "current"


def _repointed(config: Config, root: Path, name: str) -> Config:
    """``config`` with its one source pointed at ``root``, one folder above the tree it mirrors (``name``),
    and include and the sentinel re-anchored to keep the same files: one edit of sources.toml."""
    [src] = config.sources
    moved = dataclasses.replace(src, path=root, include=(f"{name}/**",), sentinel=f"{name}/{src.sentinel}")
    return dataclasses.replace(config, sources=(moved,))


def test_a_source_pointed_at_a_new_root_retires_no_row_it_has_not_listed_there(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A stored path is relative to the root it was listed under. When ``path`` moves up a folder and
    include is re-anchored in the same edit, every stored path fails the new scope, and that says nothing
    about the files. While the new root cannot be listed (a typo, a volume not mounted yet) no row is
    retired; once it can, the pages move in place and nothing is read again."""
    assert run(sample_config).exit_code == 0
    before = {rel: r.state for rel, r in _file_rows(sample_config).items()}
    assert len(before) > 5 and set(before.values()) == {RowState.LIVE}
    parent, name = local_source_dir.parent, local_source_dir.name
    fetched = _fetches(monkeypatch)
    for _pass in range(2):  # the first pass under the new root, and one after it
        absent = run(_repointed(sample_config, parent / "not-mounted-yet", name))
        [rep] = absent.sources
        assert absent.exit_code == 0 and not rep.enumeration_complete
        assert not any("retired" in a for a in rep.alarms), rep.alarms
        assert {rel: r.state for rel, r in _file_rows(sample_config).items()} == before
    there = _repointed(sample_config, parent, name)
    back = run(there)
    assert back.exit_code == 0 and back.sources[0].enumeration_complete and fetched == []
    rows = _file_rows(there)
    assert {rel: r.state for rel, r in rows.items()} == {f"{name}/{rel}": RowState.LIVE for rel in before}
    repo = there.docs_repo
    assert page(repo, slug.mirror_rel_path(SID, f"{name}/{KICKOFF}"))[0]["status"] == "current"
    assert not (repo / slug.mirror_rel_path(SID, KICKOFF)).exists()
    # Under the root it has now, an exclude is judged again on an incomplete pass.
    narrowed = _incomplete(there, "acme")
    report = run(narrowed)
    assert report.exit_code == 0 and fetched == [] and not report.sources[0].enumeration_complete
    rows = _file_rows(narrowed)
    assert [rel for rel, r in rows.items() if r.state is RowState.TOMBSTONE] == [f"{name}/{KICKOFF}"]
    assert rows[f"{name}/{KICKOFF}"].state_reason == "retired:scope-change"


def test_a_row_listed_under_the_old_root_only_is_left_to_a_complete_pass(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``path`` moves and the same edit excludes a folder. The rows under that folder are never listed under
    the new root, so their stored paths stay relative to the old one and an incomplete pass cannot judge
    them. The first complete pass retires them, as it does after any scope change."""
    assert run(sample_config).exit_code == 0
    parent, name = local_source_dir.parent, local_source_dir.name
    [src] = _repointed(sample_config, parent, name).sources
    moved = dataclasses.replace(src, exclude=(*src.exclude, "acme"))
    stopped = dataclasses.replace(sample_config, sources=(dataclasses.replace(moved, sentinel="none"),))
    fetched = _fetches(monkeypatch)
    for _pass in range(2):
        report = run(stopped)
        assert report.exit_code == 0 and not report.sources[0].enumeration_complete
        assert not any("retired" in a for a in report.sources[0].alarms)
        assert _file_rows(stopped)[KICKOFF].state is RowState.LIVE  # still under its old-root path
    complete = run(dataclasses.replace(sample_config, sources=(moved,)))
    assert complete.exit_code == 0 and complete.sources[0].enumeration_complete and fetched == []
    row = _file_rows(sample_config)[KICKOFF]
    assert (row.state, row.state_reason) == (RowState.TOMBSTONE, "retired:scope-change")
    assert governance.pending_purges(sample_config.state_paths.root) == []


@pytest.mark.parametrize("edited_in_that_run", [False, True])
def test_a_manifest_with_no_recorded_root_trusts_its_paths_unless_that_run_changed_the_scope(
    sample_config: Config, edited_in_that_run: bool
) -> None:
    """A manifest written before the root was recorded has rows under the root the source has: an exclude
    added later retires them on an incomplete pass. If sources.toml changed in the very run that makes the
    first record, ``path`` may have moved in that edit, and the rows it has not listed since are left to a
    complete pass."""
    assert run(sample_config).exit_code == 0
    with Manifest(sample_config.state_paths.db) as m:
        m._db.execute("DELETE FROM meta WHERE key = ?", (cycle_mod._SCOPE_ROOT_META + SID,))
    if not edited_in_that_run:
        assert run(_incomplete(sample_config)).exit_code == 0  # the first record, with nothing changed
    excluded = _incomplete(sample_config, "acme")
    assert run(excluded).exit_code == 0
    row = _file_rows(excluded)[KICKOFF]
    assert (row.state is RowState.TOMBSTONE) is (not edited_in_that_run), row.state
    assert run(excluded).exit_code == 0
    assert _file_rows(excluded)[KICKOFF].state is row.state  # the next incomplete pass decides the same


def test_an_inbox_file_excluded_later_is_retired_on_an_incomplete_walk(
    tmp_path: Path, local_source_dir: Path
) -> None:
    inbox = tmp_path / "inbox"
    (inbox / "old").mkdir(parents=True)
    (inbox / "keep.md").write_text("# Keep\n\nstays\n", encoding="utf-8")
    (inbox / "old" / "drop.md").write_text("# Drop\n\nleaves\n", encoding="utf-8")
    table = INBOX_SOURCE.format(path=inbox)
    assert run(config_with(tmp_path, local_source_dir, table), only=["inbox"]).exit_code == 0
    narrowed = config_with(
        tmp_path, local_source_dir, table + 'exclude = ["old"]\nsentinel = "no-such-sentinel"\n'
    )
    report = run(narrowed, only=["inbox"])
    rep = next(s for s in report.sources if s.source_id == "inbox")
    assert report.exit_code == 0 and not rep.enumeration_complete
    repo = narrowed.docs_repo
    assert page(repo, "mirror/inbox/old/drop.md")[0]["reason"] == "retired:scope-change"
    assert page(repo, "mirror/inbox/keep.md")[0]["status"] == "current"
    assert governance.pending_purges(narrowed.state_paths.root) == []


def test_a_queued_drive_file_excluded_later_is_not_downloaded_before_the_listing_reaches_it(
    drive_env: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The drive arm tombstones a known file the new globs exclude when its listing reaches it. Until then
    the queue does not download it, and the row is left for the arm to settle."""
    config, _drive, client = drive_env
    first = run(config, client=client, only=["drive"], budget_bytes=0)
    assert next(s for s in first.sources if s.source_id == "drive").deferred == 1
    sources = tuple(
        dataclasses.replace(s, exclude=(*s.exclude, "notes.md")) if s.id == "drive" else s
        for s in config.sources
    )
    narrowed = dataclasses.replace(config, sources=sources)

    def interrupted(self: DriveArm, cursor: str | None, *, full: bool) -> ScanResult:
        return ScanResult(
            source_id=self.source_id,
            pass_kind=PassKind.FULL,
            items=(),
            new_cursor=None,
            enumeration_complete=False,
        )

    monkeypatch.setattr(DriveArm, "scan", interrupted)
    fetched = _fetches(monkeypatch, DriveArm)
    second = run(narrowed, client=client, only=["drive"])
    assert second.exit_code == 0 and fetched == []
    assert not (narrowed.docs_repo / "mirror" / "drive" / "projects" / "notes.md").exists()
    [row] = _file_rows(narrowed, "drive").values()
    assert (row.state, row.last_verdict) == (RowState.LIVE, Verdict.DEFERRED)


def test_an_arm_without_a_path_scope_has_every_queued_row_worked(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The mail and Teams arms have no ``in_scope``: include/exclude are not theirs to apply, so the queue
    filters nothing for them and no row of theirs is retired for its path."""

    class NoScope:
        def __init__(self, inner: Any) -> None:
            self.source_id, self.kind, self._inner = inner.source_id, inner.kind, inner

        def scan(self, cursor: str | None, *, full: bool) -> ScanResult:
            return self._inner.scan(cursor, full=full)  # type: ignore[no-any-return]

        def fetch(self, item: Any, dest_dir: Path, budget: Any) -> Any:
            return self._inner.fetch(item, dest_dir, budget)

    deck = "projects/sample.pptx"
    _mark_online_only(monkeypatch, local_source_dir / deck)
    assert run(sample_config, budget_bytes=0).sources[0].deferred_online_only == 1
    real = cycle_mod.build_arms
    monkeypatch.setattr(
        cycle_mod, "build_arms", lambda *a, **k: {sid: NoScope(arm) for sid, arm in real(*a, **k).items()}
    )
    fetched = _fetches(monkeypatch)
    report = run(_incomplete(sample_config, "sample.pptx"))
    assert report.exit_code == 0 and fetched == [deck]
    assert _file_rows(sample_config)[deck].state is not RowState.TOMBSTONE


# ---------------------------------------------------------------------------------------------------------
# [policy] re-screen backlog: a finished re-screen stays finished
# ---------------------------------------------------------------------------------------------------------


def test_a_finished_policy_rescreen_reports_no_backlog_for_a_file_that_waits_later(
    tmp_path: Path, local_source_dir: Path, fixture_files: dict[str, Path], monkeypatch: pytest.MonkeyPatch
) -> None:
    """A finished re-screen stores "" as its marker. A label-capable file that waits afterwards for another
    reason (online-only, no download budget) is not re-screen backlog, and STATE.md does not say it is."""
    assert run(config_with(tmp_path, local_source_dir)).exit_code == 0
    policy = '\n[policy]\nexclude_label_ids = ["00000000-0000-4000-8000-00000000c0de"]\n'
    config = config_with(tmp_path, local_source_dir, policy)
    state_md = config.docs_repo / "_sync" / "STATE.md"
    assert run(config).exit_code == 0  # the changed policy re-screens every label-capable file in this run
    with Manifest(config.state_paths.db) as m:
        assert m.get_meta(cycle_mod._RESCREEN_META) == ""
    assert "## Content policy" not in state_md.read_text(encoding="utf-8")
    late = local_source_dir / "projects" / "late.pdf"
    shutil.copy2(fixture_files["sample.pdf"], late)
    _mark_online_only(monkeypatch, late)
    report = run(config, budget_bytes=0)
    assert report.exit_code == 0 and report.sources[0].deferred_online_only == 1
    assert "## Content policy" not in state_md.read_text(encoding="utf-8")
    with Manifest(config.state_paths.db) as m:
        assert m.get_meta(cycle_mod._RESCREEN_META) == ""


# ---------------------------------------------------------------------------------------------------------
# on-device OCR: the cycle looks for the engine, reads images on this Mac, within a time budget
# ---------------------------------------------------------------------------------------------------------

SITE_PLAN = "projects/Contoso Site Plan.png"


def _use_ocr(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, **kw: Any) -> Any:
    """Every cycle from now on finds this fake OCR engine, as after an install that built the helper."""
    engine = fake_engine(tmp_path / "ocr-bin", **kw)
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    return engine


def _image_page(config: Config, rel: str, sid: str = SID) -> tuple[dict[str, Any], str]:
    return page(config.docs_repo, slug.mirror_rel_path(sid, rel))


def test_an_image_on_this_mac_is_read_once_in_the_staging_folder(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    engine = _use_ocr(monkeypatch, tmp_path)
    picture(local_source_dir / SITE_PLAN, "Loading dock", "North gate: closed")
    fetched = _fetches(monkeypatch)
    assert run(sample_config).exit_code == 0
    fm, body = _image_page(sample_config, SITE_PLAN)
    assert fm["status"] == "current" and fm["converter"].startswith("image-ocr@2.0.0+ocr-paper-vision-")
    assert fm["source_title"] == "Contoso Site Plan.png", "the unit has no title: the shown name stands in"
    assert fm["summary"] == "Image 800x600 px; OCR: 2 line(s)"
    assert body.startswith("> [UNTRUSTED CONTENT]") and body.rstrip().endswith(
        "[image · 800x600 px · text read by on-device OCR (Apple Vision)]\n\nLoading dock\nNorth gate: closed"
    )
    (call,) = calls(engine.helper)
    staging = sample_config.state_paths.staging.resolve()
    assert Path(call["cwd"]).parent == staging, "the helper works under the cycle's staging folder"
    assert Path(call["args"][-1]).parent == Path(call["cwd"]) and list(staging.iterdir()) == []
    assert SITE_PLAN in fetched
    again = run(sample_config)
    assert again.commit_sha is None and fetched.count(SITE_PLAN) == 1 and len(calls(engine.helper)) == 1


def test_the_cycle_looks_for_the_engine_once_and_never_under_a_dry_run(
    sample_config: Config, monkeypatch: pytest.MonkeyPatch
) -> None:
    looked: list[Path] = []
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, cache_dir: looked.append(cache_dir))
    assert run(sample_config, mode=CycleMode.DRY_RUN).exit_code == 0 and looked == []
    assert run(sample_config).exit_code == 0 and looked == [sample_config.cache_dir]


def test_the_cycle_engine_is_what_ocr_engine_finds_and_none_when_ocr_is_off(
    sample_config: Config, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(sys, "platform", "darwin")
    helper = write_fake(ocr._helper_path(sample_config.cache_dir))
    assert cycle_mod._cycle_ocr(sample_config) is None, "AGENTSYNC_OCR=0, the suite's own switch"
    monkeypatch.delenv("AGENTSYNC_OCR")
    found = cycle_mod._cycle_ocr(sample_config)
    assert found is not None and found.helper == helper and found.spent_s == 0.0
    assert found.identity == "ocr-paper-vision-r2-h0.3.0-l1"
    off = dataclasses.replace(sample_config, convert=dataclasses.replace(sample_config.convert, ocr=False))
    assert cycle_mod._cycle_ocr(off) is None, "[convert] ocr = false"


def test_the_cycle_engine_adds_up_the_helper_time_of_reads_that_fail_too(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plain = fake_engine(tmp_path / "bin")
    metered = cycle_mod._CycleOcr(plain.helper, name=plain.name, revision=2, helper_version="0.3.0")
    ticks = iter([10.0, 12.5, 20.0, 21.0])
    monkeypatch.setattr(cycle_mod, "_ocr_clock", lambda: next(ticks))
    (frames,) = metered.read([picture(tmp_path / "a.png", "read")], work_dir=tmp_path, budget_s=60)
    assert ocr.text_lines(frames[0]) == ["read"] and metered.spent_s == 2.5
    with pytest.raises(ocr.OcrError):
        metered.read([tmp_path / "gone.png"], work_dir=tmp_path, budget_s=60)
    assert metered.spent_s == 3.5 and metered.identity == plain.identity


def test_the_ocr_budget_defers_the_images_past_it_and_a_cache_hit_costs_nothing(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A folder of thousands of screenshots is read over many cycles. Each read here 'takes' 100 s, so the
    180 s budget is passed by the second and the rest wait, as files past ``max_files`` do: the loop says to
    sync again."""
    assert run(sample_config).exit_code == 0
    engine = _use_ocr(monkeypatch, tmp_path)
    ticks = itertools.count(0.0, 100.0)
    monkeypatch.setattr(cycle_mod, "_ocr_clock", lambda: next(ticks))
    shots = [picture(local_source_dir / "shots" / f"shot {n}.png", f"screenshot {n}") for n in range(5)]
    pages = [sample_config.docs_repo / slug.mirror_rel_path(SID, f"shots/{p.name}") for p in shots]
    seen: list[tuple[int, int, int]] = []
    for _cycle in range(3):
        [rep] = run(sample_config).sources
        seen.append((rep.converted, rep.deferred, rep.deferred_online_only))
        assert sum(p.is_file() for p in pages) == sum(done for done, _d, _o in seen)
        assert (loop.next_step(sample_config).rule == 3) is (rep.deferred > 0)
    assert seen == [(2, 3, 0), (2, 1, 0), (1, 0, 0)]
    read = [Path(run[0]).name for run in reads(engine.helper)]
    assert sorted(read) == sorted(p.name for p in shots), "every image is read once, none twice"
    assert all(
        page(sample_config.docs_repo, slug.mirror_rel_path(SID, f"shots/{p.name}"))[0]["status"] == "current"
        for p in shots
    )
    # Three more files with bytes already converted: the cache serves them, so no helper time is spent on
    # them and none of them waits, although one read alone would pass the budget.
    monkeypatch.setattr(cycle_mod, "_OCR_BUDGET_S", 50.0)
    for n, shot in enumerate(shots[:3]):
        shutil.copyfile(shot, local_source_dir / "shots" / f"copy {n}.png")
    [rep] = run(sample_config).sources
    assert (rep.converted, rep.deferred) == (3, 0) and len(reads(engine.helper)) == 5


def test_a_scanned_pdf_page_is_read_in_the_staging_folder_and_past_the_ocr_budget_a_pdf_converts_without(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A document does not wait as an image does. Each read here 'takes' 100 s, so the second PDF passes
    the 180 s budget and the third is converted as on a Mac without an engine: the page it would have
    there, under that version, which is what tells a later re-read that OCR has not read it. The next
    cycle, with its own OCR time, is that re-read."""
    assert run(sample_config).exit_code == 0
    engine = shade_engine(tmp_path / "ocr-bin", {90 + n: [f"Delivery note {n}"] for n in range(3)})
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    ticks = itertools.count(0.0, 100.0)
    monkeypatch.setattr(cycle_mod, "_ocr_clock", lambda: next(ticks))
    (local_source_dir / "scans").mkdir()
    for n in range(3):
        pages = [(["The cover page has a text layer of its own"], []), ([], [page_picture(90 + n)])]
        build_picture_pdf(local_source_dir / "scans" / f"scan {n}.pdf", pages)
    [rep] = run(sample_config).sources
    assert (rep.converted, rep.deferred, rep.errors) == (3, 0, ())
    read, unread = [], []
    for n in range(3):
        fm, body = page(sample_config.docs_repo, slug.mirror_rel_path(SID, f"scans/scan {n}.pdf"))
        assert fm["status"] == "current" and body.startswith("> [UNTRUSTED CONTENT]")
        if "+ocr-paper-vision-" in fm["converter"]:
            assert fm["summary"] == "PDF: 2 page(s), 1 read by on-device OCR"
            assert body.rstrip().endswith(f"(Apple Vision)]\n\nDelivery note {n}")
            read.append(n)
        else:
            assert fm["summary"] == "PDF: 2 page(s), 1 without a text layer (scanned; OCR not run)"
            assert body.rstrip().endswith("[scanned page: no text layer]") and "Delivery" not in body
            unread.append(n)
    assert (len(read), len(unread)) == (2, 1), "two reads pass the budget; the third file does not wait"
    staging = sample_config.state_paths.staging.resolve()
    runs = calls(engine.helper)
    assert len(runs) == 2 and all(Path(c["cwd"]).parent.parent == staging for c in runs)
    assert all(Path(c["cwd"]).name.startswith(".ocr-") for c in runs) and list(staging.iterdir()) == []
    assert loop.next_step(sample_config).rule != 3, "nothing waits on this Mac"
    again = run(sample_config)
    (late,) = unread
    fm, body = page(sample_config.docs_repo, slug.mirror_rel_path(SID, f"scans/scan {late}.pdf"))
    assert again.commit_sha is not None and "+ocr-paper-vision-" in fm["converter"]
    assert body.rstrip().endswith(f"(Apple Vision)]\n\nDelivery note {late}")
    assert again.sources[0].converted == 0, "read again: the same bytes are no new conversion"
    assert len(calls(engine.helper)) == 3 and all(
        Path(c["cwd"]).parent.parent == staging for c in calls(engine.helper)
    )
    settled = run(sample_config)
    assert settled.commit_sha is None and len(calls(engine.helper)) == 3, "no file is read a third time"
    assert loop.next_step(sample_config).rule != 3


def test_a_deck_and_a_word_document_are_read_like_a_pdf_and_an_rtf_file_is_as_it_was(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The pictures of a deck and of a Word document go through the same registry as a PDF's: read under
    the staging folder, and past the OCR budget converted as on a Mac without an engine.  An .rtf file has
    one converter with or without an engine, so nothing about its page changes."""
    assert run(sample_config).exit_code == 0
    engine = _use_ocr(monkeypatch, tmp_path)
    folder, markdown = local_source_dir / "launch", "# Minutes\n\n![shot](shot.png)\n"
    folder.mkdir()
    (tmp_path / "shot.png").write_bytes(text_png("Orders by month"))
    for name in ("minutes.docx", "minutes.rtf"):
        pandoc_build(markdown, "markdown", folder / name, cwd=tmp_path)
    prs = Presentation()
    shapes = prs.slides.add_slide(prs.slide_layouts[6]).shapes
    shapes.add_picture(io.BytesIO(text_png("Units by region")), Inches(1), Inches(1))
    prs.save(str(folder / "review.pptx"))
    [rep] = run(sample_config).sources
    assert (rep.converted, rep.deferred, rep.errors) == (3, 0, ())

    def published(name: str) -> tuple[dict[str, Any], str]:
        fm, body = page(sample_config.docs_repo, slug.mirror_rel_path(SID, f"launch/{name}"))
        assert fm["status"] == "current" and body.startswith("> [UNTRUSTED CONTENT]")
        return fm, body.rstrip()

    fm, body = published("minutes.docx")
    assert fm["converter"].startswith("pandoc-gfm@1.0.0+pandoc-") and "+ocr-paper-vision-" in fm["converter"]
    assert fm["summary"] == "Word document; text of 1 picture(s) read by on-device OCR; headings: Minutes"
    assert body.endswith("read by on-device OCR (Apple Vision):]\n\nOrders by month")
    fm, body = published("review.pptx")
    assert fm["converter"].startswith("pptx-python-pptx@1.0.0+") and "+ocr-paper-vision-" in fm["converter"]
    assert body.endswith("read by on-device OCR (Apple Vision):]\nUnits by region")
    rtf, rtf_body = published("minutes.rtf")
    assert rtf["converter"].startswith("pandoc-gfm@1.0.0+pandoc-") and "ocr" not in rtf["converter"]
    assert "Orders by month" not in rtf_body
    staging = sample_config.state_paths.staging.resolve()
    runs = calls(engine.helper)
    assert len(runs) == 2 and all(Path(c["cwd"]).parent.parent == staging for c in runs)
    assert all(Path(c["cwd"]).name.startswith(".ocr-") for c in runs) and list(staging.iterdir()) == []
    assert run(sample_config).commit_sha is None and len(calls(engine.helper)) == 2
    # The cycle's OCR time is used up: a new deck and a new document do not wait, and are not read.
    monkeypatch.setattr(cycle_mod, "_OCR_BUDGET_S", 0.0)
    (tmp_path / "shot.png").write_bytes(text_png("Returns by month"))
    pandoc_build(markdown, "markdown", folder / "late.docx", cwd=tmp_path)
    pandoc_build(markdown, "markdown", folder / "late.rtf", cwd=tmp_path)
    shapes.add_picture(io.BytesIO(text_png("Returns by region")), Inches(1), Inches(3))
    prs.save(str(folder / "late.pptx"))
    [rep] = run(sample_config).sources
    assert (rep.converted, rep.deferred, rep.errors) == (3, 0, ()) and len(calls(engine.helper)) == 2
    for name in ("late.docx", "late.pptx"):
        fm, body = published(name)
        assert "ocr" not in fm["converter"] and "OCR" not in fm["summary"] + body, name
    assert published("late.rtf")[0]["converter"] == rtf["converter"]
    assert loop.next_step(sample_config).rule != 3, "nothing waits on this Mac"


def test_an_online_only_image_is_not_downloaded_for_ocr(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """No image was ever downloaded, and OCR does not start it: while reading an image means a download it
    keeps the stub it has without an engine, whatever the byte budget and even when ``materialise`` names
    it. An image already on this Mac is read."""
    engine = _use_ocr(monkeypatch, tmp_path)
    local = picture(local_source_dir / SITE_PLAN, "Loading dock")
    photo = picture(local_source_dir / "projects" / "Contoso Offsite Photo.heic", "Welcome")
    rel = "projects/Contoso Offsite Photo.heic"
    online = {photo.stat().st_ino}
    real = materialise.is_dataless

    def is_dataless(st: Any) -> bool:
        return st.st_ino in online or real(st)

    monkeypatch.setattr(materialise, "is_dataless", is_dataless)
    monkeypatch.setattr(al, "is_dataless", is_dataless)
    fetched = _fetches(monkeypatch)
    [rep] = run(sample_config).sources
    assert SITE_PLAN in fetched and rel not in fetched
    assert rep.materialised_bytes == 0 and (rep.deferred, rep.deferred_online_only) == (0, 0)
    fm, _body = _image_page(sample_config, rel)
    assert (fm["status"], fm["reason"], fm["converter"]) == ("refused", "no converter for .heic", "none@0")
    assert _image_page(sample_config, SITE_PLAN)[0]["status"] == "current"
    assert [Path(run[0]).name for run in reads(engine.helper)] == [local.name]
    assert loop.next_step(sample_config).rule != 3, "nothing waits for a download budget"
    run(sample_config)
    run(sample_config, materialise_paths=[photo], budget_bytes=10_000_000)
    assert rel not in fetched and len(reads(engine.helper)) == 1
    assert _image_page(sample_config, rel)[0]["status"] == "refused"
    # The person downloads it (Finder's Download Now). That changes the inode: its flags, so its change
    # time. The walk sees the row moved, and the file, now on this Mac, is read.
    online.clear()
    photo.chmod(0o600)
    assert run(sample_config).exit_code == 0 and fetched.count(rel) == 1
    fm, body = _image_page(sample_config, rel)
    assert fm["status"] == "current" and body.rstrip().endswith("Welcome")


def test_a_graph_image_is_not_downloaded_for_ocr(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    engine = _use_ocr(monkeypatch, tmp_path)
    config = config_with(tmp_path, local_source_dir, GRAPH_SOURCE)
    drive = FakeDrive(
        {"I1": ("notes.md", b"# Notes\n"), "I2": ("Contoso Site Plan.png", picture_bytes("Loading dock"))}
    )
    with GraphClient(
        FakeTokens(),
        user_agent="NONISV|test|agentsync/0",
        transport=httpx.MockTransport(drive.handler),
        sleep=lambda _s: None,
    ) as client:
        assert run(config, client=client, only=["drive"]).exit_code == 0
    assert any(entry.endswith("/items/I1/content") for entry in drive.log)
    assert not any("I2" in entry for entry in drive.log), drive.log
    fm, _body = page(config.docs_repo, "mirror/drive/projects/contoso-site-plan.png.md")
    assert (fm["status"], fm["reason"]) == ("refused", "no converter for .png") and reads(engine.helper) == []


def test_a_graph_document_waits_for_ocr_time_and_is_not_converted_without_it(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A file on this Mac that is converted past the cycle's OCR time is read again later.  Nothing reads a
    Graph file again (a re-read never downloads), so converted without OCR it would keep that page until its
    bytes changed.  A Graph document whose converter reads with the engine is left for the next cycle
    before it is downloaded; a file OCR has nothing to do with is not held up."""
    engine = shade_engine(tmp_path / "ocr-bin", {90: ["Delivery note 7"]})
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    config = config_with(tmp_path, local_source_dir, GRAPH_SOURCE)
    pages = [(["The cover page has a text layer of its own"], []), ([], [page_picture(90)])]
    scan = build_picture_pdf(tmp_path / "scan.pdf", pages).read_bytes()
    drive = FakeDrive({"I1": ("notes.md", b"# Notes\n"), "I2": ("Contoso delivery note.pdf", scan)})
    with GraphClient(
        FakeTokens(),
        user_agent="NONISV|test|agentsync/0",
        transport=httpx.MockTransport(drive.handler),
        sleep=lambda _s: None,
    ) as client:
        with monkeypatch.context() as spent:
            spent.setattr(cycle_mod, "_OCR_BUDGET_S", 0.0)  # an earlier source used the cycle's OCR time
            first = run(config, client=client, only=["drive"])
        rep = source_report(first, "drive")
        assert first.exit_code == 0 and (rep.deferred, rep.deferred_online_only, rep.errors) == (1, 0, ())
        assert any(entry.endswith("/items/I1/content") for entry in drive.log)
        assert not any(entry.endswith("/items/I2/content") for entry in drive.log), "not downloaded to wait"
        assert not list((config.docs_repo / "mirror" / "drive").rglob("*.pdf.md"))
        second = run(config, client=client, only=["drive"])
        assert second.exit_code == 0 and source_report(second, "drive").deferred == 0
        assert sum(entry.endswith("/items/I2/content") for entry in drive.log) == 1
    [mirrored] = (config.docs_repo / "mirror" / "drive").rglob("*.pdf.md")
    text = mirrored.read_text(encoding="utf-8")
    assert "+ocr-paper-vision-" in text and text.rstrip().endswith("(Apple Vision)]\n\nDelivery note 7")


def test_an_image_read_by_ocr_keeps_its_page_once_it_is_online_only_whatever_queues_it(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An image is read while it is on this Mac, and the provider evicts it later.  It is not downloaded
    for OCR, and that used to mean its row got the ``no converter`` stub as soon as anything made it
    pending work: a ``materialise PATH`` run that names it committed the stub over the page.  The page is
    kept, the row is settled as the online-only file it is, and the run says why nothing was read."""
    engine = _use_ocr(monkeypatch, tmp_path)
    photo = picture(local_source_dir / SITE_PLAN, "Loading dock", "North gate: closed")
    assert (
        run(sample_config).exit_code == 0 and _image_page(sample_config, SITE_PLAN)[0]["status"] == "current"
    )
    online, real = {photo.stat().st_ino}, materialise.is_dataless

    def is_dataless(st: Any) -> bool:
        return st.st_ino in online or real(st)

    monkeypatch.setattr(materialise, "is_dataless", is_dataless)
    monkeypatch.setattr(al, "is_dataless", is_dataless)
    photo.chmod(0o600)  # an eviction changes the file's flags, so its change time: the walk sees the row move
    assert run(sample_config).exit_code == 0
    before = _mirror_bytes(sample_config, SITE_PLAN)
    assert b"North gate: closed" in before[SITE_PLAN]
    fetched = _fetches(monkeypatch)
    named = run(sample_config, materialise_paths=[photo], budget_bytes=10_000_000)
    [rep] = named.sources
    assert (named.exit_code, named.commit_sha, fetched) == (0, None, [])
    assert rep.alarms == (
        f"{SITE_PLAN}: an online-only image is not downloaded for OCR; its page is kept as it was",
    )
    assert _mirror_bytes(sample_config, SITE_PLAN) == before and len(reads(engine.helper)) == 1
    row = _file_rows(sample_config)[SITE_PLAN]
    assert (row.state, row.last_verdict) == (RowState.DATALESS, Verdict.DATALESS)
    assert run(sample_config).commit_sha is None and loop.next_step(sample_config).rule != 3
    assert _mirror_bytes(sample_config, SITE_PLAN) == before


LABEL_RULE = '\n[policy]\nexclude_label_ids = ["00000000-0000-4000-8000-00000000c0de"]\n'


def test_under_a_label_rule_no_image_is_read_and_a_page_from_before_it_becomes_a_stub(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An image can carry a sensitivity label nothing here reads, so a label rule fails closed: no image
    converter. The [policy] re-screen covers images, so the page published before the rule does not stay."""
    engine = _use_ocr(monkeypatch, tmp_path)
    picture(local_source_dir / SITE_PLAN, "Loading dock")
    open_config = config_with(tmp_path, local_source_dir)
    assert run(open_config).exit_code == 0
    assert _image_page(open_config, SITE_PLAN)[0]["status"] == "current" and len(reads(engine.helper)) == 1
    labelled = config_with(tmp_path, local_source_dir, LABEL_RULE)
    picture(local_source_dir / "projects" / "Contoso Badge Scan.jpg", "Visitor")
    fetched = _fetches(monkeypatch)
    report = run(labelled)
    assert report.exit_code == 0 and not [rel for rel in fetched if rel.endswith((".png", ".jpg"))]
    for rel, suffix in ((SITE_PLAN, ".png"), ("projects/Contoso Badge Scan.jpg", ".jpg")):
        fm, body = _image_page(labelled, rel)
        assert (fm["status"], fm["reason"]) == ("refused", f"no converter for {suffix}")
        assert "Loading dock" not in body and "Visitor" not in body
    assert len(reads(engine.helper)) == 1, "the helper is not run under a label rule"
    assert governance.pending_purges(labelled.state_paths.root) == [], "no label was read: nothing to purge"
    with Manifest(labelled.state_paths.db) as m:
        assert m.get_meta(cycle_mod._RESCREEN_META) == ""


def test_a_helper_that_fails_on_everything_costs_no_image_its_reading(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The helper answers ``--version`` and then fails every read: its folder was made writable by another
    tool, it was removed in mid-cycle, it crashes on Vision.  That is no image's failure.  After the first
    failed read the helper is handed a blank image, fails on it too, and is not run again in that cycle.
    Every image gets the ``no converter`` stub it has without an engine (never a failed conversion, which
    the next cycle would settle for good), the report says OCR stopped working, and no file's re-read is
    used up.  Once the helper works, each image is read."""
    failing = fake_engine(tmp_path / "failing", fail=True)
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: failing)
    plans = [f"projects/Contoso Site Plan {n}.png" for n in range(3)]
    for n, rel in enumerate(plans):
        picture(local_source_dir / rel, f"Loading dock {n}")
    [rep] = run(sample_config).sources
    for rel in plans:
        fm, body = _image_page(sample_config, rel)
        assert (fm["status"], fm["reason"], fm["converter"]) == ("refused", "no converter for .png", "none@0")
        assert "the fake helper was told to fail" not in body + str(fm), (
            "what the helper said stays in the log"
        )
    assert (rep.errors, rep.alarms) == ((), (cycle_mod._OCR_DOWN,))
    assert len(calls(failing.helper)) == 2, "one image, then the blank image: no run for the other two"
    assert list(sample_config.state_paths.staging.iterdir()) == [], "the blank image is removed"
    assert _reread_record(sample_config) == (False, []) and loop.next_step(sample_config).rule != 3
    again = run(sample_config)
    assert again.commit_sha is None and again.sources[0].errors == ()
    assert again.sources[0].alarms == (cycle_mod._OCR_DOWN,) and len(calls(failing.helper)) == 4
    assert _reread_record(sample_config) == (False, []), "the failed read was not the file's to lose"
    assert all(_file_rows(sample_config)[rel].state is RowState.REFUSED for rel in plans)
    healthy = _use_ocr(monkeypatch, tmp_path)
    fixed = run(sample_config)
    assert fixed.commit_sha is not None and fixed.sources[0].alarms == () and len(reads(healthy.helper)) == 3
    for n, rel in enumerate(plans):
        fm, body = _image_page(sample_config, rel)
        assert fm["status"] == "current" and body.rstrip().endswith(f"Loading dock {n}")
    assert _reread_record(sample_config) == (True, [])
    assert run(sample_config).commit_sha is None and len(reads(healthy.helper)) == 3


def _run_record(config: Config) -> dict[str, int]:
    """The newest run's ``runs.counts_json``: its change counts and what ``_Cycle._run_tally`` adds."""
    with Manifest(config.state_paths.db) as m:
        (raw,) = m._db.execute("SELECT counts_json FROM runs ORDER BY run_id DESC LIMIT 1").fetchone()
    record: dict[str, int] = json.loads(raw)
    return record


LOOKED = {"reread_for": ANY}
"""In the record of a run that brought a source's re-read record up to date: what it looked for, a number
of its own (``test_a_run_record_says_what_its_re_read_looked_for``)."""


def test_a_run_record_says_what_its_re_read_looked_for(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A source's re-read record says what it was written for, and a cycle goes by it only when that is
    what it looks for.  The setup report cannot work that out (it builds no registry), so each run that
    brings a record up to date says it in its run record: the front of the same digest, as an integer.
    A new number once there is an engine, and none from a ``materialise PATH`` run, which reads no file
    again and leaves the records as they are."""

    def stored() -> int:
        with Manifest(sample_config.state_paths.db) as m:
            raw = m.get_meta(cycle_mod._REREAD_META + SID)
        assert raw is not None
        return int(json.loads(raw)[0]["for"][: cycle_mod._REREAD_FOR_DIGITS], 16)

    assert run(sample_config).exit_code == 0
    plain = _run_record(sample_config)["reread_for"]
    assert plain == stored() and isinstance(plain, int)
    assert run(sample_config).commit_sha is None and _run_record(sample_config)["reread_for"] == plain
    _use_ocr(monkeypatch, tmp_path)
    assert run(sample_config).exit_code == 0
    with_engine = _run_record(sample_config)["reread_for"]
    assert with_engine == stored() and with_engine != plain, "an engine is something else to look for"
    named = run(
        sample_config,
        materialise_paths=[local_source_dir / "projects" / "sample.md"],
        budget_bytes=10_000_000,
    )
    assert named.exit_code == 0 and "reread_for" not in _run_record(sample_config)
    assert stored() == with_engine


def test_each_run_records_its_ocr_time_and_the_images_that_waited_for_it(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The setup report reads OCR's time from the run record (CONTRACTS.md 16.28): milliseconds used, the
    budget, whether it was used up and how many images it left for a later cycle.  Integers only: a
    record never holds a name."""
    assert run(sample_config).exit_code == 0
    before = _run_record(sample_config)
    assert before["converted"] > 0 and not [key for key in before if key.startswith("ocr_")], (
        "a cycle without an engine records no OCR time"
    )
    _use_ocr(monkeypatch, tmp_path)
    ticks = itertools.count(0.0, 100.0)
    monkeypatch.setattr(cycle_mod, "_ocr_clock", lambda: next(ticks))
    for n in range(5):
        picture(local_source_dir / "shots" / f"Contoso shot {n}.png", f"screenshot {n}")
    assert run(sample_config).exit_code == 0
    record = _run_record(sample_config)
    assert {key: record[key] for key in record if key.startswith(("ocr_", "converted"))} == {
        "converted": 2,
        "ocr_budget_s": 180,
        "ocr_deferred": 3,
        "ocr_ms": 200_000,
        "ocr_over": 1,
    }
    assert all(isinstance(value, int) for value in record.values()) and "shot" not in json.dumps(record)
    # Images left for a later cycle's OCR are files on this Mac not converted yet: the NEXT line itself
    # says to sync again (rule 3), so a loop on NEXT reads them all before it goes on.
    step = loop.next_step(sample_config)
    assert (step.rule, step.step) == (
        3,
        f"1 source(s) not fully listed or converted yet ({SID}): run `{loop.BIN} sync` again",
    )
    assert run(sample_config).exit_code == 0 and run(sample_config).exit_code == 0
    last = _run_record(sample_config)
    assert (last["converted"], last["ocr_ms"], last.get("ocr_deferred"), last.get("ocr_over")) == (
        1,
        100_000,
        None,
        None,
    )
    assert loop.next_step(sample_config).rule != 3, "none waits any more"
    assert run(sample_config).commit_sha is None
    assert _run_record(sample_config) == {"converted": 0, "ocr_budget_s": 180, "ocr_ms": 0, **LOOKED}, (
        "an idle cycle with an engine"
    )


def test_the_run_record_counts_what_was_converted_without_ocr_and_why(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A document converted past the cycle's OCR time, the re-read that gives it its text, a scan longer
    than the page limit, and a helper that stops working: each is a count in the run record."""
    assert run(sample_config).exit_code == 0
    engine = shade_engine(tmp_path / "ocr-bin", {90 + n: [f"Delivery note {n}"] for n in range(4)})
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    monkeypatch.setattr(pdf_mod, "MAX_PAGES", 1)
    (local_source_dir / "scans").mkdir()
    cover = (["The cover page has a text layer of its own"], [])
    build_picture_pdf(
        local_source_dir / "scans" / "Contoso long scan.pdf",
        [cover, ([], [page_picture(90)]), ([], [page_picture(91)])],
    )
    assert run(sample_config).exit_code == 0
    record = _run_record(sample_config)
    assert (record["converted"], record["ocr_page_cap"]) == (1, 1), "the second scanned page: past the limit"
    assert "ocr_failed" not in record and "ocr_without_budget" not in record
    monkeypatch.setattr(cycle_mod, "_OCR_BUDGET_S", 0.0)
    build_picture_pdf(local_source_dir / "scans" / "Contoso late scan.pdf", [cover, ([], [page_picture(92)])])
    assert run(sample_config).exit_code == 0
    record = _run_record(sample_config)
    assert (record["converted"], record["ocr_without_budget"], record["ocr_over"]) == (1, 1, 1)
    assert "ocr_page_cap" not in record, "converted without OCR: no limit of OCR was reached"
    assert record["reread_left"] == 1 and "reread" not in record, "it waits for a sync with OCR time"
    assert loop.next_step(sample_config).notes[-1].startswith("sync again: 1 file(s) in ")
    monkeypatch.setattr(cycle_mod, "_OCR_BUDGET_S", 180.0)
    assert run(sample_config).exit_code == 0
    record = _run_record(sample_config)
    assert (record["reread"], record["converted"]) == (1, 0) and "reread_kept" not in record
    failing = fake_engine(tmp_path / "failing", fail=True)
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: failing)
    for n in range(3):
        picture(local_source_dir / "projects" / f"Contoso Site Plan {n}.png", f"Loading dock {n}")
    assert run(sample_config).exit_code == 0
    record = _run_record(sample_config)
    assert (record["ocr_failed"], record["ocr_without_down"], record["ocr_down"]) == (1, 2, 1)
    assert "Contoso" not in json.dumps(record) and "scan" not in json.dumps(record)


def test_the_run_record_counts_a_file_converted_again_from_the_same_bytes(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Whether the same files are converted run after run is a question the setup report answers from the
    run record: ``converted_again`` is a file converted from the bytes its own pages were made from, when
    the run just before had converted them too.  The loop here is a file the walk calls maybe-changed
    whose pages are never found intact."""
    assert run(sample_config).exit_code == 0
    first = _run_record(sample_config)
    assert first["converted"] > 1 and "converted_again" not in first
    assert "converted_seen" not in first, "a first run has seen no bytes before"
    assert run(sample_config).commit_sha is None
    assert _run_record(sample_config) == {"converted": 0, **LOOKED}
    monkeypatch.setattr(cycle_mod, "_pages_intact", lambda _repo, _outs: False)
    source = local_source_dir / "projects" / "sample.md"

    def touch() -> None:
        stamp = source.stat().st_mtime_ns + 5_000_000_000
        os.utime(source, ns=(stamp, stamp))

    touch()
    assert run(sample_config).exit_code == 0
    record = _run_record(sample_config)
    assert (record["converted"], record["converted_seen"]) == (1, 1)
    assert "converted_again" not in record, "the run just before converted nothing"
    touch()
    assert run(sample_config).exit_code == 0
    record = _run_record(sample_config)
    assert (record["converted"], record["converted_seen"], record["converted_again"]) == (1, 1, 1)


def test_a_copy_of_a_file_the_run_before_converted_is_no_file_converted_again(
    sample_config: Config, local_source_dir: Path
) -> None:
    """The cache row the run record went by is of the bytes and the converter, not of the file.  A second
    file with the bytes of one the run before had converted (a copy, a re-export, one attachment saved
    twice) was counted as converted again, which the setup report calls the sign of a loop.  A repeat is
    a file whose own pages were made from those bytes."""
    assert run(sample_config).exit_code == 0
    source = local_source_dir / "projects" / "sample.md"
    shutil.copyfile(source, source.with_name("Contoso sample copy.md"))
    assert run(sample_config).commit_sha is not None
    record = _run_record(sample_config)
    assert record["converted"] == 1 and not [key for key in record if key.startswith("converted_")]
    with Manifest(sample_config.state_paths.db) as m:
        twins = m._db.execute(
            "SELECT COUNT(*), COUNT(DISTINCT o.action_key) FROM outputs o JOIN items i "
            "ON i.source_id = o.source_id AND i.stable_id = o.stable_id WHERE i.name LIKE '%sample%.md'"
        ).fetchone()
    assert tuple(twins) == (2, 1), "the two files do share one action key"
    # A third file with those bytes in the run after: the key was used by the run just before, and it is
    # still no repeat.
    shutil.copyfile(source, source.with_name("Contoso sample copy 2.md"))
    assert run(sample_config).commit_sha is not None
    record = _run_record(sample_config)
    assert record["converted"] == 1 and not [key for key in record if key.startswith("converted_")]


def test_a_conversion_that_failed_is_counted_once_and_not_tried_again(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``converted_failed`` counts a conversion in the run it failed in.  The file then has its stub, and
    no later run converts it until its bytes change: a run record with failures followed by one with none
    is not a retry that worked."""
    from agentsync.convert.markdown import MarkdownConverter  # noqa: PLC0415

    real = MarkdownConverter.convert
    tries: list[str] = []

    def failing(self: Any, src: Path, *, name: str) -> Any:
        if name.startswith("Contoso broken"):
            tries.append(name)
            raise RuntimeError("no such heading")
        return real(self, src, name=name)

    monkeypatch.setattr(MarkdownConverter, "convert", failing)
    broken = local_source_dir / "projects" / "Contoso broken.md"
    broken.write_text("# Notes\n\nA page the converter fails on.\n", encoding="utf-8")
    assert run(sample_config).exit_code == 0
    first = _run_record(sample_config)
    assert first["converted_failed"] == 1 and first["converted"] > 1 and len(tries) == 1
    for _ in range(2):
        assert run(sample_config).commit_sha is None
        assert _run_record(sample_config) == {"converted": 0, **LOOKED}
    assert len(tries) == 1, "the same bytes are not converted again"
    with Manifest(sample_config.state_paths.db) as m:
        row = m.item_by_path(SID, "projects/Contoso broken.md")
    assert row is not None
    assert row.state is RowState.QUARANTINED and (row.state_reason or "").startswith("conversion failed")
    broken.write_text("# Notes\n\nNew bytes, and the converter still fails.\n", encoding="utf-8")
    assert run(sample_config).exit_code == 0
    record = _run_record(sample_config)
    assert (record["converted"], record["converted_failed"], len(tries)) == (1, 1, 2)
    assert "converted_seen" not in record


def test_the_limit_marks_are_the_converters_own_wording() -> None:
    assert cycle_mod._PAGE_CAP_MARK in pdf_mod._SCANNED_OUTCOMES[pdf_mod._OCR_OVER_LIMIT]
    assert cycle_mod._PAGE_CAP_MARK in pdf_mod._NO_TEXT_PAST_LIMIT
    assert cycle_mod._PICTURE_CAP_MARK in image_mod._PICTURES_CUT


def test_an_image_without_text_is_a_settled_stub_outside_the_curation_queue(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A page per logo would be curation work for ever. The stub is cached and settled: nothing reads the
    file again, and the curation queue, which lists every current page no topic cites, does not hold it."""
    engine = _use_ocr(monkeypatch, tmp_path)
    logo = "projects/Contoso logo.png"
    picture(local_source_dir / logo)
    picture(local_source_dir / SITE_PLAN, "Loading dock")
    assert run(sample_config).exit_code == 0
    fm, body = _image_page(sample_config, logo)
    assert (fm["status"], fm["reason"]) == ("unreadable", "no text found in the image by on-device OCR")
    assert "[image" not in body
    uncovered = curate.uncovered_mirror_pages(DocsLayout(root=sample_config.docs_repo), [])
    assert slug.mirror_rel_path(SID, SITE_PLAN) in uncovered
    assert slug.mirror_rel_path(SID, logo) not in uncovered
    assert _file_rows(sample_config)[logo].last_verdict is Verdict.QUARANTINED
    fetched = _fetches(monkeypatch)
    assert run(sample_config).commit_sha is None and fetched == [] and len(reads(engine.helper)) == 2
    assert loop.next_step(sample_config).rule != 3


def test_a_credential_read_from_an_image_is_quarantined_like_any_other_text(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The secret scan reads every page a cycle writes, and what OCR read is in the page (the summary, which
    is front matter, holds none of it): a screenshot of a key is a ``contains a credential`` stub."""
    _use_ocr(monkeypatch, tmp_path)
    key = "AKIA" + "ABCDEFGHIJKLMNOP"
    console = "projects/Contoso console.png"
    picture(local_source_dir / console, "Access keys", f"aws key {key}")
    report = run(sample_config)
    assert report.exit_code == 0 and report.commit_sha is not None
    fm, body = _image_page(sample_config, console)
    assert fm["reason"] == "contains a credential" and key not in body + str(fm)
    assert not committed_blobs_containing(sample_config.docs_repo, key.encode())


# ---------------------------------------------------------------------------------------------------------
# re-read once: a file converted before something its converter has is read again, once (CONTRACTS 16.27)
# ---------------------------------------------------------------------------------------------------------

COMMENTS = "[comments on this page (PDF annotations):]"
REVIEW = "projects/Contoso widget review.pdf"
PLAIN_PDF = "projects/sample.pdf"


def _commented(path: Path, title: str) -> Path:
    """A one-page PDF with one reviewer's note on it."""
    path.parent.mkdir(parents=True, exist_ok=True)
    note = f"/Subtype /Text /Rect [400 700 420 720] /T (Roe, John) /Contents (Check {title})"
    return build_annotated_pdf(path, [([title, "Totals by region"], [note])])


@contextlib.contextmanager
def _before_comments(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """The PDF emitter from before comments were kept: it says 2.0.0 and writes no comment."""
    with monkeypatch.context() as old:
        old.setattr(pdf_mod, "_EMITTER_VERSION", "2.0.0")
        old.setattr(pdf_mod, "_render_comments", lambda _found: [])
        yield


def _shade_png(path: Path, shade: int) -> Path:
    """A real gray PNG, every pixel ``shade``: an image file the ``shade_engine`` helper tells by it."""

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    path.parent.mkdir(parents=True, exist_ok=True)
    head = chunk(b"IHDR", struct.pack(">IIBBBBB", 96, 64, 8, 0, 0, 0, 0))
    pixels = chunk(b"IDAT", zlib.compress((b"\x00" + bytes([shade]) * 96) * 64))
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + head + pixels + chunk(b"IEND", b""))
    return path


def _reread_record(config: Config, sid: str = SID) -> tuple[bool, list[str]]:
    """(done, tried) of the re-read record the last cycle wrote for ``sid``."""
    with Manifest(config.state_paths.db) as m:
        raw = m.get_meta(cycle_mod._REREAD_META + sid)
    assert raw is not None
    newest = json.loads(raw)[0]
    return newest["done"], newest["tried"]


def _reread_notes(config: Config) -> list[str]:
    """The loop's note lines about files still to read again (``loop._reread_note``)."""
    return [note for note in loop.next_step(config).notes if "read again" in note]


def _sync_again(left: int, sid: str = SID) -> str:
    return (
        f"sync again: {left} file(s) in {sid} are still to be read again, once, for what this build's "
        "converters have gained (each sync reads about 2 minutes' worth); it does not block the next step"
    )


def _reread_failed(config: Config, sid: str = SID) -> dict[str, int]:
    """Stable id -> the cycles its re-read failed in, for the files not yet given up."""
    with Manifest(config.state_paths.db) as m:
        raw = m.get_meta(cycle_mod._REREAD_META + sid)
    assert raw is not None
    return dict(json.loads(raw)[0].get("failed", {}))


def _mirror_page(config: Config, rel: str) -> tuple[dict[str, Any], str]:
    return page(config.docs_repo, slug.mirror_rel_path(SID, rel))


def _mirror_bytes(config: Config, *rels: str) -> dict[str, bytes]:
    return {rel: (config.docs_repo / slug.mirror_rel_path(SID, rel)).read_bytes() for rel in rels}


def test_a_pdf_converted_before_comments_were_kept_gains_them_and_is_read_again_once(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A PDF mirrored by emitter 2.0.0 gets its reviewers' comments without its bytes changing: the next
    cycle reads it again, once.  A PDF with no comment is read too, and its page stays as it is, byte for
    byte, its 2.0.0 line included.  After that no cycle looks: none of the three queries is run."""
    _commented(local_source_dir / REVIEW, "Contoso widget review")
    fetched = _fetches(monkeypatch)
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
        fm, body = _mirror_page(sample_config, REVIEW)
        assert fm["converter"].startswith("pdf-pypdfium2@2.0.0+") and COMMENTS not in body
        del fetched[:]
        # The floor is above the emitter that runs here: what a re-read wrote would be below it too, so
        # nothing is outdated and no file is read every cycle.
        assert _reread_record(sample_config) == (True, [])
        assert run(sample_config).commit_sha is None and fetched == []
    plain = _mirror_bytes(sample_config, PLAIN_PDF)
    assert run(sample_config, mode=CycleMode.DRY_RUN).exit_code == 0 and fetched == [], (
        "a dry run reads nothing"
    )
    second = run(sample_config)
    [rep] = second.sources
    assert second.commit_sha is not None and sorted(fetched) == [REVIEW, PLAIN_PDF]
    assert (rep.converted, rep.deferred, rep.errors) == (0, 0, ()), "the same bytes are no new conversion"
    assert [c.path for c in second.changes] == [slug.mirror_rel_path(SID, REVIEW)]
    fm, body = _mirror_page(sample_config, REVIEW)
    assert fm["converter"].startswith("pdf-pypdfium2@2.1.0+") and fm["summary"].endswith(
        "1 comment(s) on 1 page(s)"
    )
    assert body.rstrip().endswith(f"{COMMENTS}\n- Note by Roe, John: Check Contoso widget review")
    assert _mirror_bytes(sample_config, PLAIN_PDF) == plain and b"pdf-pypdfium2@2.0.0+" in plain[PLAIN_PDF]
    assert _reread_record(sample_config) == (True, [])
    third = run(sample_config)
    assert third.commit_sha is None and len(fetched) == 2, "no file is read a second time"
    for query in ("produced_by", "reread_candidates", "reread_left", "reread_count"):
        monkeypatch.setattr(Manifest, query, crash)
    assert run(sample_config).commit_sha is None, "nothing is left, and no cycle looks"


def test_files_mirrored_before_there_was_an_engine_are_read_by_it_once(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The files OCR exists for were mirrored before it: an image refused as ``no converter``, a scan
    quarantined as ``no text layer``, a PDF whose scanned page was a marker.  With an engine each is read
    again, once.  A page the engine adds nothing to stays as it is, byte for byte.  The stub of a scan it
    reads nothing in no longer says OCR was not run."""
    scans = local_source_dir / "scans"
    scans.mkdir()
    note, agreement, blank = (
        f"scans/Contoso {n}.pdf" for n in ("delivery note", "supply agreement", "blank scan")
    )
    cover = ["The cover page has a text layer of its own"]
    _shade_png(local_source_dir / SITE_PLAN, 70)
    build_picture_pdf(local_source_dir / note, [([], [page_picture(90)])])
    build_picture_pdf(local_source_dir / agreement, [(cover, []), ([], [page_picture(91)])])
    build_picture_pdf(local_source_dir / blank, [([], [page_picture(92)])])
    assert run(sample_config).exit_code == 0
    assert _mirror_page(sample_config, SITE_PLAN)[0]["reason"] == "no converter for .png"
    for scan in (note, blank):
        fm, _body = _mirror_page(sample_config, scan)
        assert (fm["status"], fm["reason"]) == ("unreadable", pdf_mod._NO_TEXT)
    fm, body = _mirror_page(sample_config, agreement)
    assert fm["status"] == "current" and "ocr" not in fm["converter"]
    assert body.rstrip().endswith("[scanned page: no text layer]")
    assert run(sample_config).commit_sha is None and _reread_record(sample_config) == (True, [])
    documents = [PLAIN_PDF, "projects/sample.pptx", "projects/sample.docx", KICKOFF]
    before = _mirror_bytes(sample_config, *documents)

    said = {70: ["Loading dock"], 90: ["Delivery note 7"], 91: ["Signed in Rotterdam"]}
    engine = shade_engine(tmp_path / "ocr-bin", said)
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    fetched = _fetches(monkeypatch)
    second = run(sample_config)
    [rep] = second.sources
    assert sorted(fetched) == sorted([SITE_PLAN, note, agreement, blank, *documents])
    assert (rep.converted, rep.deferred, rep.errors) == (1, 0, ()), (
        "the image: its bytes were never read before"
    )
    fm, body = _mirror_page(sample_config, SITE_PLAN)
    assert fm["status"] == "current" and body.rstrip().endswith("(Apple Vision)]\n\nLoading dock")
    fm, body = _mirror_page(sample_config, note)
    assert fm["status"] == "current" and "+ocr-paper-vision-" in fm["converter"]
    assert body.rstrip().endswith("(Apple Vision)]\n\nDelivery note 7")
    fm, body = _mirror_page(sample_config, agreement)
    assert "+ocr-paper-vision-" in fm["converter"] and body.rstrip().endswith("Signed in Rotterdam")
    assert _mirror_bytes(sample_config, *documents) == before
    fm, _body = _mirror_page(sample_config, blank)
    assert (fm["status"], fm["reason"]) == ("unreadable", pdf_mod._NO_TEXT_FOUND)
    assert "+ocr-paper-vision-" in fm["converter"]
    assert sorted(c.path for c in second.changes) == sorted(
        slug.mirror_rel_path(SID, rel) for rel in (SITE_PLAN, note, agreement, blank)
    )
    rows = _file_rows(sample_config)
    assert [rows[rel].state for rel in (SITE_PLAN, note, blank)] == [
        RowState.LIVE,
        RowState.LIVE,
        RowState.QUARANTINED,
    ]
    assert _reread_record(sample_config) == (True, []) and loop.next_step(sample_config).rule != 3
    runs = len(calls(engine.helper))
    third = run(sample_config)
    assert third.commit_sha is None and len(fetched) == 8 and len(calls(engine.helper)) == runs


def test_pages_the_field_build_of_ocr_wrote_are_read_again_once(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """One Mac mirrored with a build of OCR that was never on main: image pages from emitter 1.0.0, which
    named the file in the page, and documents under versions ending ``+ocr-off``.  Both read as current.
    This build reads each such file again, once, and what it writes is not read again."""
    engine = _use_ocr(monkeypatch, tmp_path)
    picture(local_source_dir / SITE_PLAN, "Loading dock")
    real = pdf_mod.PdfConverter.version
    with monkeypatch.context() as field:
        field.setattr(image_mod, "_EMITTER_VERSION", "1.0.0")
        field.setattr(pdf_mod.PdfConverter, "version", lambda self: real(self).split("+ocr-")[0] + "+ocr-off")
        assert run(sample_config).exit_code == 0
        assert _image_page(sample_config, SITE_PLAN)[0]["converter"].startswith("image-ocr@1.0.0+ocr-")
        assert _mirror_page(sample_config, PLAIN_PDF)[0]["converter"].endswith("+ocr-off")
    fetched = _fetches(monkeypatch)
    again = run(sample_config)
    assert again.exit_code == 0 and sorted(fetched) == sorted([SITE_PLAN, PLAIN_PDF])
    assert again.sources[0].converted == 0, "the same bytes are no new conversion"
    assert len(reads(engine.helper)) == 2, "the image is read by the helper under this build's version"
    assert _reread_record(sample_config) == (True, [])
    assert run(sample_config).commit_sha is None and len(fetched) == 2


@pytest.mark.parametrize("fault", ["the helper fails on it", "the converter breaks on it"])
def test_an_image_that_cannot_be_read_at_its_re_read_keeps_the_stub_it_has(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fault: str
) -> None:
    """An image mirrored as ``no converter`` was never read, so its re-read is the first look at its bytes.
    When that fails, the stub stays: nothing is published or committed, the source gets no error line
    naming the file, and the read counts against the file, as a document's does: after two cycles of it
    the file is given up.  It used to become a failed conversion, which no later re-read selects."""
    repo = sample_config.docs_repo
    _shade_png(local_source_dir / SITE_PLAN, 70)
    assert run(sample_config).exit_code == 0
    before, head = _mirror_bytes(sample_config, SITE_PLAN), git(repo, "rev-parse", "HEAD")
    assert b"no converter for .png" in before[SITE_PLAN]
    engine = shade_engine(tmp_path / "ocr-bin", {70: "fail" if fault.startswith("the helper") else ["Dock"]})
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    if fault.startswith("the converter"):

        def broken(self: Any, src: Path, *, name: str) -> Any:
            raise RuntimeError(f"cannot decode /Users/someone/Pictures/{name}")

        monkeypatch.setattr(image_mod.ImageConverter, "convert", broken)
    report = run(sample_config)
    [rep] = report.sources
    assert (report.exit_code, report.commit_sha, rep.errors, rep.converted) == (0, None, (), 0)
    kept = "1 file(s) read again for what their converter has gained could not be converted; their pages"
    assert [a for a in rep.alarms if a.startswith(kept)] == ([] if fault.startswith("the helper") else [ANY])
    assert _mirror_bytes(sample_config, SITE_PLAN) == before
    assert git(repo, "rev-parse", "HEAD") == head and porcelain(repo) == ""
    row = _file_rows(sample_config)[SITE_PLAN]
    assert (row.state, row.state_reason) == (RowState.REFUSED, "no converter for .png")
    assert _reread_record(sample_config) == (False, []) and _reread_failed(sample_config) == {
        row.stable_id: 1
    }
    assert run(sample_config).commit_sha is None and _reread_record(sample_config) == (True, [row.stable_id])
    assert _mirror_bytes(sample_config, SITE_PLAN) == before
    fetched = _fetches(monkeypatch)
    assert run(sample_config).commit_sha is None and fetched == []


def test_under_a_label_rule_a_refused_image_is_not_read_and_without_the_rule_it_is(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An image can carry a label nothing here reads, so under a label rule it has no converter, engine or
    not: the stub it got before there was an engine stays, and the file is not read.  When the rule goes,
    the registry has a converter for it again, and the image is read like any file from before it."""
    picture(local_source_dir / SITE_PLAN, "Loading dock")
    assert run(config_with(tmp_path, local_source_dir)).exit_code == 0
    engine = _use_ocr(monkeypatch, tmp_path)
    labelled = config_with(tmp_path, local_source_dir, LABEL_RULE)
    fetched = _fetches(monkeypatch)
    assert run(labelled).exit_code == 0 and run(labelled).exit_code == 0
    assert SITE_PLAN not in fetched and not any(
        SITE_PLAN.endswith(Path(r[0]).name) for r in reads(engine.helper)
    )
    assert _image_page(labelled, SITE_PLAN)[0]["reason"] == "no converter for .png"
    opened = config_with(tmp_path, local_source_dir)
    assert run(opened).exit_code == 0 and fetched.count(SITE_PLAN) == 1
    fm, body = _image_page(opened, SITE_PLAN)
    assert fm["status"] == "current" and body.rstrip().endswith("Loading dock")
    assert run(opened).commit_sha is None and fetched.count(SITE_PLAN) == 1


def test_a_re_read_that_fails_keeps_the_page_and_a_file_is_given_up_after_two_cycles_of_it(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A converter that breaks while files are read again costs no page: nothing is published, nothing is
    committed, and the rows keep their hashes and their verdict.  A fault that passes costs nothing else
    either: the next cycle reads the files again.  A file whose re-read fails in two cycles is given up for
    what this install has, so one that cannot be converted is not read every cycle.  Something new to look
    for starts over."""
    repo = sample_config.docs_repo
    other = "projects/Contoso gadget review.pdf"
    _commented(local_source_dir / REVIEW, "Contoso widget review")
    _commented(local_source_dir / other, "Contoso gadget review")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
    pages = _mirror_bytes(sample_config, REVIEW, other, PLAIN_PDF)
    rows, head = _file_rows(sample_config), git(repo, "rev-parse", "HEAD")
    fetched = _fetches(monkeypatch)
    real = pdf_mod.PdfConverter.convert
    broken_for = {Path(rel).name for rel in pages}

    def convert(self: Any, src: Path, *, name: str) -> Any:
        if name in broken_for:
            raise RuntimeError(f"cannot open /Users/someone/Library/{name}")
        return real(self, src, name=name)

    monkeypatch.setattr(pdf_mod.PdfConverter, "convert", convert)
    report = run(sample_config)
    [rep] = report.sources
    assert (report.exit_code, report.commit_sha, rep.errors) == (0, None, ())
    assert rep.alarms == (
        "3 file(s) read again for what their converter has gained could not be converted; their pages are "
        "kept as they were",
    )
    assert _mirror_bytes(sample_config, *pages) == pages
    assert git(repo, "rev-parse", "HEAD") == head and porcelain(repo) == ""
    after = _file_rows(sample_config)
    for rel in pages:
        was, now = rows[rel], after[rel]
        assert (now.state, now.state_reason, now.last_verdict) == (RowState.LIVE, None, Verdict.UNCHANGED)
        assert (now.content_sha256, now.canonical_sha256) == (was.content_sha256, was.canonical_sha256)
    assert _reread_record(sample_config) == (False, []) and sorted(fetched) == sorted(pages)
    assert _reread_failed(sample_config) == {after[rel].stable_id: 1 for rel in pages}
    # The fault passes for one of them.  The next cycle reads all three again: that one gains its comments,
    # and the two that failed in a second cycle are given up.
    broken_for.discard(Path(REVIEW).name)
    second = run(sample_config)
    assert second.commit_sha is not None and COMMENTS in _mirror_page(sample_config, REVIEW)[1]
    assert len(fetched) == 6 and _mirror_bytes(sample_config, other, PLAIN_PDF) == {
        rel: pages[rel] for rel in (other, PLAIN_PDF)
    }
    given_up = sorted(after[rel].stable_id for rel in (other, PLAIN_PDF))
    assert _reread_record(sample_config) == (True, given_up) and _reread_failed(sample_config) == {}
    # The converter works again.  The two were given up for what this install has: they are left alone.
    broken_for.clear()
    assert run(sample_config).commit_sha is None and len(fetched) == 6
    assert loop.next_step(sample_config).rule != 3
    # A new floor is something new to look for, so the files are read once more.
    monkeypatch.setattr(pdf_mod.PdfConverter, "outdated_key", "2.1.0<2.1.0, and one thing more")
    again = run(sample_config)
    assert again.commit_sha is not None and len(fetched) == 8
    assert COMMENTS in _mirror_page(sample_config, other)[1] and _reread_record(sample_config) == (True, [])
    assert run(sample_config).commit_sha is None and len(fetched) == 8


def test_three_failed_re_reads_in_a_row_end_the_re_reads_of_that_cycle(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """What fails three files running is more likely the Mac than the files (a full disk, a converter that
    cannot start).  The cycle stops there, so a bad hour costs three files one of their attempts and the
    rest of the mirror nothing."""
    for n in range(5):
        _commented(local_source_dir / f"reviews/Contoso review {n}.pdf", f"Contoso review {n}")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0

    def broken(self: Any, src: Path, *, name: str) -> Any:
        raise RuntimeError("no space left on device")

    fetched = _fetches(monkeypatch)
    with monkeypatch.context() as fault:
        fault.setattr(pdf_mod.PdfConverter, "convert", broken)
        [rep] = run(sample_config).sources
    kept = "3 file(s) read again for what their converter has gained could not be converted; their pages"
    assert len(fetched) == 3 and rep.errors == () and rep.alarms == (f"{kept} are kept as they were",)
    assert _reread_record(sample_config) == (False, []) and len(_reread_failed(sample_config)) == 3
    assert run(sample_config).commit_sha is not None and len(fetched) == 9, "the fault passed: all six"
    assert _reread_record(sample_config) == (True, [])


def test_a_re_read_downloads_nothing_reads_nothing_out_of_scope_and_waits_for_materialise(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Three PDFs from before comments were kept.  One is online-only now, one is under a folder the config
    has since excluded, one is here.  Only the last is read again, on a walk that does not complete, and not
    by the ``materialise PATH`` run before it.  Once the person downloads the first, it is read too."""
    online_rel, excluded_rel = "projects/Contoso online review.pdf", "archive/Contoso old review.pdf"
    _commented(local_source_dir / REVIEW, "Contoso widget review")
    online = _commented(local_source_dir / online_rel, "Contoso online review")
    _commented(local_source_dir / excluded_rel, "Contoso old review")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
    with Manifest(sample_config.state_paths.db) as m:
        record = m.get_meta(cycle_mod._REREAD_META + SID)
    inos = {online.stat().st_ino}
    real = materialise.is_dataless

    def is_dataless(st: Any) -> bool:
        return st.st_ino in inos or real(st)

    monkeypatch.setattr(materialise, "is_dataless", is_dataless)
    monkeypatch.setattr(al, "is_dataless", is_dataless)
    config = _incomplete(sample_config, "archive")
    fetched = _fetches(monkeypatch)
    named = run(config, materialise_paths=[local_source_dir / "README.txt"], budget_bytes=10_000_000)
    assert named.exit_code == 0 and fetched == ["README.txt"], "materialise reads the file it names"
    with Manifest(config.state_paths.db) as m:
        assert m.get_meta(cycle_mod._REREAD_META + SID) == record, "and looks for nothing to read again"
    second = run(config)
    [rep] = second.sources
    assert not rep.enumeration_complete and sorted(fetched[1:]) == [REVIEW, PLAIN_PDF]
    assert (rep.deferred, rep.deferred_online_only, rep.materialised_bytes) == (0, 0, 0)
    assert COMMENTS in _mirror_page(config, REVIEW)[1] and COMMENTS not in _mirror_page(config, online_rel)[1]
    rows = _file_rows(config)
    assert rows[online_rel].state is RowState.DATALESS and rows[online_rel].last_verdict is Verdict.DATALESS
    assert (rows[excluded_rel].state, rows[excluded_rel].state_reason) == (
        RowState.TOMBSTONE,
        "retired:scope-change",
    )
    step = loop.next_step(config)
    assert step.rule != 3 and not any("online-only" in line for line in step.lines()), step.lines()
    assert _reread_record(config) == (True, []), "a file that is not on this Mac is not waited for"
    assert run(config).commit_sha is None and len(fetched) == 3
    # The person downloads it (Finder's Download Now changes the inode's flags, so its change time).  The
    # walk sees the row move and the work queue reads the file as it reads any touched file.  It finds the
    # bytes its page was made from and a page from before comments, and converts it there and then.
    inos.clear()
    online.chmod(0o600)
    assert run(config).commit_sha is not None and fetched[3:] == [online_rel]
    assert COMMENTS in _mirror_page(config, online_rel)[1] and _reread_record(config) == (True, [])
    assert run(config).commit_sha is None and len(fetched) == 4


def test_a_file_evicted_after_the_walk_listed_it_is_not_downloaded_for_a_re_read(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The walk lists a file as on this Mac and the provider evicts it before it is read.  The work queue
    would download it within the byte budget; a re-read never does.  The row is not deferred, the loop has
    nothing to wait for, and the file is not remembered as tried: it is asked about again."""
    review = _commented(local_source_dir / REVIEW, "Contoso widget review")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
    before = _mirror_bytes(sample_config, REVIEW)
    evicted, real = review.stat().st_ino, materialise.is_dataless
    monkeypatch.setattr(materialise, "is_dataless", lambda st: st.st_ino == evicted or real(st))
    report = run(sample_config, budget_bytes=10_000_000)
    [rep] = report.sources
    assert (rep.deferred, rep.deferred_online_only, rep.materialised_bytes, rep.errors) == (0, 0, 0, ())
    assert _mirror_bytes(sample_config, REVIEW) == before
    row = _file_rows(sample_config)[REVIEW]
    assert (row.state, row.state_reason, row.last_verdict) == (RowState.LIVE, None, Verdict.UNCHANGED)
    assert _reread_record(sample_config) == (False, []) and loop.next_step(sample_config).rule != 3
    monkeypatch.setattr(materialise, "is_dataless", real)  # on this Mac again
    assert run(sample_config).commit_sha is not None and COMMENTS in _mirror_page(sample_config, REVIEW)[1]


def test_a_file_that_cannot_be_read_again_is_asked_about_later_and_no_log_line_names_it(
    sample_config: Config,
    local_source_dir: Path,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """The file was readable when it was converted and is not now.  Its page stays, the row keeps its
    verdict (it is no pending work), and nothing is counted against the file: it never reached a converter,
    and the fault may pass (a permission, a provider that timed out).  A later cycle asks again.  What the
    cycle says about re-reads at INFO is its one line, and no line at any level holds a name or a path."""
    _commented(local_source_dir / REVIEW, "Contoso widget review")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
    before = _mirror_bytes(sample_config, REVIEW)
    real = LocalArm.fetch
    asked: list[str] = []
    unreadable = {REVIEW}

    def fetch(self: Any, item: Any, dest: Path, budget: Any) -> Any:
        asked.append(item.rel_path)
        if item.rel_path in unreadable:
            raise PermissionError(13, "Permission denied", f"/Users/someone/source/{item.rel_path}")
        return real(self, item, dest, budget)

    monkeypatch.setattr(LocalArm, "fetch", fetch)
    caplog.set_level(logging.DEBUG, logger="agentsync.cycle")
    report = run(sample_config)
    [rep] = report.sources
    assert (report.exit_code, rep.errors, rep.alarms) == (0, (), ()) and sorted(asked) == [REVIEW, PLAIN_PDF]
    assert _mirror_bytes(sample_config, REVIEW) == before
    row = _file_rows(sample_config)[REVIEW]
    assert (row.state, row.last_verdict) == (RowState.LIVE, Verdict.UNCHANGED)
    assert _reread_record(sample_config) == (False, []) and _reread_failed(sample_config) == {}
    assert loop.next_step(sample_config).rule != 3
    said = [r for r in caplog.records if r.name == "agentsync.cycle" and "read again" in r.getMessage()]
    assert [r.getMessage() for r in said if r.levelno >= logging.INFO] == [
        "1 file(s) converted before a capability this install has were read again; 0 of them could not be "
        "converted and keep the page they had"
    ]
    assert [r.getMessage() for r in said if r.levelno < logging.INFO] == [
        "a file could not be read again (PermissionError); its page is as it was"
    ]
    assert run(sample_config).commit_sha is None and asked.count(REVIEW) == 2, "asked about again"
    unreadable.clear()
    assert run(sample_config).commit_sha is not None and COMMENTS in _mirror_page(sample_config, REVIEW)[1]
    assert _reread_record(sample_config) == (True, []) and asked.count(REVIEW) == 3
    assert run(sample_config).commit_sha is None and asked.count(REVIEW) == 3


def test_an_error_while_a_file_is_read_again_does_not_fail_the_source(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Something other than the conversion goes wrong while a file is read again (here: its page cannot be
    written).  The source has synced and stays synced: its removals still run, the report carries one
    line with no name in it, and the cycle exits 0.  The read counts against the file, and as pending work
    the next pass checks its page against the manifest; the files behind it are read by that pass."""
    _commented(local_source_dir / REVIEW, "Contoso widget review")
    doomed = "projects/sample.md"
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0 and run(sample_config).commit_sha is None
    (local_source_dir / doomed).unlink()  # absent from the next two passes: the second removes it
    real = cycle_mod._Cycle._after_fetch
    read: list[str] = []

    def after_fetch(self: Any, src: Any, row: Any, fetched: Any, acc: Any, *, reread: bool = False) -> Any:
        if reread:
            read.append(row.rel_path)
            if len(read) == 1:
                raise OSError(28, "No space left on device", f"/Users/someone/docs/{row.rel_path}")
        return real(self, src, row, fetched, acc, reread=reread)

    monkeypatch.setattr(cycle_mod._Cycle, "_after_fetch", after_fetch)
    first = run(sample_config)
    [rep] = first.sources
    assert first.exit_code == 0 and rep.errors == ("reading files again stopped: OSError",)
    assert any("1 file(s) absent from this complete pass" in a for a in rep.alarms), rep.alarms
    (stopped,) = read
    row = _file_rows(sample_config)[stopped]
    assert row.last_verdict is Verdict.MAYBE_CHANGED and _reread_record(sample_config) == (False, [])
    assert _reread_failed(sample_config) == {row.stable_id: 1}, "one of its two attempts"
    # The next pass's work reads it, finds its page whole, and reads it again there for what it lacks; the
    # file behind it is read once that pass has listed it.
    second = run(sample_config)
    [rep] = second.sources
    (other,) = {REVIEW, PLAIN_PDF} - {stopped}
    assert second.exit_code == 0 and rep.errors == () and sorted(read) == sorted([stopped, stopped, other])
    assert _file_rows(sample_config)[stopped].last_verdict in (Verdict.UNCHANGED, Verdict.OUTPUT_UNCHANGED)
    assert COMMENTS in _mirror_page(sample_config, REVIEW)[1]
    assert _file_rows(sample_config)[doomed].state is RowState.TOMBSTONE
    assert _reread_record(sample_config) == (True, []) and _reread_failed(sample_config) == {}
    assert run(sample_config).exit_code == 0 and len(read) == 3


def test_a_re_read_leaves_the_deletion_breaker_and_its_held_files_as_they_were(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Files absent from a complete listing are deletion candidates, held while the breaker is tripped.
    They are not read again and not tried, the breaker counts what it counted, and nothing is removed; the
    files the pass did list are read again all the same."""
    config = config_with(
        tmp_path, local_source_dir, "\n[breaker]\nfraction = 0.2\nfloor = 2\nhold_days = 7\n"
    )
    _commented(local_source_dir / REVIEW, "Contoso widget review")
    victims = [f"reviews/Contoso review {n}.pdf" for n in range(6)]
    for n, rel in enumerate(victims):
        _commented(local_source_dir / rel, f"Contoso review {n}")
    with _before_comments(monkeypatch):
        assert run(config).commit_sha is not None
    for rel in victims:
        (local_source_dir / rel).unlink()
    fetched = _fetches(monkeypatch)
    report = run(config)
    rep = source_report(report)
    assert rep.breaker_tripped and any("breaker TRIPPED: 6 absent file(s) held" in a for a in rep.alarms)
    assert sorted(fetched) == [REVIEW, PLAIN_PDF]
    assert [c.path for c in report.changes] == [slug.mirror_rel_path(SID, REVIEW)]
    rows = _file_rows(config)
    assert all(rows[rel].state is RowState.LIVE for rel in victims)
    assert all(
        page(config.docs_repo, slug.mirror_rel_path(SID, rel))[0]["status"] == "current" for rel in victims
    )
    with Manifest(config.state_paths.db) as m:
        srow = m.get_source(SID)
        assert (
            srow is not None
            and srow.breaker_candidates == 6
            and m.breaker_active(SID, "2026-09-29T12:00:00Z")
        )
    assert _reread_record(config) == (False, []), "the held files may come back: the look stays open"
    held = run(config)
    assert source_report(held).breaker_tripped and held.changes == () and len(fetched) == 2
    assert all(_file_rows(config)[rel].state is RowState.LIVE for rel in victims)


def test_a_file_an_incomplete_pass_did_not_list_is_not_read_again(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An incomplete pass holds deletions: a file it did not list is unknown, never gone.  It is not read
    again either (there may be nothing to read), and the files the pass did list are."""
    absent_rel = "projects/Contoso absent review.pdf"
    absent = _commented(local_source_dir / absent_rel, "Contoso absent review")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
    absent.unlink()
    config = _incomplete(sample_config)
    fetched = _fetches(monkeypatch)
    report = run(config)
    [rep] = report.sources
    assert not rep.enumeration_complete and not rep.breaker_tripped and fetched == [PLAIN_PDF]
    assert report.changes == () and report.commit_sha is None, "read again, the same page: nothing to commit"
    assert _file_rows(config)[absent_rel].state is RowState.LIVE
    assert _mirror_page(config, absent_rel)[0]["status"] == "current"
    assert _reread_record(config) == (False, []), "it may be listed again: the look stays open"
    assert run(config).commit_sha is None and fetched == [PLAIN_PDF]


def test_a_file_the_engine_fails_on_twice_is_given_up_and_an_engine_that_comes_back_reads_what_is_new(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """OCR never fails a document: a PDF the engine fails on keeps the page it has without OCR, under the
    version without OCR.  That version is what a re-read looks for, so the file is given up after the
    second cycle in which the engine fails on it, and is not read every cycle.  An engine that goes away
    and comes back reads the file converted meanwhile, and neither the one it read before nor the one it
    failed on."""
    scans = local_source_dir / "scans"
    scans.mkdir()

    def scan(name: str, shade: int) -> str:
        pages = [([f"{name}: the cover page has a text layer of its own"], []), ([], [page_picture(shade)])]
        build_picture_pdf(scans / f"Contoso {name}.pdf", pages)
        return f"scans/Contoso {name}.pdf"

    good, bad = scan("good scan", 91), scan("bad scan", 95)
    assert run(sample_config).exit_code == 0
    said = {91: ["Signed in Rotterdam"], 93: ["Counted in Antwerp"], 95: "fail"}
    engine = shade_engine(tmp_path / "ocr-bin", said)
    found: list[Any] = [engine]
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: found[0])
    fetched = _fetches(monkeypatch)
    unread = _mirror_bytes(sample_config, bad)
    second = run(sample_config)
    assert second.exit_code == 0 and second.sources[0].errors == () and second.sources[0].alarms == ()
    assert _mirror_page(sample_config, good)[1].rstrip().endswith("Signed in Rotterdam")
    assert _mirror_bytes(sample_config, bad) == unread and b"ocr" not in unread[bad]
    bad_id = _file_rows(sample_config)[bad].stable_id
    assert _reread_record(sample_config) == (False, []) and _reread_failed(sample_config) == {bad_id: 1}
    # The helper reads the blank image, so the failure is the file's.  A second cycle of it gives it up.
    assert run(sample_config).commit_sha is None and _reread_record(sample_config) == (True, [bad_id])
    runs = len(calls(engine.helper))
    assert run(sample_config).commit_sha is None and len(calls(engine.helper)) == runs
    assert (fetched.count(good), fetched.count(bad)) == (1, 2)
    # The engine goes away.  A scan mirrored meanwhile is converted as on a Mac without one.
    found[0] = None
    late = scan("late scan", 93)
    assert (
        run(sample_config).exit_code == 0 and "ocr" not in _mirror_page(sample_config, late)[0]["converter"]
    )
    assert run(sample_config).commit_sha is None and len(calls(engine.helper)) == runs
    # It comes back: the late scan is read again, the other two are not.
    found[0] = engine
    assert run(sample_config).commit_sha is not None
    assert _mirror_page(sample_config, late)[1].rstrip().endswith("Counted in Antwerp")
    assert (fetched.count(good), fetched.count(bad), fetched.count(late)) == (1, 2, 2)
    assert _mirror_bytes(sample_config, bad) == unread and _reread_record(sample_config) == (True, [bad_id])
    assert run(sample_config).commit_sha is None and fetched.count(late) == 2
    # A new scan the engine fails on is converted without OCR.  That leaves a file to read again, so the
    # two cycles after read it once more each, and then it is given up too.
    worse = scan("worse scan", 95)
    assert run(sample_config).exit_code == 0 and fetched.count(worse) == 1
    assert "ocr" not in _mirror_page(sample_config, worse)[0]["converter"]
    assert _reread_record(sample_config) == (False, [bad_id])
    assert run(sample_config).commit_sha is None and fetched.count(worse) == 2
    assert _reread_record(sample_config) == (False, [bad_id])
    assert run(sample_config).commit_sha is None and fetched.count(worse) == 3
    worse_id = _file_rows(sample_config)[worse].stable_id
    assert _reread_record(sample_config) == (True, sorted([bad_id, worse_id]))
    assert run(sample_config).commit_sha is None and fetched.count(worse) == 3


def test_re_reads_stop_at_their_time_and_go_on_in_the_next_cycle_past_a_file_that_fails(
    sample_config: Config,
    local_source_dir: Path,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Six PDFs to read again, each read 'taking' 25 s of a cycle's 60: three a cycle.  One of them fails
    every time it is converted.  It costs its two reads and holds nobody up: the other five are read once
    each, and after three cycles nothing is left.  A cycle says what it did in one line, a count, with no
    name in it."""
    reviews = [f"reviews/Contoso review {n}.pdf" for n in range(5)]
    for n, rel in enumerate(reviews):
        _commented(local_source_dir / rel, f"Contoso review {n}")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
    stuck = _mirror_bytes(sample_config, reviews[0])
    real = pdf_mod.PdfConverter.convert

    def convert(self: Any, src: Path, *, name: str) -> Any:
        if name == "Contoso review 0.pdf":
            raise RuntimeError("this file breaks the converter")
        return real(self, src, name=name)

    monkeypatch.setattr(pdf_mod.PdfConverter, "convert", convert)
    ticks = itertools.count(0.0, 25.0)
    monkeypatch.setattr(cycle_mod, "_reread_clock", lambda: next(ticks))
    monkeypatch.setattr(cycle_mod, "_REREAD_BUDGET_S", 60.0)
    monkeypatch.setattr(cycle_mod, "_REREAD_BATCH", 2)
    fetched = _fetches(monkeypatch)
    caplog.set_level(logging.INFO, logger="agentsync.cycle")
    first = run(sample_config)
    assert first.exit_code == 0 and len(fetched) == 3 and not _reread_record(sample_config)[0]
    record = _run_record(sample_config)
    assert record["reread_left"] == 3 + record.get("reread_kept", 0), "six, less the ones that were read"
    assert _reread_notes(sample_config) == [_sync_again(record["reread_left"])]
    assert run(sample_config).exit_code == 0 and len(fetched) == 6
    assert run(sample_config).exit_code == 0 and len(fetched) == 7
    assert "reread_left" not in _run_record(sample_config) and _reread_notes(sample_config) == []
    assert sorted(fetched) == sorted([*reviews, reviews[0], PLAIN_PDF]), "each once, the one that fails twice"
    stuck_id = _file_rows(sample_config)[reviews[0]].stable_id
    assert _reread_record(sample_config) == (True, [stuck_id])
    assert _mirror_bytes(sample_config, reviews[0]) == stuck
    assert all(COMMENTS in _mirror_page(sample_config, rel)[1] for rel in reviews[1:])
    assert run(sample_config).commit_sha is None and len(fetched) == 7
    said = [
        r.getMessage()
        for r in caplog.records
        if r.name == "agentsync.cycle" and "read again" in r.getMessage()
    ]
    line = re.compile(
        r"(\d) file\(s\) converted before a capability this install has were read again; "
        r"(\d) of them could not be converted and keep the page they had"
    )
    counts = [tuple(map(int, found.groups())) for found in map(line.fullmatch, said) if found]
    assert len(counts) == len(said) == 3 and [read for read, _kept in counts] == [3, 3, 1]
    assert sum(kept for _read, kept in counts) == 2
    assert "Contoso" not in "".join(said) and "sample" not in "".join(said)


def test_a_file_given_up_is_read_again_once_its_bytes_have_changed(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """What a file's re-reads did says nothing about other bytes.  A scan the engine failed on twice is
    given up; then the file is rewritten in place (the same stable id) with a page the engine can read, in
    a cycle whose OCR time is used, so it is converted without OCR.  The id used to stay among the tried,
    and the same cycle stored ``done`` again: the page never got its text.  The work queue's conversion
    takes the id out, and the next cycle reads the file."""
    rel = "scans/Contoso supply agreement.pdf"
    cover = ["The cover page has a text layer of its own"]
    (local_source_dir / "scans").mkdir()
    build_picture_pdf(local_source_dir / rel, [(cover, []), ([], [page_picture(95)])])
    assert run(sample_config).exit_code == 0
    engine = shade_engine(tmp_path / "ocr-bin", {91: ["Signed in Rotterdam"], 95: "fail"})
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    assert run(sample_config).exit_code == 0 and run(sample_config).exit_code == 0
    stable = _file_rows(sample_config)[rel].stable_id
    assert _reread_record(sample_config) == (True, [stable])
    build_picture_pdf(local_source_dir / rel, [(cover, []), ([], [page_picture(91)])])
    with monkeypatch.context() as spent:
        spent.setattr(cycle_mod, "_OCR_BUDGET_S", 0.0)
        assert run(sample_config).exit_code == 0
    assert _file_rows(sample_config)[rel].stable_id == stable, "rewritten in place"
    assert "ocr" not in _mirror_page(sample_config, rel)[0]["converter"]
    assert _reread_record(sample_config) == (False, []) and _reread_failed(sample_config) == {}
    assert run(sample_config).commit_sha is not None and _reread_record(sample_config) == (True, [])
    assert _mirror_page(sample_config, rel)[1].rstrip().endswith("Signed in Rotterdam")


def test_a_file_the_walk_calls_maybe_changed_in_every_pass_is_read_again_by_the_work_queue(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A volume that reports no generation count: every pass calls every file maybe-changed, so the work
    queue reads each one every cycle and no pass ever lists one unchanged.  Such a file was never read
    again, and its source's record stayed open for ever.  The queue has the file in hand, with the bytes
    its page was made from: it converts it there."""
    real = al._file_attrs
    monkeypatch.setattr(al, "_file_attrs", lambda path: dataclasses.replace(real(path), gen_count=None))
    _commented(local_source_dir / REVIEW, "Contoso widget review")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
        assert run(sample_config).commit_sha is None
    assert COMMENTS not in _mirror_page(sample_config, REVIEW)[1]
    plain = _mirror_bytes(sample_config, PLAIN_PDF)
    fetched = _fetches(monkeypatch)
    second = run(sample_config)
    [rep] = second.sources
    assert second.commit_sha is not None and fetched.count(REVIEW) == 1, "read once, by the queue itself"
    assert rep.converted == 0 and rep.errors == (), "the same bytes are no new conversion"
    assert COMMENTS in _mirror_page(sample_config, REVIEW)[1]
    assert _mirror_bytes(sample_config, PLAIN_PDF) == plain, "no comment in it: its page is as it was"
    assert _file_rows(sample_config)[PLAIN_PDF].last_verdict is Verdict.OUTPUT_UNCHANGED
    assert _reread_record(sample_config) == (True, [])
    for query in ("produced_by", "reread_candidates", "reread_left", "reread_count"):
        monkeypatch.setattr(Manifest, query, crash)
    assert run(sample_config).commit_sha is None, "nothing is left, and no cycle looks"


def test_a_cycle_that_died_while_a_file_was_read_again_counts_against_that_file(
    sample_config: Config, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A re-read that hangs or crashes the process (a page PDFium cannot render, the launcher's watchdog)
    takes the manifest transaction with it, so nothing in it can remember the file, and every later cycle
    would start with the same one.  Before a file is read the record says so, in a write committed on its
    own; a cycle that finds such a mark counts one failed read.  Two of them give the file up, and the
    files behind it are read."""
    _commented(local_source_dir / REVIEW, "Contoso widget review")
    with _before_comments(monkeypatch):
        assert run(sample_config).exit_code == 0
    key, review_id = cycle_mod._REREAD_META + SID, _file_rows(sample_config)[REVIEW].stable_id
    real = cycle_mod._Cycle._reread
    marks: list[tuple[str, object]] = []

    def reread(self: Any, src: Any, arm: Any, row: Any, acc: Any) -> Any:
        with Manifest(sample_config.state_paths.db) as other:  # what a new process would find on disk
            marks.append((row.stable_id, json.loads(other.get_meta(key) or "[{}]")[0].get("reading")))
        return real(self, src, arm, row, acc)

    monkeypatch.setattr(cycle_mod._Cycle, "_reread", reread)
    assert run(sample_config).commit_sha is not None and len(marks) == 2
    assert all(reading == stable for stable, reading in marks), "committed before the read starts"
    with Manifest(sample_config.state_paths.db) as m:
        done = json.loads(m.get_meta(key) or "")
        assert "reading" not in done[0] and done[0]["done"] is True
        # What two cycles that died in the read of one file leave: the mark of the second, the count of
        # the first.  (The files are outdated again: the capabilities are what the record is for.)
        left = [
            {
                "done": False,
                "for": done[0]["for"],
                "tried": [],
                "failed": {review_id: 1},
                "reading": review_id,
            }
        ]
        m.set_meta(key, json.dumps(left, sort_keys=True))
    monkeypatch.setattr(pdf_mod.PdfConverter, "outdated", lambda self, produced, reason=None: True)
    del marks[:]
    fetched = _fetches(monkeypatch)
    assert run(sample_config).exit_code == 0
    assert REVIEW not in fetched and fetched == [PLAIN_PDF], "given up; the file behind it is read"
    assert [stable for stable, _reading in marks] != [review_id]


def test_a_file_a_materialise_run_converted_without_ocr_is_read_again_by_the_next_sync(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A ``materialise PATH`` run reads nothing again, and it can still leave a file to read again: the
    file it names, converted after the cycle's OCR time was used.  The source's record says so in that
    run, so the next sync looks although the run before it had found nothing left."""
    engine = shade_engine(tmp_path / "ocr-bin", {91: ["Signed in Rotterdam"], 93: ["Counted in Antwerp"]})
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    rel = "scans/Contoso supply agreement.pdf"
    cover = ["The cover page has a text layer of its own"]
    (local_source_dir / "scans").mkdir()
    build_picture_pdf(local_source_dir / rel, [(cover, []), ([], [page_picture(91)])])
    assert run(sample_config).exit_code == 0 and _reread_record(sample_config) == (True, [])
    assert _mirror_page(sample_config, rel)[1].rstrip().endswith("Signed in Rotterdam")
    build_picture_pdf(local_source_dir / rel, [(cover, []), ([], [page_picture(93)])])
    with monkeypatch.context() as spent:
        spent.setattr(cycle_mod, "_OCR_BUDGET_S", 0.0)
        named = run(sample_config, materialise_paths=[local_source_dir / rel], budget_bytes=10_000_000)
    fm, body = _mirror_page(sample_config, rel)
    assert named.exit_code == 0 and "ocr" not in fm["converter"] and "Antwerp" not in body
    assert _reread_record(sample_config) == (False, [])
    assert run(sample_config).commit_sha is not None and _reread_record(sample_config) == (True, [])
    assert _mirror_page(sample_config, rel)[1].rstrip().endswith("Counted in Antwerp")
    assert run(sample_config).commit_sha is None


def test_a_converter_that_cannot_run_costs_its_files_neither_a_page_nor_their_one_re_read(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """pandoc cannot be run in the cycle that would read the Word files again.  That is no failure of those
    files: they are not read, their pages stay, and they are not remembered as tried, so the cycle in which
    pandoc runs again reads them.  The files of the other converters are read meanwhile."""
    assert run(sample_config).exit_code == 0
    words = ["projects/sample.docx", KICKOFF]
    before = _mirror_bytes(sample_config, *words)
    _use_ocr(monkeypatch, tmp_path)
    fetched = _fetches(monkeypatch)

    def missing(self: Any) -> str:
        raise ConversionError("cannot run pandoc: [Errno 2] No such file or directory")

    with monkeypatch.context() as fault:
        fault.setattr(pandoc_mod._PandocRunner, "version", missing)
        report = run(sample_config)
    [rep] = report.sources
    assert (report.exit_code, rep.errors, rep.alarms) == (0, (), ())
    assert sorted(fetched) == [PLAIN_PDF, "projects/sample.pptx"]
    assert _mirror_bytes(sample_config, *words) == before and _reread_record(sample_config) == (False, [])
    assert run(sample_config).exit_code == 0 and sorted(fetched[2:]) == sorted(words)
    assert _mirror_bytes(sample_config, *words) == before, "no picture in them: the pages stay as they are"
    assert _reread_record(sample_config) == (True, [])
    assert run(sample_config).commit_sha is None and len(fetched) == 4


def test_the_loop_says_sync_again_while_a_sync_reads_more_of_what_is_left_and_not_after_one_read_none(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Each run records how many files its re-read left, and ``sync`` prints that as a note.  The note
    says "sync again" while the last sync read some of them.  With pandoc missing the two Word files are
    never read: the first such sync still read the others, the second read nothing, and from then the note
    no longer says "sync again", so an agent told to follow that note stops.  The setup prompt is."""
    assert run(sample_config).exit_code == 0
    assert "reread_left" not in _run_record(sample_config) and _reread_notes(sample_config) == []
    _use_ocr(monkeypatch, tmp_path)

    def missing(self: Any) -> str:
        raise ConversionError("cannot run pandoc: [Errno 2] No such file or directory")

    with monkeypatch.context() as fault:
        fault.setattr(pandoc_mod._PandocRunner, "version", missing)
        assert run(sample_config).exit_code == 0
        record = _run_record(sample_config)
        assert (record["reread"], record["reread_left"]) == (2, 2), "the PDF and the deck; two Word files"
        assert _reread_notes(sample_config) == [_sync_again(2)]
        assert run(sample_config).exit_code == 0
        record = _run_record(sample_config)
        assert "reread" not in record and record["reread_left"] == 2
        assert _reread_notes(sample_config) == [
            f"2 file(s) in {SID} wait to be read again, and the last sync read none of them (a converter or "
            "on-device OCR that cannot run, or a folder that could not be listed): another sync does not "
            "clear it; it does not block the next step"
        ]
    assert run(sample_config).exit_code == 0, "pandoc runs again: whichever sync comes next reads them"
    record = _run_record(sample_config)
    assert record["reread"] == 2 and "reread_left" not in record and _reread_notes(sample_config) == []
    assert all(isinstance(value, int) for value in record.values())


def test_downloads_that_take_the_ocr_time_end_the_sync_again_note_and_a_sync_without_them_reads_on(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A folder of scans that are still online-only.  Each sync downloads two: the first uses up the
    cycle's OCR time, the second is converted without OCR and joins the files to read again, and the
    re-read never starts.  The note said "sync again: N file(s)" with a larger N after every such sync,
    and the setup prompt ran twelve more of them.  It now says the downloads took the time, and does not
    say "sync again".  A sync that downloads nothing (``--materialise-budget 0``, what the setup prompt
    runs before its report) has the OCR time for them: the count falls, the note says "sync again" while
    it does, and it ends."""
    assert run(sample_config).exit_code == 0
    engine = shade_engine(tmp_path / "ocr-bin", {90 + n: [f"Delivery note {n}"] for n in range(8)})
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: engine)
    assert run(sample_config).exit_code == 0 and _reread_notes(sample_config) == [], "the samples: read"
    (local_source_dir / "scans").mkdir()
    cover = (["The cover page has a text layer of its own"], [])
    scans = [
        build_picture_pdf(
            local_source_dir / "scans" / f"Contoso scan {n}.pdf", [cover, ([], [page_picture(90 + n)])]
        )
        for n in range(8)
    ]
    online = {scan.stat().st_ino for scan in scans}
    real_dataless, real_materialise = materialise.is_dataless, al.materialise

    def is_dataless(st: Any) -> bool:  # mocked SF_DATALESS: no File Provider in a test
        return st.st_ino in online or real_dataless(st)

    def download(src: Path, dest: Path, budget: Any) -> Any:
        result = real_materialise(src, dest, budget)
        online.discard(src.stat().st_ino)  # a file that was read is on this Mac from now on
        return result

    monkeypatch.setattr(materialise, "is_dataless", is_dataless)
    monkeypatch.setattr(al, "is_dataless", is_dataless)
    monkeypatch.setattr(al, "materialise", download)
    monkeypatch.setattr(cycle_mod, "_OCR_BUDGET_S", 1e-9)  # the first scanned page of a sync uses it up
    two = 2 * scans[0].stat().st_size

    def sync(budget: int) -> tuple[int | None, int | None, list[str]]:
        """One sync with this download budget: (files read again, files left, the loop's notes)."""
        assert run(sample_config, budget_bytes=budget).exit_code == 0
        record = _run_record(sample_config)
        assert record["ocr_over"] == 1
        return record.get("reread"), record.get("reread_left"), _reread_notes(sample_config)

    crowded = (
        "{n} file(s) in local-fixture wait to be read again, and the last sync read none of them: new and "
        "changed files took its OCR time, more files joined, and more downloads wait. They are read once a "
        "sync has OCR time left; it does not block the next step"
    )
    # A file downloaded in a sync is counted from the next one, whose listing finds it on this Mac.
    assert sync(two) == (None, None, [])
    assert sync(two) == (None, 1, [crowded.format(n=1)])
    assert sync(two) == (None, 2, [crowded.format(n=2)]), "two more scans are still online-only"
    assert len(online) == 2 and "sync again" not in crowded
    # No download: each sync reads one scan again before its OCR time is used, and the count falls.
    assert sync(0) == (1, 2, [_sync_again(2)]), "three were left: the one the third sync converted is counted"
    assert sync(0) == (1, 1, [_sync_again(1)])
    assert sync(0) == (1, None, [])
    assert len(online) == 2, "nothing was downloaded for it"
    assert run(sample_config, budget_bytes=0).commit_sha is None and _reread_notes(sample_config) == []


def test_a_sync_whose_listing_macos_holds_does_not_leave_the_sync_again_note_standing(
    sample_config: Config, local_source_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A sync read some files again and left two.  Then a macOS privacy prompt holds the folder's listing:
    each sync waits out the listing's time limit and never gets to the re-read, so its run record says
    nothing about one.  The note went by the last run that did, and "sync again: 2 file(s)" stood through
    every held sync.  It now goes by the source's newest pass, and says another sync does not clear it.
    Once the listing is back, the next sync reads the files and the note is gone."""
    assert run(sample_config).exit_code == 0
    _use_ocr(monkeypatch, tmp_path)

    def missing(self: Any) -> str:
        raise ConversionError("cannot run pandoc: [Errno 2] No such file or directory")

    with monkeypatch.context() as fault:
        fault.setattr(pandoc_mod._PandocRunner, "version", missing)
        assert run(sample_config).exit_code == 0
    assert _reread_notes(sample_config) == [_sync_again(2)], "the two Word files: pandoc was missing"
    real = al._list_dir
    release = threading.Event()

    def held(path: Path, *, dir_dataless: bool) -> list[os.DirEntry[str]]:
        if path == local_source_dir / "projects":
            release.wait(5.0)
        return real(path, dir_dataless=dir_dataless)

    with monkeypatch.context() as prompt:
        prompt.setattr(al, "_list_dir", held)
        prompt.setattr(al, "LISTING_TIMEOUT_S", 0.2, raising=False)
        try:
            for _sync in range(2):
                report = run(sample_config)
                assert any("click Allow" in alarm for alarm in report.sources[0].alarms)
                assert "reread_for" not in _run_record(sample_config), "it never got to the re-read"
                assert _reread_notes(sample_config) == [
                    f"file(s) in {SID} wait to be read again, and the last sync read none of them (a "
                    "converter or on-device OCR that cannot run, or a folder that could not be listed): "
                    "another sync does not clear it; it does not block the next step"
                ]
        finally:
            release.set()
    assert run(sample_config).exit_code == 0
    assert _run_record(sample_config)["reread"] == 2 and _reread_notes(sample_config) == []


def test_a_graph_file_is_never_read_again(
    tmp_path: Path, local_source_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Every read of a Graph item is a download, and nothing is downloaded for a re-read: a PDF a drive
    source mirrored before comments were kept gains them when its bytes next change."""
    config = config_with(tmp_path, local_source_dir, GRAPH_SOURCE)
    data = _commented(tmp_path / "review.pdf", "Contoso widget review").read_bytes()
    drive = FakeDrive({"I1": ("Contoso widget review.pdf", data)})
    with GraphClient(
        FakeTokens(),
        user_agent="NONISV|test|agentsync/0",
        transport=httpx.MockTransport(drive.handler),
        sleep=lambda _s: None,
    ) as client:
        with _before_comments(monkeypatch):
            assert run(config, client=client, only=["drive"]).exit_code == 0
        downloads = [entry for entry in drive.log if entry.endswith("/content")]
        assert len(downloads) == 1
        assert run(config, client=client, only=["drive"]).commit_sha is None
        assert [entry for entry in drive.log if entry.endswith("/content")] == downloads
    [mirrored] = (config.docs_repo / "mirror" / "drive").rglob("*.pdf.md")
    text = mirrored.read_text(encoding="utf-8")
    assert "pdf-pypdfium2@2.0.0+" in text and COMMENTS not in text
    with Manifest(config.state_paths.db) as m:
        assert m.get_meta(cycle_mod._REREAD_META + "drive") is None


# ---------------------------------------------------------------------------------------------------------
# meeting recordings: the recording pass after every source (spec S0 rules 3 to 8, 4.1, S10 rule 2)
# ---------------------------------------------------------------------------------------------------------

PIECE_MS = 300_000  # five minutes of the picture track (spec 4.1)
STUB_MAGIC = b"STUB-REC:"


class FakePieces(PieceStore):
    """The piece store's contract (``convert/pieces.py``) on disk: one file per piece, ``progress.json``."""

    def load(self, sha256: str, piece: int, key: str) -> bytes | None:
        try:
            return (self.root / sha256 / f"{piece:04d}-{key}").read_bytes()
        except OSError:
            return None

    def save(self, sha256: str, piece: int, key: str, data: bytes) -> None:
        folder = self.root / sha256
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob(f"{piece:04d}-*"):
            old.unlink()
        (folder / f"{piece:04d}-{key}").write_bytes(data)

    def set_progress(self, sha256: str, *, done_ms: int, total_ms: int) -> None:
        (self.root / sha256).mkdir(parents=True, exist_ok=True)
        (self.root / sha256 / "progress.json").write_text(json.dumps([done_ms, total_ms]))

    def progress(self, sha256: str) -> tuple[int, int] | None:
        try:
            done, total = json.loads((self.root / sha256 / "progress.json").read_text())
        except (OSError, ValueError):
            return None
        return int(done), int(total)

    def pending(self) -> list[str]:
        return sorted(p.name for p in self.root.iterdir() if p.is_dir()) if self.root.is_dir() else []

    def remove(self, sha256: str) -> bool:
        if not (self.root / sha256).is_dir():
            return False
        shutil.rmtree(self.root / sha256)
        return True

    def prune(self, keep: Any) -> int:
        kept = set(keep)
        return sum(self.remove(sha) for sha in self.pending() if sha not in kept)


class FakeMedia(MediaEngine):
    """A media helper that answers ``--version`` while ``up``."""

    def __init__(self) -> None:
        super().__init__(Path("/nonexistent/media-frames"), name="paper-media", helper_version="0.1.0")
        self.up = True
        self.asked = 0

    def alive(self) -> bool:
        self.asked += 1
        return self.up


@dataclasses.dataclass
class Rec:
    """What the stub recording converter did, and the faults a test sets for it."""

    media: FakeMedia
    allowances: list[Allowance] = dataclasses.field(default_factory=list)  # open work allowances
    opened: list[float | None] = dataclasses.field(default_factory=list)  # every allowance opened
    signals: list[RecordingNotFinished] = dataclasses.field(default_factory=list)
    worked: list[tuple[str, int]] = dataclasses.field(default_factory=list)  # (name, piece) read
    converts: list[str] = dataclasses.field(default_factory=list)  # every convert call, by name
    timeout_at: set[tuple[str, int]] = dataclasses.field(default_factory=set)
    crash_at: set[tuple[str, int]] = dataclasses.field(default_factory=set)
    media_fails: bool = False

    def stop(self, signal: RecordingNotFinished) -> None:
        self.signals.append(signal)
        raise signal


class StubRecording(RecordingConverter):
    """The recording converter's contract with the cycle, without AVFoundation: a file ``STUB-REC:{spec}``
    is ``minutes`` long, read in pieces of five minutes, each costing ``piece_s`` of the work allowance and
    ``ocr_s`` of the cycle's OCR time; ``text`` is on screen, ``sidecars`` ({name: text or hex}) are its
    keyframes and full text."""

    def __init__(self, rec: Rec, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self.rec = rec

    def version(self) -> str:
        return "1.0.0+stub-h0.1.0-s1"

    def options(self) -> dict[str, Any]:
        return {"piece_ms": PIECE_MS}

    @property
    def outdated_key(self) -> str:
        return "1.0.0<1.0.0|-|-"

    def outdated(self, produced: str, reason: str | None = None) -> bool:
        return False

    def convert(self, src: Path, *, name: str) -> tuple[RenderedUnit, ...]:
        self.rec.converts.append(name)
        data = src.read_bytes()
        spec = json.loads(data.split(STUB_MAGIC, 1)[1])
        sha, total = hashlib.sha256(data).hexdigest(), spec["minutes"] * 60_000
        allowance = self.rec.allowances[-1] if self.rec.allowances else None
        assert self._pieces is not None, "the cycle hands the converter its piece store"
        for n in range(-(-total // PIECE_MS)):
            if self._pieces.load(sha, n, "k1") is not None:
                continue
            if allowance is not None and allowance.used_up:
                self.rec.stop(RecordingNotFinished(done_ms=n * PIECE_MS, total_ms=total, timed_out=False))
            if (name, n) in self.rec.crash_at:
                raise Crash
            if (name, n) in self.rec.timeout_at:
                self.rec.stop(RecordingNotFinished(done_ms=n * PIECE_MS, total_ms=total, timed_out=True))
            if self.rec.media_fails:
                raise MediaError("the media helper exited 3")
            self.rec.worked.append((name, n))
            if isinstance(self._ocr, cycle_mod._CycleOcr):
                self._ocr.spent_s += spec.get("ocr_s", 0.0)
            self._pieces.save(sha, n, "k1", f"{name}:{n}".encode())
            self._pieces.set_progress(sha, done_ms=min((n + 1) * PIECE_MS, total), total_ms=total)
            if allowance is not None:
                allowance.spent_s += spec.get("piece_s", 70.0)
        lines = [f"[{n * 5:02d}:00] SCREEN {spec.get('text', 'Contoso quarterly review')}" for n in range(3)]
        body = f"# {name}\n\n" + "\n".join(lines) + "\n"
        sidecars = tuple(
            (key, bytes.fromhex(value[4:]) if value.startswith("hex:") else value.encode())
            for key, value in spec.get("sidecars", {}).items()
        )
        unit = RenderedUnit(
            unit_id="whole",
            kind=UnitKind.WHOLE,
            index=0,
            of=1,
            name="",
            file_stem="",
            title=name,
            summary=f"Recording of {spec['minutes']} minutes",
            tokens_estimate=estimate_tokens(body),
            body=body,
            rendered_sha256=rendered_sha256(body),
            sidecars=sidecars,
        )
        return (unit,)


def stub_recording(path: Path, minutes: int, **spec: Any) -> Path:
    """A recording the stub converter reads (see :class:`StubRecording`)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    spec.setdefault("text", f"Contoso review: {path.stem}")  # two recordings never share their bytes
    path.write_bytes(
        b"\x00\x00\x00\x18ftypmp42" + STUB_MAGIC + json.dumps({"minutes": minutes, **spec}).encode()
    )
    return path


@pytest.fixture
def rec(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Rec:
    """Every cycle from now on finds an OCR engine, a media helper and the piece store, and its registry
    reads recordings with :class:`StubRecording` (the cycle's side of the recording contracts)."""
    found = Rec(FakeMedia())
    _use_ocr(monkeypatch, tmp_path)
    monkeypatch.setattr(cycle_mod.media, "engine", lambda _convert, _cache_dir: found.media)
    monkeypatch.setattr(
        PieceStore, "under", classmethod(lambda cls, cache_dir: FakePieces(cache_dir / "recordings"))
    )

    def registry(config: Config, policy: Any, engine: Any, media: Any, pieces: Any) -> Registry:
        plain = Registry.default(config.convert, policy=policy, ocr=engine)
        if media is None:
            return plain
        stub = StubRecording(
            found, config.convert, engine, media, pieces=pieces, label_rule=policy.labels_active
        )
        built = Registry([*(c.inner for c in plain.converters()), stub], policy=policy, banner=True)  # type: ignore[attr-defined]
        built._without_ocr = plain.without_ocr
        return built

    monkeypatch.setattr(cycle_mod, "_cycle_registry", registry)

    @contextlib.contextmanager
    def work_allowance(seconds: float | None) -> Iterator[Allowance]:
        allowance = Allowance(seconds)
        found.opened.append(seconds)
        found.allowances.append(allowance)
        try:
            yield allowance
        finally:
            found.allowances.pop()

    monkeypatch.setattr(recording_mod, "work_allowance", work_allowance)
    real_convert = cycle_mod.convert_file

    def convert_file(*args: Any, **kwargs: Any) -> Any:
        """``convert_file`` re-raises the "not finished" signal (the recording teammate's contract)."""
        found.signals.clear()
        result = real_convert(*args, **kwargs)
        if found.signals:
            raise found.signals.pop()
        return result

    monkeypatch.setattr(cycle_mod, "convert_file", convert_file)
    return found


MEETING = "meetings/Contoso weekly sync.mp4"


@pytest.fixture
def meetings(tmp_path: Path) -> Path:
    """A small local source: the sentinel and one note (no converter needs pandoc)."""
    root = tmp_path / "meetings-source"
    root.mkdir()
    (root / "README.txt").write_text("sentinel: this file must always be present\n", encoding="utf-8")
    (root / "notes.md").write_text("# Notes\n\nContoso agenda\n", encoding="utf-8")
    return root


def _row(config: Config, rel: str, sid: str = SID) -> Any:
    return _file_rows(config, sid)[rel]


def _record(config: Config, sid: str = SID) -> dict[str, Any]:
    with Manifest(config.state_paths.db) as m:
        raw = m.get_meta(cycle_mod._RECORDING_META + sid)
    return json.loads(raw) if raw else {}


def _progress(config: Config, rel: str, sid: str = SID) -> str | None:
    stable = _row(config, rel, sid).stable_id
    with Manifest(config.state_paths.db) as m:
        return m.get_meta(f"{cycle_mod.RECORDING_PROGRESS_META}{sid}:{stable}")


def _recording_page(config: Config, rel: str, sid: str = SID) -> tuple[dict[str, Any], str]:
    return page(config.docs_repo, slug.mirror_rel_path(sid, rel))


def _pieces_of(config: Config, path: Path) -> Path:
    return config.cache_dir / "recordings" / hashlib.sha256(path.read_bytes()).hexdigest()


def test_an_interactive_sync_starts_no_recording_and_prints_the_count(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """Ruling 5: a tool's timeout may end an interactive sync, so it reads no recording and opens no work
    allowance; the recording waits (``RECORDING_WAITS``) and its document neighbours are published.  The next
    background sync reads it."""
    config = config_with(tmp_path, meetings)
    stub_recording(meetings / MEETING, 12)
    report = run(config, mode=None)
    assert report.exit_code == 0 and rec.converts == [] and rec.opened == []
    row = _row(config, MEETING)
    assert row.state_reason == cycle_mod.RECORDING_WAITS and row.last_verdict is Verdict.DEFERRED
    assert _row(config, "notes.md").last_verdict is Verdict.UNCHANGED
    assert not (config.docs_repo / slug.mirror_rel_path(SID, MEETING)).exists()
    assert run(config).exit_code == 0 and rec.opened == [cycle_mod._RECORDING_BUDGET_S]
    fm, body = _recording_page(config, MEETING)
    assert fm["status"] == "current" and "SCREEN Contoso review: Contoso weekly sync" in body
    assert _row(config, MEETING).state_reason is None and not _progress(config, MEETING)


def test_a_background_cycle_works_pieces_until_180_s_and_materialise_path_works_to_the_end(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """A background cycle works pieces of one recording at a time until 180 s of recording work is used; the
    piece in flight finishes and the next recording waits.  ``materialise PATH`` reads what it names to the
    end with no allowance."""
    config = config_with(tmp_path, meetings)
    for name in ("a.mp4", "b.mp4"):
        stub_recording(meetings / "meetings" / name, 30, piece_s=70.0)
    assert run(config).exit_code == 0
    (first,) = rec.converts  # the allowance is used up: the other one waits
    other = "b.mp4" if first == "a.mp4" else "a.mp4"
    assert rec.worked == [(first, 0), (first, 1), (first, 2)], "70 + 70 + 70 s: the third passes 180"
    assert _progress(config, f"meetings/{first}") == f"{3 * PIECE_MS} {30 * 60_000}"
    for name in (first, other):
        assert _row(config, f"meetings/{name}").state_reason == cycle_mod.RECORDING_WAITS
    rec.worked.clear()
    named = run(config, materialise_paths=[meetings / "meetings" / other])
    assert named.exit_code == 0 and rec.opened[-1] is None
    assert rec.worked == [(other, n) for n in range(6)], (
        "every piece of the named recording, none of the other"
    )
    assert _recording_page(config, f"meetings/{other}")[0]["status"] == "current"
    assert _row(config, f"meetings/{first}").state_reason == cycle_mod.RECORDING_WAITS
    assert _pieces_of(config, meetings / "meetings" / first).is_dir()
    assert not _pieces_of(config, meetings / "meetings" / other).exists()


def test_a_stored_piece_is_never_redone_and_three_cycles_equal_one_byte_for_byte(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """Spec 4.1: a recording read over three cycles reads each piece once, and its page is the page one pass
    makes; its piece folder goes once the page is in the converter cache."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 40, piece_s=70.0)
    for _ in range(3):
        assert run(config).exit_code == 0
    assert rec.worked == [("Contoso weekly sync.mp4", n) for n in range(8)]
    assert not _pieces_of(config, path).exists()
    over_three = _recording_page(config, MEETING)[1]
    other = tmp_path / "one-pass"
    other.mkdir()
    shutil.copytree(meetings, other / "meetings-source")
    one = config_with(other, other / "meetings-source")
    assert run(one, materialise_paths=[other / "meetings-source" / MEETING]).exit_code == 0
    assert _recording_page(one, MEETING)[1] == over_three, "the body; the front matter names the file"


def test_a_piece_that_times_out_leaves_the_recording_waiting_never_failed(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """Ruling 4: a piece past its deadline is discarded and the recording waits, with no error and no
    stub; the read is counted once, and the next cycle finishes it and clears the count."""
    config = config_with(tmp_path, meetings)
    stub_recording(meetings / MEETING, 15)
    rec.timeout_at.add(("Contoso weekly sync.mp4", 1))
    report = run(config)
    assert report.exit_code == 0 and source_report(report).errors == ()
    row = _row(config, MEETING)
    assert row.state is RowState.LIVE and row.state_reason == cycle_mod.RECORDING_WAITS
    assert row.last_verdict is Verdict.DEFERRED and _progress(config, MEETING) == f"{PIECE_MS} {15 * 60_000}"
    assert _record(config)["failed"] == {row.stable_id: 1}
    rec.timeout_at.clear()
    assert run(config).exit_code == 0
    assert _recording_page(config, MEETING)[0]["status"] == "current" and _record(config) == {}


def test_a_piece_stopped_in_two_cycles_settles_as_a_stub_on_both_paths(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """O13: a recording whose read is killed or times out in two cycles becomes a stub that says to run
    ``agentsync materialise`` on it, whether it was being read for the first time or was a ``no converter``
    stub from before recordings were read.  Its stored pieces are kept."""
    config = config_with(tmp_path, meetings)
    first = stub_recording(meetings / "meetings/first.mp4", 15)
    rec.timeout_at.add(("first.mp4", 1))
    assert (
        run(config).exit_code == 0
        and _row(config, "meetings/first.mp4").state_reason == cycle_mod.RECORDING_WAITS
    )
    assert run(config).exit_code == 0
    fm, body = _recording_page(config, "meetings/first.mp4")
    row = _row(config, "meetings/first.mp4")
    assert row.state is RowState.QUARANTINED and row.state_reason == cycle_mod._RECORDING_STOPPED
    assert fm["status"] == "unreadable" and "run agentsync materialise on it from a terminal" in body
    assert "first.mp4" not in cycle_mod._RECORDING_STOPPED and _pieces_of(config, first).is_dir()
    assert _progress(config, "meetings/first.mp4") == "" and row.last_verdict is Verdict.QUARANTINED
    # the stub path: a ``no converter`` stub made without a media helper, killed in two cycles
    monkeypatch.setattr(cycle_mod.media, "engine", lambda _convert, _cache_dir: None)
    stub_recording(meetings / "meetings/old.mp4", 15)
    assert run(config).exit_code == 0
    assert _row(config, "meetings/old.mp4").state_reason == "no converter for .mp4"
    monkeypatch.setattr(cycle_mod.media, "engine", lambda _convert, _cache_dir: rec.media)
    rec.crash_at.add(("old.mp4", 0))
    for _ in range(2):
        with pytest.raises(Crash):
            run(config)
    rec.crash_at.clear()
    converts = len(rec.converts)
    assert run(config).exit_code == 0 and len(rec.converts) == converts, "given up without a third read"
    old = _row(config, "meetings/old.mp4")
    assert old.state is RowState.QUARANTINED and old.state_reason == cycle_mod._RECORDING_STOPPED


def test_materialise_path_resumes_a_recording_given_up_after_two_kills_from_its_pieces(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """O13: ``agentsync materialise PATH`` clears the count of a recording given up after two kills and
    resumes it from the pieces stored before them."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 15)
    rec.crash_at.add(("Contoso weekly sync.mp4", 1))
    for _ in range(2):
        with pytest.raises(Crash):
            run(config)
    assert run(config).exit_code == 0
    assert _row(config, MEETING).state_reason == cycle_mod._RECORDING_STOPPED
    assert rec.worked == [("Contoso weekly sync.mp4", 0)], "piece 0 was stored before the first kill"
    rec.crash_at.clear()
    assert run(config, materialise_paths=[path]).exit_code == 0
    assert rec.worked[1:] == [("Contoso weekly sync.mp4", 1), ("Contoso weekly sync.mp4", 2)]
    row = _row(config, MEETING)
    assert row.state is RowState.LIVE and row.state_reason is None and _record(config) == {}
    assert _recording_page(config, MEETING)[0]["status"] == "current"


def test_a_cycle_killed_in_a_recording_loses_no_other_row(tmp_path: Path, meetings: Path, rec: Rec) -> None:
    """The recording pass runs after every source's queue, one recording per transaction: a kill in it
    loses nothing the queues settled, and the next cycle counts the killed read once."""
    config = config_with(tmp_path, meetings)
    assert run(config).exit_code == 0
    (meetings / "notes.md").write_text("# Notes\n\nContoso agenda, second draft\n", encoding="utf-8")
    (meetings / "plan.md").write_text("# Plan\n\nContoso launch\n", encoding="utf-8")
    stub_recording(meetings / MEETING, 60, piece_s=70.0)
    rec.crash_at.add(("Contoso weekly sync.mp4", 0))
    with pytest.raises(Crash):
        run(config)
    rows = _file_rows(config)
    assert (
        rows["notes.md"].last_verdict is Verdict.UNCHANGED
        and rows["plan.md"].last_verdict is Verdict.UNCHANGED
    )
    assert rows[MEETING].last_verdict is Verdict.DEFERRED
    rec.crash_at.clear()
    with Manifest(config.state_paths.db) as m:
        assert (
            json.loads(m.get_meta(cycle_mod._RECORDING_META + SID) or "{}")["reading"]
            == rows[MEETING].stable_id
        )
    assert run(config).exit_code == 0
    assert _record(config)["failed"] == {rows[MEETING].stable_id: 1}, "the kill, counted once"
    assert _row(config, MEETING).state_reason == cycle_mod.RECORDING_WAITS, "three of twelve pieces read"
    assert "second draft" in page(config.docs_repo, slug.mirror_rel_path(SID, "notes.md"))[1]
    assert "Contoso launch" in page(config.docs_repo, slug.mirror_rel_path(SID, "plan.md"))[1]


def test_a_media_helper_that_fails_on_everything_costs_no_recording_its_reading(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """Spec S0 rule 8: after a MediaError the cycle asks the helper for ``--version``.  When it does not
    answer, the failure is no file's: nothing is published or counted, no further recording is staged, and
    the source's report says once to run scripts/install.sh.  When it answers, the failure is the file's."""
    config = config_with(tmp_path, meetings)
    both = ("meetings/a.mp4", "meetings/b.mp4")
    for rel in both:
        stub_recording(meetings / rel, 10, piece_s=10.0)
    rec.media_fails, rec.media.up = True, False
    report = run(config)
    assert report.exit_code == 0 and len(rec.converts) == 1 and rec.media.asked == 1
    assert source_report(report).alarms.count(cycle_mod._MEDIA_DOWN) == 1
    assert "scripts/install.sh" in cycle_mod._MEDIA_DOWN
    for rel in both:
        row = _row(config, rel)
        assert row.state_reason == cycle_mod.RECORDING_WAITS and row.last_verdict is Verdict.DEFERRED
        assert not (config.docs_repo / slug.mirror_rel_path(SID, rel)).exists()
    assert _record(config) == {}, "no read was counted against either file"
    rec.media.up = True  # the helper answers: the failure is the file's, and counts against it
    assert run(config).exit_code == 0
    assert all(_row(config, rel).state_reason == "no converter for .mp4" for rel in both)
    assert _record(config)["failed"] == {_row(config, rel).stable_id: 1 for rel in both}
    rec.media_fails = False
    assert run(config).exit_code == 0
    assert all(_recording_page(config, rel)[0]["status"] == "current" for rel in both)
    assert _record(config) == {}


def test_no_recording_is_staged_once_either_helper_is_down(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A recording is read by both helpers: once the OCR helper failed on everything in this cycle, or the
    media helper stopped answering, the recording pass stages no recording (no copy, no hash)."""
    config = config_with(tmp_path, meetings)
    stub_recording(meetings / MEETING, 10)
    failing = fake_engine(tmp_path / "failing-ocr", fail=True)
    monkeypatch.setattr(cycle_mod.ocr, "engine", lambda _convert, _cache_dir: failing)
    picture(meetings / "shots" / "Contoso board.png", "Roadmap")  # the OCR helper fails on it, then a blank
    fetched = _fetches(monkeypatch)
    assert run(config).exit_code == 0
    assert "shots/Contoso board.png" in fetched and MEETING not in fetched and rec.converts == []
    assert _row(config, MEETING).state_reason == cycle_mod.RECORDING_WAITS and _record(config) == {}
    _use_ocr(monkeypatch, tmp_path)
    rec.media_fails, rec.media.up = True, False
    stub_recording(meetings / "meetings/second.mp4", 10)
    fetched.clear()
    assert run(config).exit_code == 0
    assert fetched.count(MEETING) + fetched.count("meetings/second.mp4") == 1, "the second is never staged"


def test_a_no_converter_mp4_stub_is_read_in_the_recording_pass(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A recording stubbed ``no converter for .mp4`` before this Mac read recordings is read once the
    converter is registered, by the recording pass of a background cycle, never by an interactive sync or
    the source's re-read pass."""
    config = config_with(tmp_path, meetings)
    monkeypatch.setattr(cycle_mod.media, "engine", lambda _convert, _cache_dir: None)
    stub_recording(meetings / MEETING, 10)
    assert run(config).exit_code == 0
    row = _row(config, MEETING)
    assert row.state is RowState.REFUSED and row.state_reason == "no converter for .mp4"
    monkeypatch.setattr(cycle_mod.media, "engine", lambda _convert, _cache_dir: rec.media)
    assert run(config, mode=None).exit_code == 0 and rec.converts == []
    assert run(config).exit_code == 0 and rec.converts == ["Contoso weekly sync.mp4"]
    row = _row(config, MEETING)
    assert row.state is RowState.LIVE and row.state_reason is None
    assert _recording_page(config, MEETING)[0]["status"] == "current"


SHOTS_SOURCE = """
[[source]]
id = "shots"
kind = "local"
path = "{path}"
sentinel = "README.txt"
"""


def test_a_recording_takes_no_ocr_time_from_a_later_source(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Recordings are read after every source: a recording in the first source that uses the cycle's whole
    OCR time leaves the screenshots of a later source theirs, all read by OCR in the same cycle."""
    shots = tmp_path / "shots-source"
    (shots / "board").mkdir(parents=True)
    (shots / "README.txt").write_text("sentinel\n", encoding="utf-8")
    for n in range(3):
        picture(shots / "board" / f"Contoso slide {n}.png", f"Milestone {n}")
    config = config_with(tmp_path, meetings, SHOTS_SOURCE.format(path=shots))
    stub_recording(meetings / MEETING, 10, ocr_s=10 * cycle_mod._OCR_BUDGET_S)
    assert run(config).exit_code == 0 and rec.converts == ["Contoso weekly sync.mp4"]
    for n in range(3):
        fm, body = page(config.docs_repo, slug.mirror_rel_path("shots", f"board/Contoso slide {n}.png"))
        assert fm["converter"].startswith("image-ocr@") and f"Milestone {n}" in body
    assert _run_record(config).get("ocr_deferred") is None


def test_screenshots_beside_a_pending_recording_keep_their_ocr_time(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """A recording read in part, with screenshots beside it: the screenshots are read by OCR in the cycle
    that sees them, before the recording's pieces take any of the cycle's OCR time."""
    config = config_with(tmp_path, meetings)
    stub_recording(meetings / MEETING, 60, piece_s=70.0, ocr_s=cycle_mod._OCR_BUDGET_S)
    assert run(config).exit_code == 0 and _row(config, MEETING).state_reason == cycle_mod.RECORDING_WAITS
    for n in range(3):
        picture(meetings / "meetings" / f"Contoso screenshot {n}.png", f"Action item {n}")
    assert run(config).exit_code == 0 and rec.converts[-1] == "Contoso weekly sync.mp4"
    for n in range(3):
        fm, body = page(config.docs_repo, slug.mirror_rel_path(SID, f"meetings/Contoso screenshot {n}.png"))
        assert fm["converter"].startswith("image-ocr@") and f"Action item {n}" in body
    assert _run_record(config).get("ocr_deferred") is None


def test_a_graph_document_in_a_later_source_converts_in_the_cycle_that_reads_a_recording(
    tmp_path: Path, meetings: Path, fixture_files: dict[str, Path], rec: Rec
) -> None:
    """A Graph document whose converter reads with OCR waits once the cycle's OCR time is used.  A
    recording in an earlier source that uses all of it is read after the Graph source, so the document
    converts in the same cycle."""
    config = config_with(tmp_path, meetings, GRAPH_SOURCE)
    drive = FakeDrive({"I1": ("Contoso brief.pdf", fixture_files["sample.pdf"].read_bytes())})
    client = GraphClient(
        FakeTokens(),
        user_agent="NONISV|test|agentsync/0",
        transport=httpx.MockTransport(drive.handler),
        sleep=lambda _s: None,
    )
    stub_recording(meetings / MEETING, 10, ocr_s=10 * cycle_mod._OCR_BUDGET_S)
    try:
        assert run(config, client=client).exit_code == 0
    finally:
        client.close()
    assert rec.converts == ["Contoso weekly sync.mp4"]
    (brief,) = _file_rows(config, "drive").values()
    assert brief.last_verdict is Verdict.UNCHANGED and brief.state is RowState.LIVE
    assert _run_record(config).get("ocr_deferred") is None


AWS_KEY = "AKIA" + "ABCDEFGHIJKLMNOP"


def test_a_credential_on_screen_stubs_the_recording_with_its_keyframes(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """O7: a credential read off the screen makes the whole recording a ``contains a credential`` stub, and
    its keyframes and full text go with the page."""
    config = config_with(tmp_path, meetings)
    keyframe = (b"\xff\xd8\xff\xe0" + b"frame" * 20 + b"\xff\xd9").hex()
    sidecars = {"t000010.jpg": f"hex:{keyframe}", "full-text.txt": f"aws key {AWS_KEY}\n"}
    stub_recording(meetings / MEETING, 10, text=f"aws_access_key_id = {AWS_KEY}", sidecars=sidecars)
    assert run(config).exit_code == 0
    fm, body = _recording_page(config, MEETING)
    assert fm["reason"] == "contains a credential" and AWS_KEY not in body + str(fm)
    page_path = config.docs_repo / slug.mirror_rel_path(SID, MEETING)
    assert not [p for p in page_path.parent.rglob("*") if p.suffix in (".jpg", ".txt")]
    assert AWS_KEY not in git(config.docs_repo, "log", "-p", "--all")


def test_a_jpeg_sidecar_does_not_stub_its_recording(tmp_path: Path, meetings: Path, rec: Rec) -> None:
    """S10 rule 2: the secret scan reads pages and text sidecars only.  Bytes in a keyframe that look like a
    credential when decoded as text do not stub the recording."""
    config = config_with(tmp_path, meetings)
    keyframe = (b"\xff\xd8\xff\xe0" + f"PWd={AWS_KEY}".encode() + b"\xff\xd9").hex()
    stub_recording(
        meetings / MEETING, 10, sidecars={"t000010.jpg": f"hex:{keyframe}", "full-text.txt": "x\n"}
    )
    report = run(config)
    assert report.exit_code == 0 and report.commit_sha is not None
    fm, _body = _recording_page(config, MEETING)
    assert fm["status"] == "current" and _row(config, MEETING).state is RowState.LIVE


def test_gc_keeps_the_pieces_of_a_pending_recording_and_purge_removes_them(
    tmp_path: Path, meetings: Path, rec: Rec
) -> None:
    """Spec 4.1: reconcile removes the piece folder of a hash no row waits on and keeps the pieces of a
    recording still being read (the converter cache's own gc never reaches the store).  Purge removing a
    purged recording's folder is governance's (``tests/test_governance.py``)."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 60, piece_s=70.0)
    assert run(config).exit_code == 0 and _pieces_of(config, path).is_dir()
    orphan = config.cache_dir / "recordings" / ("0" * 64)
    orphan.mkdir(parents=True)
    (orphan / "0000-k1").write_bytes(b"left by a recording since deleted")
    assert run(config, mode=CycleMode.RECONCILE).exit_code == 0
    assert not orphan.exists() and _pieces_of(config, path).is_dir()
    assert _row(config, MEETING).state_reason == cycle_mod.RECORDING_WAITS


def test_a_policy_change_marks_no_recording(tmp_path: Path, meetings: Path, rec: Rec) -> None:
    """A recording's label cannot be read (ruling 2): its suffixes are not label-capable, so a changed
    ``[policy]`` re-screens no recording and nothing is read again for it."""
    assert not any(ext in cycle_mod._LABEL_CAPABLE for ext in RecordingConverter.extensions)
    stub_recording(meetings / MEETING, 10)
    assert run(config_with(tmp_path, meetings)).exit_code == 0 and len(rec.converts) == 1
    policy = '\n[policy]\nexclude_label_ids = ["00000000-0000-4000-8000-00000000c0de"]\n'
    config = config_with(tmp_path, meetings, policy)
    assert run(config).exit_code == 0 and len(rec.converts) == 1
    row = _row(config, MEETING)
    assert row.last_verdict is Verdict.UNCHANGED and row.state is RowState.LIVE


def _online_only(monkeypatch: pytest.MonkeyPatch, *paths: Path) -> set[int]:
    """Mock SF_DATALESS on these files (a test has no File Provider).  A file ``materialise`` reads is on this
    Mac from then on, as after a download; the returned set holds the inodes still online-only."""
    online = {p.stat().st_ino for p in paths}
    real_dataless, real_materialise = materialise.is_dataless, al.materialise

    def is_dataless(st: Any) -> bool:
        return st.st_ino in online or real_dataless(st)

    def download(src: Path, dest: Path, budget: Any, **kwargs: Any) -> Any:
        result = real_materialise(src, dest, budget, **kwargs)
        online.discard(src.stat().st_ino)
        return result

    monkeypatch.setattr(materialise, "is_dataless", is_dataless)
    monkeypatch.setattr(al, "is_dataless", is_dataless)
    monkeypatch.setattr(al, "materialise", download)
    return online


def test_an_online_only_recording_is_downloaded_by_reconcile_outside_the_document_budget(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Ruling 1: a reconcile job downloads an online-only recording under an allowance of its own.  A run
    with no document download budget still reads it, charges nothing to the source's budget, and leaves the
    online-only document beside it for a run whose budget allows it."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 10, piece_s=10.0)
    brief = meetings / "brief.md"
    brief.write_text("# Brief\n\nContoso launch plan\n", encoding="utf-8")
    online = _online_only(monkeypatch, path, brief)
    report = run(config, mode=CycleMode.RECONCILE, budget_bytes=0)
    assert report.exit_code == 0 and rec.converts == ["Contoso weekly sync.mp4"]
    assert _recording_page(config, MEETING)[0]["status"] == "current"
    assert source_report(report).materialised_bytes == 0, "the recording is not charged to the budget"
    assert online == {brief.stat().st_ino} and _row(config, "brief.md").last_verdict is Verdict.DEFERRED


def test_a_poll_cycle_and_an_interactive_sync_download_no_recording(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Rule 3 and L9: a poll job's watchdog is too short for a download, and a tool's timeout may end an
    interactive sync, so neither downloads a recording; the reconcile job does."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 10, piece_s=10.0)
    online = _online_only(monkeypatch, path)
    fetched = _fetches(monkeypatch)
    assert run(config).exit_code == 0 and run(config, mode=None).exit_code == 0
    assert MEETING not in fetched and rec.converts == [] and online == {path.stat().st_ino}
    row = _row(config, MEETING)
    assert row.dataless and row.state_reason == cycle_mod.RECORDING_WAITS
    assert row.last_verdict is Verdict.DEFERRED
    assert run(config, mode=CycleMode.RECONCILE).exit_code == 0
    assert fetched.count(MEETING) == 1 and _recording_page(config, MEETING)[0]["status"] == "current"


def test_one_recording_download_per_cycle_newest_first_and_none_past_four_gib(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """At most one recording is downloaded per cycle, the newest first; one over the size limit is never
    downloaded and is counted as one that could not be (``HYDRATION_REFUSED``)."""
    monkeypatch.setattr(cycle_mod, "_RECORDING_MAX_BYTES", 2000)
    config = config_with(tmp_path, meetings)
    now = time.time()
    paths = {
        "old": stub_recording(meetings / "meetings/old.mp4", 10, piece_s=10.0),
        "new": stub_recording(meetings / "meetings/new.mp4", 10, piece_s=10.0),
        "huge": stub_recording(meetings / "meetings/huge.mp4", 10, piece_s=10.0, pad="x" * 3000),
    }
    for age, key in ((3000, "old"), (1000, "new"), (0, "huge")):
        os.utime(paths[key], (now - age, now - age))
    _online_only(monkeypatch, *paths.values())
    fetched = _fetches(monkeypatch)
    assert run(config, mode=CycleMode.RECONCILE).exit_code == 0
    assert [f for f in fetched if f.endswith(".mp4")] == ["meetings/new.mp4"]
    assert _row(config, "meetings/huge.mp4").state_reason == cycle_mod.HYDRATION_REFUSED
    assert _row(config, "meetings/old.mp4").state_reason == cycle_mod.RECORDING_WAITS
    assert run(config, mode=CycleMode.RECONCILE).exit_code == 0
    assert [f for f in fetched if f.endswith(".mp4")] == ["meetings/new.mp4", "meetings/old.mp4"]
    assert _row(config, "meetings/huge.mp4").state_reason == cycle_mod.HYDRATION_REFUSED


def test_a_recording_download_waits_for_twice_its_size_and_five_gib_free(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A download starts only while the volume keeps twice the recording's size and 5 GiB free."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 10, piece_s=10.0)
    _online_only(monkeypatch, path)
    fetched = _fetches(monkeypatch)
    need = 2 * path.stat().st_size + 5 * 1024**3
    monkeypatch.setattr(cycle_mod, "_free_bytes", lambda _path: need - 1)
    assert run(config, mode=CycleMode.RECONCILE).exit_code == 0 and MEETING not in fetched
    assert _row(config, MEETING).state_reason == cycle_mod.RECORDING_WAITS
    monkeypatch.setattr(cycle_mod, "_free_bytes", lambda _path: need)
    assert run(config, mode=CycleMode.RECONCILE).exit_code == 0 and fetched.count(MEETING) == 1


def test_a_download_that_fails_with_errno_89_or_passes_its_deadline_is_a_counted_finder_note_never_rule_three(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A download that passes its deadline or that the OS refuses (``errno 89``) sets ``HYDRATION_REFUSED``
    (the counted Finder note), never an error or a stub.  It is tried once more in a later reconcile, then
    only when the file's stat changes; a ``materialise PATH`` run that names it says what to do in Finder."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 10, piece_s=10.0)
    _online_only(monkeypatch, path)
    copy, downloads = al.materialise, list[str]()

    def slow(src: Path, dest: Path, budget: Any, **kwargs: Any) -> Any:
        if src.suffix != ".mp4":
            return copy(src, dest, budget, **kwargs)
        downloads.append(src.name)
        time.sleep(0.5)
        raise AssertionError("past the deadline: nobody waits for this")

    def canceled(src: Path, dest: Path, budget: Any, **kwargs: Any) -> Any:
        if src.suffix != ".mp4":
            return copy(src, dest, budget, **kwargs)
        downloads.append(src.name)
        raise DatalessRefusedError(str(src), 89, "File Provider canceled the read (ECANCELED)")

    monkeypatch.setattr(cycle_mod, "_RECORDING_DEADLINE_S", 0.05)
    monkeypatch.setattr(cycle_mod, "_RECORDING_FLOOR_BPS", 10**12)
    monkeypatch.setattr(al, "materialise", slow)
    report = run(config, mode=CycleMode.RECONCILE)
    row = _row(config, MEETING)
    assert report.exit_code == 0 and source_report(report).errors == () and downloads == [path.name]
    assert row.state is RowState.DATALESS and row.state_reason == cycle_mod.HYDRATION_REFUSED
    assert row.last_verdict is Verdict.DEFERRED and rec.converts == []
    monkeypatch.setattr(al, "materialise", canceled)
    for _ in range(2):
        assert run(config, mode=CycleMode.RECONCILE).exit_code == 0
    assert len(downloads) == 2, "tried once more, then not while its stat is the same"
    assert _row(config, MEETING).state_reason == cycle_mod.HYDRATION_REFUSED
    later = time.time() + 60
    os.utime(path, (later, later))
    assert run(config, mode=CycleMode.RECONCILE).exit_code == 0 and len(downloads) == 3
    named = run(config, materialise_paths=[path])
    assert len(downloads) == 4 and cycle_mod._RECORDING_NOT_DOWNLOADED.format(path=MEETING) in (
        source_report(named).alarms
    )
    assert _row(config, MEETING).state is RowState.DATALESS, "never quarantined, never a stub"


def test_a_document_beside_an_online_only_recording_is_never_held_back(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An online-only recording never takes the source's document download budget: an online-only
    document beside it downloads in the same cycle, whatever the order of their rows."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 10, piece_s=10.0)
    brief = meetings / "brief.md"
    brief.write_text("# Brief\n\nContoso launch plan\n", encoding="utf-8")
    _online_only(monkeypatch, path, brief)
    assert run(config, budget_bytes=brief.stat().st_size).exit_code == 0
    assert "Contoso launch plan" in page(config.docs_repo, slug.mirror_rel_path(SID, "brief.md"))[1]
    assert _row(config, MEETING).state_reason == cycle_mod.RECORDING_WAITS and rec.converts == []


def test_a_recording_downloaded_with_its_stat_unchanged_is_read_by_the_cycle_that_sees_it(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Downloaded in Finder (only its online-only flag changes), a recording that waited, or a ``no
    converter`` stub made before recordings were read, is read by the next cycle that lists it on this Mac."""
    config = config_with(tmp_path, meetings)
    monkeypatch.setattr(cycle_mod.media, "engine", lambda _convert, _cache_dir: None)
    old = stub_recording(meetings / "meetings/old.mp4", 10, piece_s=10.0)
    online = _online_only(monkeypatch, old)
    assert run(config).exit_code == 0
    assert _row(config, "meetings/old.mp4").state_reason == "no converter for .mp4"
    monkeypatch.setattr(cycle_mod.media, "engine", lambda _convert, _cache_dir: rec.media)
    new = stub_recording(meetings / "meetings/new.mp4", 10, piece_s=10.0)
    online.add(new.stat().st_ino)
    assert run(config).exit_code == 0 and rec.converts == []
    assert _row(config, "meetings/new.mp4").state_reason == cycle_mod.RECORDING_WAITS
    online.clear()  # Download Now in Finder: same size, same modification time
    assert run(config).exit_code == 0 and sorted(rec.converts) == ["new.mp4", "old.mp4"]
    for rel in ("meetings/old.mp4", "meetings/new.mp4"):
        assert _recording_page(config, rel)[0]["status"] == "current"


def test_a_recording_read_while_local_keeps_its_pages_once_online_only(
    tmp_path: Path, meetings: Path, rec: Rec, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A recording read while it was on this Mac keeps its page once OneDrive evicts it: no reconcile and
    no ``materialise PATH`` downloads it again to read the same bytes."""
    config = config_with(tmp_path, meetings)
    path = stub_recording(meetings / MEETING, 10, piece_s=10.0)
    assert run(config).exit_code == 0 and rec.converts == ["Contoso weekly sync.mp4"]
    before = (config.docs_repo / slug.mirror_rel_path(SID, MEETING)).read_bytes()
    _online_only(monkeypatch, path)
    fetched = _fetches(monkeypatch)
    assert run(config, mode=CycleMode.RECONCILE).exit_code == 0
    assert run(config, materialise_paths=[path]).exit_code == 0
    assert MEETING not in fetched and rec.converts == ["Contoso weekly sync.mp4"]
    assert (config.docs_repo / slug.mirror_rel_path(SID, MEETING)).read_bytes() == before
    assert _row(config, MEETING).dataless and _row(config, MEETING).state_reason is None
