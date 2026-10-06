"""Determinism: double conversion per converter, rebuild/resave stability, name independence, cache hits."""

from __future__ import annotations

import hashlib
import io
import shutil
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from agentsync.config import ConvertConfig
from agentsync.convert import ConverterCache, Registry, convert_file, double_conversion_differs
from agentsync.convert.canonical import canonical_hash
from agentsync.model import ConversionResult, ConversionStatus
from fixtures.make_fixtures import make_fixtures
from test_convert_builders import (
    PdfPicture,
    build_commented_pdf,
    build_docx_image,
    build_docx_merged,
    build_pdf,
    build_picture_pdf,
    build_pptx_rich,
    build_xlsx_rich,
    eml_bytes,
    page_picture,
    pandoc_build,
    shade_engine,
    teams_doc,
    teams_msg,
)
from test_convert_image import picture, text_png
from test_ocr import calls, fake_engine

CFG = ConvertConfig()


@pytest.fixture(scope="module")
def registry() -> Registry:
    return Registry.default(CFG)


@pytest.fixture(scope="module")
def extra_inputs(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Path]:
    """Generated inputs beyond the shared fixtures, covering every converter and its harder paths."""
    d = tmp_path_factory.mktemp("determinism")
    out: dict[str, Path] = {
        "image.docx": build_docx_image(d),
        "merged.docx": build_docx_merged(d),
        "rich.xlsx": build_xlsx_rich(d / "rich.xlsx", rows=40),
        "rich.pptx": build_pptx_rich(d / "rich.pptx"),
        "two.pdf": build_pdf(d / "two.pdf", [["alpha", "beta"], []]),
        "commented.pdf": build_commented_pdf(d / "commented.pdf"),
        "minutes.odt": pandoc_build("# Minutes\n\ntext\n", "markdown", d / "minutes.odt"),
        "minutes.rtf": pandoc_build("# Minutes\n\ntext\n", "markdown", d / "minutes.rtf"),
    }
    (d / "html.eml").write_bytes(eml_bytes(plain=None, html="<p>Hi <b>there</b></p>"))
    out["html.eml"] = d / "html.eml"
    teams = teams_doc(
        [teams_msg(str(i), f"2026-09-{i + 1:02d}T10:00:00Z", f"<p>message <i>{i}</i>") for i in range(12)]
    )
    (d / "busy.teams.json").write_text(teams, encoding="utf-8")
    out["busy.teams.json"] = d / "busy.teams.json"
    (d / "notes.yaml").write_text("a: 1\nb: [x, y]\n", encoding="utf-8")
    out["notes.yaml"] = d / "notes.yaml"
    (d / "table.tsv").write_text("k\tv\na\t1\n", encoding="utf-8")
    out["table.tsv"] = d / "table.tsv"
    return out


def _all_inputs(fixture_files: dict[str, Path], extra_inputs: dict[str, Path]) -> dict[str, Path]:
    return {**fixture_files, **extra_inputs}


def _convert(src: Path, name: str, registry: Registry, cache_root: Path) -> ConversionResult:
    return convert_file(
        src,
        name=name,
        content_sha256=hashlib.sha256(src.read_bytes()).hexdigest(),
        canonical_sha256=canonical_hash(src, suffix=Path(name).suffix).sha256,
        registry=registry,
        cache=ConverterCache(cache_root),
    )


def test_every_converter_is_covered(
    registry: Registry, fixture_files: dict[str, Path], extra_inputs: dict[str, Path]
) -> None:
    used = {registry.for_name(n).converter_id for n in _all_inputs(fixture_files, extra_inputs)}  # type: ignore[union-attr]
    assert used == {c.converter_id for c in registry.converters()}


def test_double_conversion_is_identical_for_every_input(
    registry: Registry, fixture_files: dict[str, Path], extra_inputs: dict[str, Path]
) -> None:
    differs = [
        name
        for name, path in sorted(_all_inputs(fixture_files, extra_inputs).items())
        if double_conversion_differs(path, name=name, registry=registry)
    ]
    assert differs == []


def test_convert_file_twice_with_cold_caches_is_byte_identical(
    registry: Registry, fixture_files: dict[str, Path], extra_inputs: dict[str, Path], tmp_path: Path
) -> None:
    for name, path in sorted(_all_inputs(fixture_files, extra_inputs).items()):
        a = _convert(path, name, registry, tmp_path / f"a-{name}")
        b = _convert(path, name, registry, tmp_path / f"b-{name}")
        assert a.status is ConversionStatus.OK, (name, a.reason)
        assert a == b, name
        warm = _convert(path, name, registry, tmp_path / f"a-{name}")
        assert warm.from_cache and warm.units == a.units, name


def test_rebuilt_fixtures_render_identically(registry: Registry, tmp_path: Path) -> None:
    """Container metadata differs between two builds (zip stamps, docProps); the rendered units must not."""
    one = make_fixtures(tmp_path / "one")
    two = make_fixtures(tmp_path / "two")
    for name in sorted(one):
        conv = registry.for_name(name)
        assert conv is not None
        h2_one = [(u.unit_id, u.rendered_sha256) for u in conv.convert(one[name], name=name)]
        h2_two = [(u.unit_id, u.rendered_sha256) for u in conv.convert(two[name], name=name)]
        assert h2_one == h2_two, name


def test_office_style_resave_is_free(
    registry: Registry, fixture_files: dict[str, Path], tmp_path: Path
) -> None:
    """sample-resaved.xlsx: new bytes, same H1: the cache serves it, every H2 is equal (design 4.1 #2)."""
    a = _convert(fixture_files["sample.xlsx"], "sample.xlsx", registry, tmp_path / "c")
    b = _convert(fixture_files["sample-resaved.xlsx"], "sample.xlsx", registry, tmp_path / "c")
    assert a.content_sha256 != b.content_sha256
    assert b.from_cache and a.action_key == b.action_key
    assert [u.rendered_sha256 for u in a.units] == [u.rendered_sha256 for u in b.units]


def test_output_depends_on_bytes_and_suffix_only(
    registry: Registry, fixture_files: dict[str, Path], tmp_path: Path
) -> None:
    """The action key has no name component and a rename keeps bodies, so the name must not leak in."""
    for name, path in sorted(fixture_files.items()):
        suffix = ".teams.json" if name.endswith(".teams.json") else Path(name).suffix
        renamed = tmp_path / f"Completely Different Name{suffix}"
        shutil.copyfile(path, renamed)
        conv = registry.for_name(name)
        assert conv is not None
        assert conv.convert(path, name=name) == conv.convert(renamed, name=renamed.name), name


def test_an_image_converts_the_same_twice_and_under_any_name(
    fixture_files: dict[str, Path], extra_inputs: dict[str, Path], tmp_path: Path
) -> None:
    """The ninth converter, there only with an OCR engine, by the rules of the other eight: the same bytes
    under a second name are a cache hit, so the page must hold neither name."""
    engine = fake_engine(tmp_path / "bin")
    registry = Registry.default(CFG, ocr=engine)
    first = picture(tmp_path / "a" / "Contoso Roadmap.png", "Milestones", "Beta in spring")
    second = tmp_path / "b" / "Fabrikam Org Chart.png"
    second.parent.mkdir()
    shutil.copyfile(first, second)
    inputs = {*_all_inputs(fixture_files, extra_inputs), first.name}
    assert {registry.for_name(n).converter_id for n in inputs} == {  # type: ignore[union-attr]
        c.converter_id for c in registry.converters()
    }
    assert not double_conversion_differs(first, name=first.name, registry=registry)
    a = _convert(first, first.name, registry, tmp_path / "cache")
    b = _convert(second, second.name, registry, tmp_path / "cache")
    assert a.status is ConversionStatus.OK and a.converter_id == "image-ocr" and not a.from_cache
    assert b.from_cache and b.action_key == a.action_key and b.units == a.units
    assert _convert(second, second.name, registry, tmp_path / "cold") == a, "a cold cache gives the same page"
    (unit,) = b.units
    for word in ("Contoso", "Roadmap", "Fabrikam", "Org Chart", ".png"):
        assert word not in unit.body + unit.title + unit.summary, word
    assert "Milestones\nBeta in spring\n" in unit.body
    assert len(calls(engine.helper)) == 4, "two for the double conversion, one each for the two cold caches"


def test_teams_messages_render_independently_of_neighbours(registry: Registry, tmp_path: Path) -> None:
    def month(second_body: str) -> str:
        return teams_doc(
            [
                teams_msg("1", "2026-09-01T10:00:00Z", "<p>stable <b>first</b>"),
                teams_msg("2", "2026-09-02T10:00:00Z", second_body),
                teams_msg("3", "2026-09-03T10:00:00Z", "<p>stable third</p>"),
            ]
        )

    conv = registry.for_name("x.teams.json")
    assert conv is not None
    bodies = []
    for i, second in enumerate(["<p>plain</p>", "<p>unclosed <i>italic <table><tr><td>cell"]):
        p = tmp_path / f"{i}.teams.json"
        p.write_text(month(second), encoding="utf-8")
        bodies.append(conv.convert(p, name=p.name)[0].body)
    for text in ("stable **first**", "stable third"):
        assert text in bodies[0] and text in bodies[1]
    first_a = bodies[0].split("### 2026-09-02")[0]
    first_b = bodies[1].split("### 2026-09-02")[0]
    assert first_a == first_b


def test_options_change_moves_the_key(fixture_files: dict[str, Path], tmp_path: Path) -> None:
    src = fixture_files["sample.csv"]
    a = _convert(src, "sample.csv", Registry.default(CFG), tmp_path / "c")
    b = _convert(src, "sample.csv", Registry.default(replace(CFG, max_rows_per_sheet=1)), tmp_path / "c")
    assert a.options_hash != b.options_hash and a.action_key != b.action_key
    assert not b.from_cache


def test_a_pdf_read_by_ocr_converts_the_same_twice_and_under_any_name(tmp_path: Path) -> None:
    """With an engine a PDF's pages hold what OCR read: the same bytes still give the same page, whatever
    the file is called, and the page names neither file."""
    said = {90: ["Delivery note", "Pallets | 14"], 40: ["Gate B"]}
    engine = shade_engine(tmp_path / "bin", said)
    registry = Registry.default(CFG, ocr=engine)
    text = ["A page with a text layer and a picture below it"]
    pages = [([], [page_picture(90)]), (text, [PdfPicture(40, "96 0 0 64 150 20")])]
    (tmp_path / "a").mkdir()
    first = build_picture_pdf(tmp_path / "a" / "Contoso Scan.pdf", pages)
    second = tmp_path / "b" / "Fabrikam Copy.pdf"
    second.parent.mkdir()
    shutil.copyfile(first, second)
    assert not double_conversion_differs(first, name=first.name, registry=registry)
    a = _convert(first, first.name, registry, tmp_path / "cache")
    b = _convert(second, second.name, registry, tmp_path / "cache")
    assert a.status is ConversionStatus.OK and a.converter_id == "pdf-pypdfium2" and not a.from_cache
    assert a.converter_version.endswith("+ocr-paper-vision-r2-h0.3.0-l1")
    assert b.from_cache and b.action_key == a.action_key and b.units == a.units
    assert _convert(second, second.name, registry, tmp_path / "cold") == a, "a cold cache gives the same page"
    (unit,) = a.units
    for word in ("Contoso", "Scan", "Fabrikam", "Copy", ".pdf"):
        assert word not in unit.body + unit.title + unit.summary, word
    assert "Delivery note\nPallets | 14\n" in unit.body and "\nGate B\n" in unit.body


def test_a_deck_read_by_ocr_converts_the_same_twice_and_under_any_name(tmp_path: Path) -> None:
    """With an engine a deck's page holds what OCR read in its pictures: the same bytes still give the same
    page, whatever the file is called, and the page names neither file."""
    engine = fake_engine(tmp_path / "bin")
    registry = Registry.default(CFG, ocr=engine)
    prs = Presentation()
    shapes = prs.slides.add_slide(prs.slide_layouts[6]).shapes
    shapes.add_picture(io.BytesIO(text_png("Delivery note", "Pallets | 14")), Inches(1), Inches(1))
    (tmp_path / "a").mkdir()
    first = tmp_path / "a" / "Contoso Review.pptx"
    prs.save(str(first))
    second = tmp_path / "b" / "Fabrikam Copy.pptx"
    second.parent.mkdir()
    shutil.copyfile(first, second)
    assert not double_conversion_differs(first, name=first.name, registry=registry)
    a = _convert(first, first.name, registry, tmp_path / "cache")
    b = _convert(second, second.name, registry, tmp_path / "cache")
    assert a.status is ConversionStatus.OK and a.converter_id == "pptx-python-pptx" and not a.from_cache
    assert a.converter_version.endswith("+ocr-paper-vision-r2-h0.3.0-l1")
    assert b.from_cache and b.action_key == a.action_key and b.units == a.units
    assert _convert(second, second.name, registry, tmp_path / "cold") == a, "a cold cache gives the same page"
    (unit,) = a.units
    for word in ("Contoso", "Review", "Fabrikam", "Copy", ".pptx"):
        assert word not in unit.body + unit.title + unit.summary, word
    assert "\nDelivery note\nPallets | 14\n" in unit.body


@pytest.mark.parametrize("suffix", [".docx", ".odt"])
def test_a_word_document_read_by_ocr_converts_the_same_twice_and_under_any_name(
    tmp_path: Path, suffix: str
) -> None:
    """With an engine the page holds what OCR read in the file's pictures: the same bytes still give the
    same page, whatever the file is called, and the page names neither file."""
    engine = fake_engine(tmp_path / "bin")
    registry = Registry.default(CFG, ocr=engine)
    (tmp_path / "shot.png").write_bytes(text_png("Delivery note", "Pallets | 14"))
    (tmp_path / "a").mkdir()
    first = tmp_path / "a" / f"Contoso Review{suffix}"
    pandoc_build("# Minutes\n\n![shot](shot.png)\n", "markdown", first, cwd=tmp_path)
    second = tmp_path / "b" / f"Fabrikam Copy{suffix}"
    second.parent.mkdir()
    shutil.copyfile(first, second)
    assert not double_conversion_differs(first, name=first.name, registry=registry)
    a = _convert(first, first.name, registry, tmp_path / "cache")
    b = _convert(second, second.name, registry, tmp_path / "cache")
    assert a.status is ConversionStatus.OK and a.converter_id == "pandoc-gfm" and not a.from_cache
    assert a.converter_version.endswith("+ocr-paper-vision-r2-h0.3.0-l1")
    assert b.from_cache and b.action_key == a.action_key and b.units == a.units
    assert _convert(second, second.name, registry, tmp_path / "cold") == a, "a cold cache gives the same page"
    (unit,) = a.units
    for word in ("Contoso", "Review", "Fabrikam", "Copy", suffix):
        assert word not in unit.body + unit.title + unit.summary, word
    assert "\n\nDelivery note\\\nPallets \\| 14\n" in unit.body
