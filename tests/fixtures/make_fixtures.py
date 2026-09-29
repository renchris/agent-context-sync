"""Build small, real sample files for every converter at test time (nothing binary is committed).

Usage: ``python tests/fixtures/make_fixtures.py OUTDIR`` or ``make_fixtures(Path)`` from conftest.
Every builder is deterministic in content; container metadata (zip timestamps, docProps) may differ per build,
which is exactly what the H1 canonical-hash tests need (see ``sample-resaved.xlsx``).
"""

from __future__ import annotations

import datetime as dt
import json
import re
import sys
import zipfile
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

import pypandoc
from openpyxl import Workbook
from pptx import Presentation

MARKDOWN_SOURCE = """\
# Quarterly Plan

Intro paragraph with the term **purchase order** and an accent: café.

## Budget

| Region | Spend | Forecast |
|--------|------:|---------:|
| North  |   100 |      120 |
| South  |    80 |       75 |

- first bullet
- second bullet
"""


def build_docx(dest: Path) -> Path:
    """Write sample.docx via pandoc (pypandoc_binary), from MARKDOWN_SOURCE."""
    out = dest / "sample.docx"
    pypandoc.convert_text(MARKDOWN_SOURCE, "docx", format="md", outputfile=str(out))
    return out


def _workbook(modified: dt.datetime) -> object:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Q3 Budget"
    ws.append(["Region", "Spend", "Forecast", "Variance"])
    ws.append(["North", 100, 120, "=C2-B2"])
    ws.append(["South", 80, 75, "=C3-B3"])
    ws.append(["Total", "=SUM(B2:B3)", "=SUM(C2:C3)", None])
    ws.merge_cells("A6:D6")
    ws["A6"] = "Merged note across four columns"
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Key", "Value"])
    ws2.append(["owner", "finance"])
    ws2.append(["empty", None])
    ws2.append(["ratio", 0.25])
    wb.properties.creator = "agentsync-fixtures"
    wb.properties.created = dt.datetime(2026, 9, 1, 12, 0, 0)
    wb.properties.modified = modified
    return wb


def build_xlsx(dest: Path) -> tuple[Path, Path]:
    """Write sample.xlsx and sample-resaved.xlsx: identical parts except docProps/core.xml, new zip stamps.

    The re-save is simulated at the container level (openpyxl stamps ``modified`` with the clock on save, so
    two saves in one second are byte-identical): every entry's zip mod-time moves by 2 s and core.xml's
    ``dcterms:modified`` changes, which is what a no-edit Office re-save does to a workbook (design 4.1 #2).
    """
    a = dest / "sample.xlsx"
    b = dest / "sample-resaved.xlsx"
    _workbook(dt.datetime(2026, 9, 2, 9, 0, 0)).save(a)  # type: ignore[attr-defined]
    with zipfile.ZipFile(a) as src, zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "docProps/core.xml":
                pattern = rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)"
                data = re.sub(pattern, rb"\g<1>2026-09-03T17:30:00Z\g<2>", data)
            y, mo, d, h, mi, sec = info.date_time
            moved = zipfile.ZipInfo(info.filename, date_time=(y, mo, d, h, mi, min(sec + 2, 58)))
            moved.compress_type = zipfile.ZIP_DEFLATED
            out.writestr(moved, data)
    return a, b


def build_pptx(dest: Path) -> Path:
    """Write sample.pptx: two slides with titles, bullets and speaker notes."""
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for i, (title, bullets, notes) in enumerate(
        [
            ("Kickoff", ["Scope agreed", "Pricing pending"], "Mention the renewal date."),
            ("Next steps", ["Send the purchase order", "Book the review"], "Owner: finance."),
        ],
        start=1,
    ):
        slide = prs.slides.add_slide(layout)
        assert slide.shapes.title is not None
        slide.shapes.title.text = f"{i}. {title}"
        body = slide.placeholders[1]
        tf = body.text_frame  # type: ignore[attr-defined]
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        slide.notes_slide.notes_text_frame.text = notes  # type: ignore[union-attr]
    out = dest / "sample.pptx"
    prs.save(str(out))
    return out


def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf(dest: Path) -> Path:
    """Write a minimal two-page born-digital PDF (Helvetica text) with a correct xref table."""
    pages = [
        ["Quarterly Plan - page 1", "The purchase order was approved."],
        ["Budget - page 2", "North spend 100, forecast 120."],
    ]
    objs: list[bytes] = []
    n_pages = len(pages)
    # 1 catalog, 2 pages, 3 font, then (page, content) pairs
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, lines in enumerate(pages):
        content_lines = ["BT", "/F1 14 Tf", "72 720 Td", "18 TL"]
        for line in lines:
            content_lines.append(f"({_pdf_escape(line)}) Tj T*")
        content_lines.append("ET")
        stream = "\n".join(content_lines).encode()
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path = dest / "sample.pdf"
    path.write_bytes(bytes(out))
    return path


def build_eml(dest: Path) -> Path:
    """Write sample.eml: plain-text body plus one small CSV attachment."""
    msg = EmailMessage()
    msg["From"] = "Jane Doe <jane.doe@example.com>"
    msg["To"] = "Team <team@example.com>"
    msg["Subject"] = "Acme kickoff recap"
    msg["Date"] = format_datetime(dt.datetime(2026, 9, 14, 9, 12, tzinfo=dt.UTC))
    msg["Message-ID"] = "<kickoff-recap@example.com>"
    msg.set_content("Hi all,\n\nThe purchase order is approved. Renewal is in March.\n\nJane\n")
    msg.add_attachment(b"region,spend\nNorth,100\n", maintype="text", subtype="csv", filename="spend.csv")
    path = dest / "sample.eml"
    path.write_bytes(bytes(msg))
    return path


def build_teams_month(dest: Path) -> Path:
    """Write 2026-09.teams.json in the model.TEAMS_MONTH_SCHEMA shape."""
    doc = {
        "schema": "agentsync.teams-month/1",
        "team_id": "team-1",
        "team_name": "Acme",
        "channel_id": "19:abc@thread.tacv2",
        "channel_name": "General",
        "month": "2026-09",
        "messages": [
            {
                "id": "1001",
                "reply_to_id": None,
                "created": "2026-09-14T09:00:00Z",
                "last_modified": "2026-09-14T09:00:00Z",
                "from": "Jane Doe",
                "subject": "Kickoff",
                "body_html": "<p>Kickoff is <b>Monday</b>.</p>",
                "deleted": False,
                "attachments": [],
            },
            {
                "id": "1002",
                "reply_to_id": "1001",
                "created": "2026-09-14T09:05:00Z",
                "last_modified": "2026-09-14T09:05:00Z",
                "from": "John Roe",
                "subject": None,
                "body_html": "<p>I will send the purchase order.</p>",
                "deleted": False,
                "attachments": [],
            },
        ],
    }
    path = dest / "2026-09.teams.json"
    path.write_text(json.dumps(doc, sort_keys=True, indent=1) + "\n", encoding="utf-8")
    return path


def make_fixtures(dest: Path) -> dict[str, Path]:
    """Build every sample into ``dest`` (created if needed); return {file name: path}."""
    dest.mkdir(parents=True, exist_ok=True)
    files: list[Path] = [
        build_docx(dest),
        *build_xlsx(dest),
        build_pptx(dest),
        build_pdf(dest),
        build_eml(dest),
    ]
    files.append(build_teams_month(dest))
    md = dest / "sample.md"
    md.write_text("---\ntitle: Sample\n---\n" + MARKDOWN_SOURCE, encoding="utf-8")
    txt = dest / "sample.txt"
    txt.write_text("Plain text notes.\nSecond line with purchase order.\n", encoding="utf-8")
    html = dest / "sample.html"
    html.write_text(
        "<html><body><h1>Title</h1><p>Hello <em>world</em>.</p></body></html>\n", encoding="utf-8"
    )
    csv = dest / "sample.csv"
    csv.write_text("region,spend\nNorth,100\nSouth,80\n", encoding="utf-8")
    files += [md, txt, html, csv]
    return {p.name: p for p in sorted(files)}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.stderr.write("usage: make_fixtures.py OUTDIR\n")
        sys.exit(2)
    for name, path in make_fixtures(Path(sys.argv[1])).items():
        sys.stdout.write(f"{name}\t{path}\n")
