"""Builders for converter test inputs (generated at test time; nothing binary is committed).

Imported by the other ``test_convert_*`` modules; the two tests at the bottom only check the builders.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import re
import struct
import subprocess
import zipfile
import zlib
from collections.abc import Mapping
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from typing import Any, NamedTuple

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from agentsync.convert.ocr import OcrEngine
from agentsync.convert.pandoc import _bundled_pandoc
from test_ocr import fake_engine

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def png_bytes() -> bytes:
    """A valid 1x1 RGB PNG (no tIME chunk, so byte-stable)."""

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
        + chunk(b"IEND", b"")
    )


def pandoc_build(text: str, src_fmt: str, out: Path, *, cwd: Path | None = None) -> Path:
    """Write ``out`` from ``text`` with the bundled pandoc (the writer is chosen by the out suffix)."""
    subprocess.run(
        [str(_bundled_pandoc()), "-f", src_fmt, "-s", "-o", str(out)],
        input=text.encode("utf-8"),
        check=True,
        cwd=str(cwd or out.parent),
        capture_output=True,
    )
    return out


def build_docx_image(dest: Path) -> Path:
    """A docx with a figure (alt text) and an inline image."""
    (dest / "pic.png").write_bytes(png_bytes())
    md = "# Doc\n\n![a chart](pic.png)\n\nText ![inline](pic.png) here.\n"
    return pandoc_build(md, "markdown", dest / "image.docx", cwd=dest)


def build_docx_merged(dest: Path) -> Path:
    """A docx whose table has a merged (gridSpan) cell (design C13 fixture shape)."""
    html = (
        "<h1>Scope</h1><table><thead><tr><th>Region</th><th>Spend</th><th>Total</th></tr></thead><tbody>"
        '<tr><td colspan="2">West (merged)</td><td>300</td></tr><tr><td>East</td><td>5</td><td>6</td></tr>'
        "</tbody></table><h2>Notes</h2><p>closing paragraph</p>"
    )
    return pandoc_build(html, "html", dest / "merged.docx")


def make_zip(
    path: Path, parts: dict[str, bytes], *, date: tuple[int, int, int, int, int, int] | None = None
) -> Path:
    """Write a ZIP with exactly ``parts`` (insertion order), optionally stamping every entry's date."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in parts.items():
            info = zipfile.ZipInfo(name, date_time=date or (2026, 9, 1, 12, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zf.writestr(info, data)
    return path


def ole_encrypted(path: Path) -> Path:
    """Bytes shaped like an Office-encrypted package: OLE header + an EncryptedPackage directory entry."""
    body = OLE_MAGIC + b"\x00" * 504 + "EncryptedPackage".encode("utf-16-le") + b"\x00" * 1024
    path.write_bytes(body)
    return path


def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf(path: Path, pages: list[list[str]], *, encrypt: bool = False) -> Path:
    """A minimal born-digital PDF (Helvetica) with a correct xref; ``encrypt`` adds a Standard security
    handler whose user password is not empty (so opening needs a password)."""
    objs: list[bytes] = []
    n_pages = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, lines in enumerate(pages):
        content = ["BT", "/F1 12 Tf", "72 720 Td", "16 TL"]
        content += [f"({_pdf_escape(line)}) Tj T*" for line in lines]
        content.append("ET")
        stream = "\n".join(content).encode()
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    if encrypt:
        objs.append(
            b"<< /Filter /Standard /V 1 /R 2 /Length 40 /P -4 "
            b"/O <" + b"11" * 32 + b"> /U <" + b"22" * 32 + b"> >>"
        )
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    extra = ""
    if encrypt:
        extra = f" /Encrypt {len(objs)} 0 R /ID [<{'ab' * 16}> <{'ab' * 16}>]"
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R{extra} >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def pdf_text(s: str) -> str:
    """A PDF text string: a literal when printable ASCII, else UTF-16BE with a BOM in hex (how a viewer
    writes a comment that has an arrow, an accent or a line break in it)."""
    if s.isascii() and s.isprintable():
        return f"({_pdf_escape(s)})"
    return f"<FEFF{s.encode('utf-16-be').hex().upper()}>"


def build_annotated_pdf(path: Path, pages: list[tuple[list[str], list[str]]], *, leading: int = 12) -> Path:
    """A born-digital PDF whose pages carry annotations.

    A page is ``(text lines, annotation dictionary bodies)``.  Text is 12 pt Helvetica from x = 72; line
    ``i`` has its baseline at y = 720 - leading * i, and PDFium's loose character box runs from 2.5 below
    the baseline to 10.9 above it.  The default leading is single spacing, where a marked line touches its
    neighbours.  ``{N}`` in a body is the reference to annotation ``N`` of the same page (/IRT, /Popup).
    """
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # the page tree, written once the page numbers are known
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids: list[str] = []
    for lines, annots in pages:
        page_no = len(objs) + 1
        refs = [f"{page_no + 2 + i} 0 R" for i in range(len(annots))]
        ops = ["BT", "/F1 12 Tf", "72 720 Td", f"{leading} TL"]
        ops += [f"({_pdf_escape(line)}) Tj T*" for line in lines]
        stream = "\n".join([*ops, "ET"]).encode()
        kids.append(f"{page_no} 0 R")
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {page_no + 1} 0 R /Annots [{' '.join(refs)}] >>".encode()
        )
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        objs += [f"<< /Type /Annot {body.format(*refs)} >>".encode() for body in annots]
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()
    out = bytearray(b"%PDF-1.6\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def pdf_line_quad(line: int, *, left: int = 70, right: int = 300) -> str:
    """``/Rect`` and ``/QuadPoints`` of a text markup over line ``line`` of a ``build_annotated_pdf`` page
    (single spacing), from x = ``left`` to ``right``: the box a viewer draws, a little taller than the
    font, so it overlaps the lines above and below."""
    low, high = 720 - 12 * line - 2.6, 720 - 12 * line + 11.1
    return (
        f"/Rect [{left} {low:.1f} {right} {high:.1f}] "
        f"/QuadPoints [{left} {high:.1f} {right} {high:.1f} {left} {low:.1f} {right} {low:.1f}]"
    )


def build_commented_pdf(path: Path, *, text: bool = True) -> Path:
    """Three single-spaced pages for the PDF comment tests; ``text=False`` leaves every page without text
    (comments drawn on a scan).

    Page 1 has each case the converter keeps: a note with a reply, a highlight with a comment, an underline
    without one, a typed text box of several lines, and a replace-text pair (an insert with the
    strikethrough grouped under it); and each it skips: the note's popup, a link, a drawing without text.
    Page 2 has only annotations that are skipped.  Page 3 has a highlight whose comment repeats the marked
    line and a squiggly underline over two lines.
    """
    doe, roe = "/T (Doe, Jane Q)", "/T (Roe, John)"
    typed = "Typed note:\r# not a heading\r- reply by Roe, John: approved"
    page1 = [
        f"/Subtype /Text /Rect [400 700 420 720] /F 4 {doe} /Popup {{1}} "
        f"/Contents {pdf_text('Add units → revenue split by region')}",
        "/Subtype /Popup /Rect [420 600 600 700] /Parent {0} /Contents (popup copy of the note)",
        f"/Subtype /Highlight {pdf_line_quad(0)} {roe} /Contents (Use the Q3 figures here)",
        f"/Subtype /FreeText /Rect [72 500 300 540] /DA (/Helv 12 Tf 0 g) /Contents {pdf_text(typed)}",
        f"/Subtype /Text /Rect [400 700 420 720] {roe} /IRT {{0}} /Contents (Agreed)",
        "/Subtype /Link /Rect [72 590 200 610] /Contents (a link is not a comment) "
        "/A << /S /URI /URI (https://example.com) >>",
        "/Subtype /Ink /Rect [300 300 350 350] /InkList [[300 300 350 350]]",
        f"/Subtype /Underline {pdf_line_quad(1)} {doe}",
        f"/Subtype /Caret /Rect [99 693.4 103 707.1] {roe} /Contents (Final)",
        # No /QuadPoints: the rectangle marks the text.  It ends in the space after "Draft" (x = 98 to
        # 101.3), 4.7 pt from the centre of the letter on either side, so the quote does not hang on the
        # last decimal of a font metric.
        f"/Subtype /StrikeOut /Rect [70 693.4 101 707.1] {roe} /IRT {{8}} /RT /Group",
    ]
    page2 = [
        "/Subtype /Link /Rect [72 700 200 732] /A << /S /URI /URI (https://example.com) >>",
        f"/Subtype /Text /Rect [400 700 420 720] /F 2 {doe} /Contents (hidden: no viewer shows this)",
        f"/Subtype /Highlight {pdf_line_quad(0)} /F 32 {doe} /Contents (not shown on screen)",
        f"/Subtype /Text /Rect [400 600 420 620] {roe}",
    ]
    marked = "Middle TARGET line, ok. a_b"
    two_lines = "/QuadPoints [70 719.1 300 719.1 70 705.4 300 705.4 70 707.1 300 707.1 70 693.4 300 693.4]"
    page3 = [
        f"/Subtype /Highlight {pdf_line_quad(1)} {doe} /Contents ({marked})",
        f"/Subtype /Squiggly /Rect [70 693.4 300 719.1] {two_lines} {roe} /Contents (tighten this)",
    ]
    lines = [
        ["Contoso widget overview", "Quarterly totals by region", "Draft wording of the summary"],
        ["Second page: nothing on it is a comment."],
        ["Alpha gyp line above, jq.", marked, "Omega line below"],
    ]
    pages = [(ln if text else [], annots) for ln, annots in zip(lines, (page1, page2, page3), strict=True)]
    return build_annotated_pdf(path, pages)


class PdfPicture(NamedTuple):
    """One picture of a ``build_picture_pdf`` page: an 8-bit gray image whose every pixel is ``shade``.

    ``matrix`` is the PDF matrix it is drawn with (``"a b c d e f"``: a wide and d high from (e, f) when b
    and c are 0).  ``px`` is its stored size.  With ``form`` it is drawn inside a form XObject that is itself
    placed with that matrix.  ``pixels``, when given, are its gray pixels row by row in place of the one
    shade (a real picture of text, for the real helper).  Pictures of one shade, size and pixels are one
    image object."""

    shade: int
    matrix: str
    px: tuple[int, int] = (96, 64)
    form: str = ""
    pixels: bytes = b""


PICTURE_PAGE = (300, 200)
"""The page size of ``build_picture_pdf``, in points: small, so a page renders fast at 300 dpi."""


def page_picture(shade: int) -> PdfPicture:
    """A picture that covers a whole ``build_picture_pdf`` page: what a scanner writes."""
    return PdfPicture(shade, f"{PICTURE_PAGE[0]} 0 0 {PICTURE_PAGE[1]} 0 0")


def build_picture_pdf(path: Path, pages: list[tuple[list[str], list[PdfPicture]]]) -> Path:
    """A PDF whose pages hold text and pictures, for the OCR tests.

    A page is ``(text lines, pictures)``, ``PICTURE_PAGE`` points in size.  Text is 12 pt Helvetica from
    x = 20; line ``i`` has its baseline at y = 180 - 14 * i.  A page with no text line and a ``page_picture``
    is a scanned page.  Rendered, such a page is the picture's shade from corner to corner, and a page with
    no picture is white: the fake OCR helper of the converter tests tells images apart by that."""
    width, height = PICTURE_PAGE
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # the page tree, written once the page numbers are known
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    images: dict[tuple[int, tuple[int, int], bytes], int] = {}
    for _lines, pictures in pages:
        for pic in pictures:
            if (pic.shade, pic.px, pic.pixels) not in images:
                data = zlib.compress(pic.pixels or bytes([pic.shade]) * (pic.px[0] * pic.px[1]))
                objs.append(
                    f"<< /Type /XObject /Subtype /Image /Width {pic.px[0]} /Height {pic.px[1]} "
                    f"/ColorSpace /DeviceGray /BitsPerComponent 8 /Filter /FlateDecode "
                    f"/Length {len(data)} >>\nstream\n".encode()
                    + data
                    + b"\nendstream"
                )
                images[pic.shade, pic.px, pic.pixels] = len(objs)
    kids: list[str] = []
    for lines, pictures in pages:
        ops = ["BT", "/F1 12 Tf", "20 180 Td", "14 TL"]
        ops += [f"({_pdf_escape(line)}) Tj T*" for line in lines]
        ops.append("ET")
        xobjects: list[str] = []
        for n, pic in enumerate(pictures):
            image = images[pic.shade, pic.px, pic.pixels]
            if pic.form:  # the picture is the form's content; the form is what the page draws
                inner = f"q {pic.matrix} cm /Im Do Q".encode()
                head = (
                    f"<< /Type /XObject /Subtype /Form /BBox [0 0 {width} {height}] "
                    f"/Resources << /XObject << /Im {image} 0 R >> >> /Length {len(inner)} >>"
                )
                objs.append(head.encode() + b"\nstream\n" + inner + b"\nendstream")
                image, matrix = len(objs), pic.form
            else:
                matrix = pic.matrix
            xobjects.append(f"/X{n} {image} 0 R")
            ops.append(f"q {matrix} cm /X{n} Do Q")
        stream = "\n".join(ops).encode()
        page_no = len(objs) + 1
        kids.append(f"{page_no} 0 R")
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] /Resources << /Font "
            f"<< /F1 3 0 R >> /XObject << {' '.join(xobjects)} >> >> /Contents {page_no + 1} 0 R >>".encode()
        )
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()
    out = bytearray(b"%PDF-1.6\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


# The fake OCR helper of tests/test_ocr.py reads what a test wrote into an image file.  A page PDFium rendered
# and a picture it decoded hold only pixels, so this hook tells images apart by the gray of their first
# pixel and answers from a table beside the helper.
SHADE_HOOK = """\
if "--" in args:
    import struct, zlib
    table = json.loads(Path(sys.argv[0] + ".shades").read_text())
    results = []
    for index, p in enumerate(args[args.index("--") + 1:]):
        data = Path(p).read_bytes()
        width, height = struct.unpack(">II", data[16:24])
        said = table.get(str(zlib.decompressobj().decompress(data[data.index(b"IDAT") + 4:], 2)[1]), [])
        if said == "fail":
            sys.exit(3)
        item = {"index": index, "frame": 0, "frames": 1, "width": width, "height": height, "lines": [],
                "skipped": False}
        if isinstance(said, dict):
            item.update(said)
        else:
            item["lines"] = [{"text": t, "confidence": 1.0, "x": 0.1, "y": 0.1 + 0.06 * n, "w": 0.6,
                              "h": 0.05} for n, t in enumerate(said)]
        results.append(item)
    print(json.dumps({"results": results}))
    sys.exit(0)
"""


def shade_engine(folder: Path, said: Mapping[int, Any]) -> OcrEngine:
    """An OCR engine for ``build_picture_pdf`` files: its fake helper 'reads' a PNG by the shade of its
    first pixel.  ``said`` maps a shade to the lines read there, to the fields of an image the helper
    reports instead (``{"error": "recognition failed"}``, ``{"skipped": True}``), or to ``"fail"``, which
    makes a run that holds such an image exit 3.  A shade it does not name has no text.  255 is a page with
    no picture on it."""
    engine = fake_engine(folder)
    text = engine.helper.read_text(encoding="utf-8")
    at = 'if "--version" in args:'
    assert at in text
    engine.helper.write_text(text.replace(at, SHADE_HOOK + at, 1), encoding="utf-8")
    Path(str(engine.helper) + ".shades").write_text(json.dumps(said), encoding="utf-8")
    return engine


def build_xlsx_rich(path: Path, *, rows: int = 10) -> Path:
    """Workbook with formulas (no cached values), a comment, a chart, a hidden sheet, a defined name, odd
    values (integral float, date, bool, pipe/newline text) and ``rows`` data rows."""

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(["Item", "Qty", "Price", "Total", "Misc"])
    for i in range(rows):
        r = i + 2
        ws.append([f"item-{i:03d}", i + 1, 2.5, f"=B{r}*C{r}", None])
    ws["E2"] = 3.0
    ws["E3"] = dt.datetime(2026, 9, 1)
    ws["E4"] = True
    ws["E5"] = "pipe|text\nsecond line"
    ws["B2"].comment = Comment("check this quantity", "Jane")
    chart = BarChart()
    chart.title = "Qty by item"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=rows + 1), titles_from_data=True)
    ws.add_chart(chart, "G2")
    hidden = wb.create_sheet("Secret")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "hidden value"
    wb.defined_names.add(DefinedName("TaxRate", attr_text="Data!$C$2"))
    wb.save(path)
    return path


def with_cached_values(src: Path, dest: Path, sheet_part: str, values: dict[str, str]) -> Path:
    """Copy an openpyxl workbook, inserting ``<v>`` cached results for the given formula cells (what Excel
    writes on save; openpyxl never does)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == sheet_part:
                text = data.decode("utf-8")
                for ref, value in values.items():
                    pattern = re.compile(rf'(<c r="{ref}"[^>]*>)(<f>[^<]*</f>)(<v\s*/>|<v></v>)?')
                    text, n = pattern.subn(lambda m, v=value: f"{m.group(1)}{m.group(2)}<v>{v}</v>", text)
                    assert n == 1, f"formula cell {ref} not found in {sheet_part}"
                data = text.encode("utf-8")
            zout.writestr(info, data)
    return dest


def build_pptx_rich(path: Path) -> Path:
    """Deck: text boxes out of XML order, a table with a merged cell, a picture with alt text, a chart, a
    group shape, a hidden slide."""

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    assert s1.shapes.title is not None
    s1.shapes.title.text = "Numbers"
    right = s1.shapes.add_textbox(Inches(5), Inches(2), Inches(2), Inches(1))
    right.text_frame.text = "RIGHT box"
    left = s1.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))
    left.text_frame.text = "LEFT box"
    table = s1.shapes.add_table(3, 2, Inches(1), Inches(3), Inches(4), Inches(1.5)).table
    for (r, c), text in {
        (0, 0): "k",
        (0, 1): "v",
        (1, 0): "rate",
        (1, 1): "4.2%",
        (2, 0): "merged row",
    }.items():
        table.cell(r, c).text = text
    table.cell(2, 0).merge(table.cell(2, 1))
    pic = s1.shapes.add_picture(io.BytesIO(png_bytes()), Inches(8), Inches(5))
    pic._element.nvPicPr.cNvPr.set("descr", "Company logo")
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Revenue", (10.0, 12.5))
    s1.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(5), Inches(4), Inches(2), data)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    group = s2.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "inside group"
    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    assert s3.shapes.title is not None
    s3.shapes.title.text = "Hidden one"
    s3._element.set("show", "0")
    prs.save(str(path))
    return path


def eml_bytes(
    *,
    subject: str = "Hello",
    plain: str | None = "Body text.\n",
    html: str | None = None,
    attachments: tuple[tuple[str, bytes, str], ...] = (),
) -> bytes:
    """An RFC 822 message with optional plain/html alternatives and attachments."""
    msg = EmailMessage()
    msg["From"] = "Jane Doe <jane.doe@example.com>"
    msg["To"] = "Team <team@example.com>"
    msg["Cc"] = "Boss <boss@example.com>"
    msg["Subject"] = subject
    msg["Date"] = format_datetime(dt.datetime(2026, 9, 14, 9, 12, tzinfo=dt.UTC))
    msg["Message-ID"] = "<m1@example.com>"
    msg["In-Reply-To"] = "<m0@example.com>"
    msg["References"] = "<m0@example.com>"
    if plain is not None:
        msg.set_content(plain)
        if html is not None:
            msg.add_alternative(html, subtype="html")
    elif html is not None:
        msg.set_content(html, subtype="html")
    for filename, data, ctype in attachments:
        maintype, subtype = ctype.split("/")
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
    return bytes(msg)


def teams_doc(messages: list[dict[str, object]], *, schema: str = "agentsync.teams-month/1") -> str:
    """A TEAMS_MONTH_SCHEMA JSON document."""

    doc = {
        "schema": schema,
        "team_id": "team-1",
        "team_name": "Acme",
        "channel_id": "19:abc@thread.tacv2",
        "channel_name": "General",
        "month": "2026-09",
        "messages": messages,
    }
    return json.dumps(doc, sort_keys=True, indent=1) + "\n"


def teams_msg(mid: str, created: str, body: str, **kw: object) -> dict[str, object]:
    """One message dict with defaults."""
    msg: dict[str, object] = {
        "id": mid,
        "reply_to_id": None,
        "created": created,
        "last_modified": created,
        "from": "Jane Doe",
        "subject": None,
        "body_html": body,
        "deleted": False,
        "attachments": [],
    }
    msg.update(kw)
    return msg


def test_builders_produce_valid_containers(tmp_path: Path) -> None:
    assert zipfile.is_zipfile(build_docx_image(tmp_path))
    assert zipfile.is_zipfile(build_docx_merged(tmp_path))
    assert zipfile.is_zipfile(build_xlsx_rich(tmp_path / "r.xlsx"))
    assert zipfile.is_zipfile(build_pptx_rich(tmp_path / "r.pptx"))
    assert build_pdf(tmp_path / "a.pdf", [["x"]]).read_bytes().startswith(b"%PDF-1.4")
    logo = PdfPicture(40, "96 0 0 64 150 20")
    pictured = build_picture_pdf(tmp_path / "p.pdf", [(["x"], [logo]), ([], [logo, page_picture(90)])])
    assert pictured.read_bytes().startswith(b"%PDF-1.6") and pictured.read_bytes().count(b"/Image") == 2
    commented = build_commented_pdf(tmp_path / "c.pdf").read_bytes()
    assert commented.startswith(b"%PDF-1.6") and commented.count(b"/Type /Annot") == 16
    assert (
        b") Tj" in commented
        and b") Tj" not in build_commented_pdf(tmp_path / "s.pdf", text=False).read_bytes()
    )


def test_png_is_stable() -> None:
    assert png_bytes() == png_bytes()
    assert png_bytes().startswith(b"\x89PNG")
